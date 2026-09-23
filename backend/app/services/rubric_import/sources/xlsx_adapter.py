"""Excel → 原文单元（解析重构方案 §4 第 1 层）。只读取，不做任何语义判断。

合并区域只登记左上角单元；行视图 ``rows`` 与原 ``_rows_with_merged_values`` 一致，
把合并区域展开为同值，供抽取器沿用既有的按行解析逻辑。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from backend.app.services.rubric_import.sources.units import SourceLedger
from backend.app.services.rubric_import.sources.units import SourceUnit


def cell_unit_id(sheet_title: str, row: int, column: int) -> str:
    return f"xlsx:{sheet_title}!R{row}C{column}"


def _text(value) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


@dataclass
class SheetView:
    title: str
    rows: list[tuple] = field(default_factory=list)
    _units: dict[tuple[int, int], str] = field(default_factory=dict)

    def unit_at(self, row_index: int, column_index: int) -> str | None:
        """0 基行列坐标 → 单元 id；合并区域内任意位置都指向左上角单元。"""

        return self._units.get((row_index, column_index))

    def row_units(self, row_index: int) -> list[str | None]:
        width = len(self.rows[row_index]) if row_index < len(self.rows) else 0
        return [self.unit_at(row_index, column) for column in range(width)]


def load_xlsx(data: bytes, ledger: SourceLedger, *, doc_id: str, doc_role: str) -> list[SheetView]:
    try:
        workbook = load_workbook(BytesIO(data), data_only=True)
    except Exception as exc:  # openpyxl 对非法文件抛多种异常
        raise ValueError("无法读取 Excel 文件，请确认为 .xlsx / .xlsm 格式") from exc

    views = []
    for sheet in workbook.worksheets:
        anchors: dict[tuple[int, int], tuple[int, int, str]] = {}
        for merged in sheet.merged_cells.ranges:
            ref = f"{get_column_letter(merged.min_col)}{merged.min_row}:{get_column_letter(merged.max_col)}{merged.max_row}"
            for row in range(merged.min_row, merged.max_row + 1):
                for column in range(merged.min_col, merged.max_col + 1):
                    anchors[(row, column)] = (merged.min_row, merged.min_col, ref)
        view = SheetView(title=sheet.title)
        for row_cells in sheet.iter_rows():
            values = []
            for cell in row_cells:
                anchor = anchors.get((cell.row, cell.column))
                if anchor is not None:
                    anchor_row, anchor_col, ref = anchor
                    value = sheet.cell(anchor_row, anchor_col).value
                else:
                    anchor_row, anchor_col, ref = cell.row, cell.column, None
                    value = cell.value
                values.append(value)
                if value is None or not _text(value):
                    continue
                unit_id = cell_unit_id(sheet.title, anchor_row, anchor_col)
                if (anchor_row, anchor_col) == (cell.row, cell.column):
                    context = {"sheet": sheet.title, "row": cell.row, "col": cell.column,
                               "value_type": "num" if isinstance(value, (int, float)) else "text"}
                    if ref:
                        context["merged"] = ref
                    ledger.register(SourceUnit(unit_id, doc_id, doc_role, "cell", _text(value), context))
                view._units[(cell.row - 1, cell.column - 1)] = unit_id
            view.rows.append(tuple(values))
        views.append(view)
    return views
