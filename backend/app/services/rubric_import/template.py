from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.styles import Border
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.styles import Side
from openpyxl.worksheet.datavalidation import DataValidation


HEADERS = ["编号", "评分项", "分值", "评分说明", "证据提示", "扣分规则", "顺序"]

EXAMPLE_ROWS = [
    ["C01", "选题意义", 10, "选题具有理论意义、现实意义，问题明确。", "绪论；研究背景；研究意义", "意义表述笼统扣 1-3 分", 1],
    ["C02", "文献综述", 15, "文献覆盖充分，能归纳研究现状和研究空白。", "文献综述；国内外研究现状；相关工作", "文献覆盖不足扣 2-5 分", 2],
    ["C03", "研究方法", 20, "方法合理，实验设计清楚，数据来源可靠。", "研究方法；实验设计；数据来源", "方法说明不清晰扣 2-6 分", 3],
    ["C04", "论文创新性", 15, "有明确创新点或改进贡献。", "创新点；贡献；改进", "创新性不足扣 2-5 分", 4],
    ["C05", "论证与分析", 20, "论证逻辑完整，数据分析能支撑结论。", "实验结果；结果分析；讨论", "论证链条不完整扣 2-6 分", 5],
    ["C06", "写作规范", 10, "结构完整，格式、图表、语言规范。", "摘要；关键词；目录；结论", "结构或格式缺项扣 1-4 分", 6],
    ["C07", "参考文献", 10, "参考文献数量、格式和引用规范。", "参考文献；引用", "参考文献数量或格式不足扣 1-4 分", 7],
]

INSTRUCTIONS = [
    ["填写说明"],
    ["1. 必填列：评分项、分值。其他列可选，但建议完整填写以提升证据召回质量。"],
    ["2. 证据提示和扣分规则可以用中文分号、英文分号、顿号或换行分隔。"],
    ["3. 合计、总分、总计行会被系统自动忽略。"],
    ["4. 导入后系统会按分值自动汇总总分，并创建草稿评分标准。"],
]


def build_rubric_import_template():
    workbook = Workbook()
    rules = workbook.active
    rules.title = "评分规则"
    rules.append(HEADERS)
    for row in EXAMPLE_ROWS:
        rules.append(row)
    rules.append(["总计", "", "=SUM(C2:C8)", "", "", "", ""])

    _style_rules_sheet(rules)
    _add_validation(rules)
    _add_instructions_sheet(workbook)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _style_rules_sheet(sheet):
    header_fill = PatternFill("solid", fgColor="DCEBFF")
    total_fill = PatternFill("solid", fgColor="F8FAFC")
    border = Border(bottom=Side(style="thin", color="D7DEE8"))

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:G8"
    sheet.column_dimensions["A"].width = 12
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 10
    sheet.column_dimensions["D"].width = 42
    sheet.column_dimensions["E"].width = 34
    sheet.column_dimensions["F"].width = 34
    sheet.column_dimensions["G"].width = 10

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="17202A")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border

    for cell in sheet[sheet.max_row]:
        cell.fill = total_fill
        cell.font = Font(bold=True)


def _add_validation(sheet):
    score_validation = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=False)
    score_validation.error = "分值必须是大于 0 的数字。"
    score_validation.errorTitle = "分值格式错误"
    sheet.add_data_validation(score_validation)
    score_validation.add("C2:C200")

    order_validation = DataValidation(type="whole", operator="greaterThan", formula1="0", allow_blank=True)
    order_validation.error = "顺序应填写正整数。"
    order_validation.errorTitle = "顺序格式错误"
    sheet.add_data_validation(order_validation)
    order_validation.add("G2:G200")


def _add_instructions_sheet(workbook):
    sheet = workbook.create_sheet("填写说明")
    for row in INSTRUCTIONS:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 92
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A1"].fill = PatternFill("solid", fgColor="DCEBFF")
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
