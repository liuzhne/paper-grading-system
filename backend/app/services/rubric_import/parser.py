import re
from dataclasses import dataclass, field
from io import BytesIO

from docx import Document
from openpyxl import load_workbook

from backend.app.services.document_parser.format_resolver import resolve_default_format
from backend.app.services.rubric_import.compiler import annotations_for_criterion
from backend.app.services.rubric_import.compiler import compile_criterion_rules
from backend.app.services.rubric_import.docx_comments import parse_comments


HEADER_ALIASES = {
    "code": ["编号", "指标编号", "评分项编号", "代码", "code", "criterion_code"],
    "name": ["评分项", "评分指标", "指标", "评价项目", "项目", "name", "criterion"],
    "max_score": ["分值", "满分", "分数", "最高分", "权重分", "max_score", "score", "points"],
    "weight": ["权重", "weight"],
    "description": ["说明", "评分说明", "评价标准", "评分标准", "标准说明", "描述", "description"],
    "evidence_hints": ["依据", "证据", "证据提示", "章节依据", "相关章节", "关键词", "evidence_hints"],
    "deduction_rules": ["扣分规则", "扣分点", "扣分说明", "扣分原因", "deduction_rules"],
    "display_order": ["顺序", "排序", "display_order"],
    "criterion_type": ["类型", "判定类型", "评分类型", "判定方式", "type"],
    "applies_to": ["适用范围", "适用章节", "作用范围", "范围", "applies_to"],
    "rubric_levels": ["分档", "档位", "等级标准", "评分档次", "rubric_levels"],
    "dimension": ["维度", "评价维度", "评分维度", "所属维度", "dimension"],
    "sub_checks": ["子检查", "子项", "子检查项", "混合子项", "子项检查", "sub_checks"],
}

TEMPLATE_HINTS = [
    "中文摘要",
    "英文摘要",
    "关键词",
    "目录",
    "绪论",
    "研究背景",
    "研究意义",
    "文献综述",
    "国内外研究现状",
    "相关工作",
    "研究方法",
    "实验设计",
    "数据来源",
    "结果分析",
    "讨论",
    "创新点",
    "结论",
    "参考文献",
    "致谢",
]


@dataclass
class ImportedCriterion:
    code: str
    name: str
    max_score: float
    weight: float | None
    description: str | None
    evidence_hints: list[str]
    deduction_rules: list[str]
    display_order: int
    criterion_type: str = "llm_judgment"
    scoring_mode: str = "llm_direct"
    applies_to: str = "global"
    rubric_levels: list = field(default_factory=list)
    sub_checks: list = field(default_factory=list)
    dimension: str | None = None
    deduction_rules_structured: list = field(default_factory=list)


@dataclass
class RubricImport:
    total_score: float
    criteria: list[ImportedCriterion]
    template_summary: dict
    warnings: list[str]
    format_spec: dict = field(default_factory=dict)


def parse_rubric_files(rules_bytes: bytes, template_bytes: bytes | None = None, scorer=None):
    warnings = []
    template_summary = (
        parse_word_template(template_bytes)
        if template_bytes
        else {"section_titles": [], "hints": [], "paragraph_count": 0, "format_spec": {}, "annotations": []}
    )
    criteria, excel_warnings = parse_excel_rules(rules_bytes)
    warnings.extend(excel_warnings)
    enriched = [_enrich_with_template(item, template_summary) for item in criteria]
    # §5 编译：把扣分规则编译成结构化规则（Excel 显式优先，否则小模型归一化批注/自由文本）。
    annotations = template_summary.get("annotations", [])
    for criterion in enriched:
        criterion.deduction_rules_structured = compile_criterion_rules(
            criterion, annotations_for_criterion(criterion, annotations), scorer
        )
    total_score = round(sum(item.max_score for item in enriched), 2)
    return RubricImport(
        total_score=total_score,
        criteria=enriched,
        template_summary=template_summary,
        warnings=warnings,
        format_spec=template_summary.get("format_spec") or {},
    )


def parse_word_template(template_bytes):
    document = Document(BytesIO(template_bytes))
    paragraphs = []
    section_titles = []

    for paragraph in document.paragraphs:
        text = _normalize(paragraph.text)
        if not text:
            continue
        paragraphs.append(text)
        style_name = getattr(paragraph.style, "name", "") or ""
        if _looks_like_template_title(text, style_name):
            section_titles.append(text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = _normalize(cell.text)
                if not text:
                    continue
                paragraphs.append(text)
                if _looks_like_template_title(text, ""):
                    section_titles.append(text)

    all_text = "\n".join(paragraphs)
    hints = [hint for hint in TEMPLATE_HINTS if hint in all_text]
    for title in section_titles:
        for token in re.findall(r"[\u4e00-\u9fffA-Za-z0-9]{2,}", title):
            if token not in hints and len(token) <= 12:
                hints.append(token)

    return {
        "section_titles": _dedupe(section_titles)[:30],
        "hints": _dedupe(hints)[:40],
        "paragraph_count": len(paragraphs),
        "format_spec": resolve_default_format(template_bytes),
        "annotations": parse_comments(template_bytes),
    }


def parse_excel_rules(rules_bytes):
    workbook = load_workbook(BytesIO(rules_bytes), data_only=True)
    warnings = []
    criteria = []

    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        header_index, mapping = _find_header(rows)
        if header_index is None:
            warnings.append("工作表 %s 未识别到评分规则表头，已跳过。" % sheet.title)
            continue
        for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            criterion = _criterion_from_row(row, mapping, len(criteria) + 1)
            if criterion is None:
                continue
            criteria.append(criterion)
        if criteria:
            break

    if not criteria:
        raise ValueError("Excel 未解析到有效评分项，请确认包含评分项名称和分值列。")
    return criteria, warnings


def _find_header(rows):
    for index, row in enumerate(rows[:15]):
        normalized = [_normalize(cell) for cell in row]
        mapping = {}
        for field, aliases in HEADER_ALIASES.items():
            for alias in aliases:
                alias_norm = alias.lower()
                for col, value in enumerate(normalized):
                    if value and alias_norm in value.lower():
                        mapping[field] = col
                        break
                if field in mapping:
                    break
        if "name" in mapping and "max_score" in mapping:
            return index, mapping
    return None, {}


def _criterion_from_row(row, mapping, order):
    name = _value(row, mapping.get("name"))
    if not name or name in {"合计", "总分", "总计"}:
        return None
    max_score = _parse_score(_value(row, mapping.get("max_score")))
    if max_score is None:
        return None

    code = _value(row, mapping.get("code")) or "C%02d" % order
    description = _value(row, mapping.get("description"))
    evidence_hints = _split_items(_value(row, mapping.get("evidence_hints")))
    deduction_rules = _split_items(_value(row, mapping.get("deduction_rules")))
    weight = _parse_score(_value(row, mapping.get("weight"))) if "weight" in mapping else None
    parsed_order = _parse_score(_value(row, mapping.get("display_order")))
    display_order = int(parsed_order if parsed_order is not None else order)
    criterion_type = _parse_type(_value(row, mapping.get("criterion_type")))
    applies_to = _parse_applies_to(_value(row, mapping.get("applies_to")))
    rubric_levels = _parse_bands(_value(row, mapping.get("rubric_levels")))
    scoring_mode = "banded" if rubric_levels else "llm_direct"
    dimension = _value(row, mapping.get("dimension"))
    sub_checks = _parse_sub_checks(_raw_value(row, mapping.get("sub_checks")))
    if sub_checks:
        criterion_type = "hybrid"  # 提供子检查即启用混合制（设计§2/§6.3）
    return ImportedCriterion(
        code=str(code),
        name=str(name),
        max_score=max_score,
        weight=weight,
        description=description,
        evidence_hints=evidence_hints,
        deduction_rules=deduction_rules,
        display_order=display_order,
        criterion_type=criterion_type,
        scoring_mode=scoring_mode,
        applies_to=applies_to,
        rubric_levels=rubric_levels,
        sub_checks=sub_checks,
        dimension=dimension,
    )


def _enrich_with_template(criterion, template_summary):
    hints = list(criterion.evidence_hints)
    matched = _match_template_hints(criterion, template_summary.get("hints", []))
    for hint in matched:
        if hint not in hints:
            hints.append(hint)

    if matched:
        template_note = "Word 模板解析提示：建议重点查看 %s。" % "、".join(matched[:6])
        description = "%s\n%s" % (criterion.description, template_note) if criterion.description else template_note
    else:
        description = criterion.description

    return ImportedCriterion(
        code=criterion.code,
        name=criterion.name,
        max_score=criterion.max_score,
        weight=criterion.weight,
        description=description,
        evidence_hints=hints,
        deduction_rules=criterion.deduction_rules,
        display_order=criterion.display_order,
        criterion_type=criterion.criterion_type,
        scoring_mode=criterion.scoring_mode,
        applies_to=criterion.applies_to,
        rubric_levels=criterion.rubric_levels,
        sub_checks=criterion.sub_checks,
        dimension=criterion.dimension,
        deduction_rules_structured=criterion.deduction_rules_structured,
    )


def _match_template_hints(criterion, template_hints):
    text = "%s %s %s" % (criterion.name, criterion.description or "", " ".join(criterion.evidence_hints))
    matched = []
    for hint in template_hints:
        if hint in text:
            matched.append(hint)
    rules = [
        ("文献", ["文献综述", "国内外研究现状", "相关工作"]),
        ("方法", ["研究方法", "实验设计", "数据来源"]),
        ("创新", ["创新点"]),
        ("写作", ["中文摘要", "英文摘要", "关键词", "目录", "结论"]),
        ("规范", ["中文摘要", "英文摘要", "关键词", "目录", "结论"]),
        ("参考文献", ["参考文献"]),
        ("选题", ["绪论", "研究背景", "研究意义"]),
        ("意义", ["绪论", "研究背景", "研究意义"]),
        ("论证", ["结果分析", "讨论", "结论"]),
        ("分析", ["结果分析", "讨论"]),
    ]
    for keyword, hints in rules:
        if keyword in text:
            for hint in hints:
                if hint in template_hints and hint not in matched:
                    matched.append(hint)
    return matched[:8]


def _looks_like_template_title(text, style_name):
    if "Heading" in style_name or "标题" in style_name:
        return True
    if len(text) > 40:
        return False
    if re.match(r"^(第[一二三四五六七八九十\d]+[章节]|[一二三四五六七八九十\d]+[、.．])", text):
        return True
    return any(hint == text or hint in text for hint in TEMPLATE_HINTS)


def _value(row, index):
    if index is None or index >= len(row):
        return None
    value = row[index]
    if value is None:
        return None
    return _normalize(value)


def _raw_value(row, index):
    """取原始单元格文本，**保留换行**（子检查按行分隔，不能被空白折叠）。"""
    if index is None or index >= len(row):
        return None
    value = row[index]
    return None if value is None else str(value)


def _parse_type(value):
    if not value:
        return "llm_judgment"
    text = str(value)
    if any(token in text for token in ["确定", "自动", "规则", "deterministic"]):
        return "deterministic"
    if any(token in text for token in ["混合", "hybrid"]):
        return "hybrid"
    return "llm_judgment"


def _parse_sub_checks(value):
    """解析混合制子检查列。每行一个子检查，字段以 `|`（或 `｜`）分隔：`名称 | 类型 | 分值`。

    类型缺省为语义判断（llm_judgment），含"确定/规则/自动/deterministic"则为确定性子检查（不调 LLM）。
    产出引擎可消费的 `{kind,name,criteria,max_points}` 列表（见 engine._make_sub_criterion）。
    """
    if not value:
        return []
    items = []
    for line in re.split(r"[\n;；]+", str(value)):
        line = line.strip()
        if not line:
            continue
        parts = [part.strip() for part in re.split(r"[|｜]", line)]
        name = parts[0] if parts else ""
        if not name:
            continue
        kind = _parse_sub_kind(parts[1]) if len(parts) > 1 else "llm_judgment"
        max_points = _parse_score(parts[2]) if len(parts) > 2 else 0.0
        items.append(
            {"kind": kind, "name": name, "criteria": name, "max_points": float(max_points or 0)}
        )
    return items


def _parse_sub_kind(value):
    if value and any(token in str(value) for token in ["确定", "规则", "自动", "deterministic", "det"]):
        return "deterministic"
    return "llm_judgment"


def _parse_applies_to(value):
    if not value:
        return "global"
    text = _normalize(value)
    if text in {"全局", "整体", "全文", "通用", "global"}:
        return "global"
    return text


def _parse_bands(value):
    if not value:
        return []
    text = str(value)
    pairs = re.findall(r"([一-鿿A-Za-z]+)\s*[:：]?\s*(\d+(?:\.\d+)?)", text)
    if pairs:
        return [{"label": label, "points": float(points)} for label, points in pairs]
    return [{"raw": _normalize(text)}]


def _parse_score(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"\d+(?:\.\d+)?", str(value))
    if not match:
        return None
    return float(match.group(0))


def _split_items(value):
    if not value:
        return []
    parts = re.split(r"[\n\r;；、]+", str(value))
    return [part.strip() for part in parts if part and part.strip()]


def _normalize(value):
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def _dedupe(items):
    result = []
    for item in items:
        if item and item not in result:
            result.append(item)
    return result
