from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.styles import Border
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.styles import Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session
from sqlalchemy.orm import selectinload

from backend.app.core.config import settings
from backend.app.db.models import GradingBatch
from backend.app.db.models import Paper
from backend.app.db.models import ReviewLog
from backend.app.db.models import ScoreItem
from backend.app.db.models import ScoringRun
from backend.app.db.models import SpreadsheetWriteLog
from backend.app.services.scoring.profiles.registry import get_workbench_profile
from backend.app.services.scoring.profiles.registry import profile_key_for_batch
from backend.app.services.scoring.profiles.thesis import ThesisProfile


# 身份列由文档类型决定（论文：学号、姓名…；通用文档：文件名、标题…），其余列所有类型一致。
# 列宽、换行与数字格式按表头名称配置：论文导出的列、顺序、宽度与改造前完全一样。
_SUMMARY_HEAD = (("批次名称", "batch_name"),)
_SUMMARY_TAIL = (
    ("评分标准版本", "rubric_version"),
    ("AI 总分", "ai_total"),
    ("最终总分", "final_total"),
    ("等级", "grade"),
    ("主要扣分点", "main_deductions"),
    ("是否人工修改", "changed"),
    ("是否需要复核", "need_manual_review"),
    ("评分时间", "finished_at"),
    ("复核教师", "reviewer"),
    ("复核意见", "review_notes"),
    ("报告链接", "report_link"),
)
_DETAIL_TAIL = (
    ("评分项", "criterion_name"),
    ("满分", "max_score"),
    ("AI 得分", "ai_score"),
    ("最终得分", "final_score"),
    ("扣分原因", "deductions"),
    ("原文依据", "evidence_quotes"),
    ("依据位置", "evidence_locations"),
    ("修改建议", "suggestion"),
    ("置信度", "confidence"),
)

_WIDTHS = {
    "批次名称": 22, "学号": 16, "姓名": 14, "学院": 18, "专业": 20, "论文题目": 38,
    "文件名": 30, "标题": 38, "归属单位": 18, "类别": 20,
    "评分标准版本": 18, "AI 总分": 12, "最终总分": 12, "等级": 10, "主要扣分点": 48,
    "是否人工修改": 14, "是否需要复核": 14, "评分时间": 20, "复核教师": 18,
    "复核意见": 40, "报告链接": 34,
    "评分项": 28, "满分": 11, "AI 得分": 11, "最终得分": 11, "扣分原因": 44,
    "原文依据": 64, "依据位置": 38, "修改建议": 48, "置信度": 12,
}
_NO_WRAP_SUMMARY = {"学号", "姓名", "评分标准版本", "AI 总分", "最终总分", "等级",
                    "是否人工修改", "是否需要复核", "评分时间", "复核教师"}
_NO_WRAP_DETAIL = {"学号", "姓名", "满分", "AI 得分", "最终得分", "置信度"}
_SCORE_HEADERS = {"AI 总分", "最终总分", "满分", "AI 得分", "最终得分"}
_DATE_HEADERS = {"评分时间"}
_CONFIDENCE_HEADERS = {"置信度"}

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_ALT_ROW_FILL = PatternFill("solid", fgColor="F3F8FC")
_THIN_BORDER = Border(bottom=Side(style="thin", color="D9E2F3"))


def summary_columns(profile) -> list[tuple[str, str]]:
    """总分表的 (表头, 投影字段)。"""

    return [*_SUMMARY_HEAD, *profile.spreadsheet_summary_identity, *_SUMMARY_TAIL]


def detail_columns(profile) -> list[tuple[str, str]]:
    """评分明细表的 (表头, 投影字段)。"""

    return [*profile.spreadsheet_detail_identity, *_DETAIL_TAIL]


def batch_profile(batch):
    return get_workbench_profile(profile_key_for_batch(batch))


#: 论文的表头（在线表格写入器与旧调用方按名称引用）。
SUMMARY_HEADERS = [header for header, _ in summary_columns(ThesisProfile)]
DETAIL_HEADERS = [header for header, _ in detail_columns(ThesisProfile)]


def _layout(headers, no_wrap):
    def columns(names):
        return {index for index, header in enumerate(headers, start=1) if header in names}

    return {
        "widths": [_WIDTHS.get(header, 20) for header in headers],
        "wrap_columns": {
            index for index, header in enumerate(headers, start=1) if header not in no_wrap
        },
        "score_columns": columns(_SCORE_HEADERS),
    }


def export_batch_excel(db: Session, batch_id: str):
    batch = db.scalar(
        select(GradingBatch)
        .where(GradingBatch.id == batch_id)
        .options(selectinload(GradingBatch.rubric), selectinload(GradingBatch.rubric_version))
    )
    if batch is None:
        raise ValueError("batch not found")
    profile = batch_profile(batch)
    summary_spec = summary_columns(profile)
    detail_spec = detail_columns(profile)
    summary_headers = [header for header, _ in summary_spec]
    detail_headers = [header for header, _ in detail_spec]

    runs = db.scalars(
        select(ScoringRun)
        .join(Paper)
        .where(Paper.batch_id == batch_id)
        .options(
            selectinload(ScoringRun.paper),
            selectinload(ScoringRun.rubric),
            selectinload(ScoringRun.items),
            selectinload(ScoringRun.items).selectinload(ScoreItem.criterion),
        )
    ).all()
    review_logs_by_run = _review_logs_by_run(db, runs)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "总分表"
    summary.append(summary_headers)

    detail = workbook.create_sheet("评分明细表")
    detail.append(detail_headers)

    for run in runs:
        projection = profile.build_spreadsheet_projection(
            batch=batch,
            run=run,
            review_logs=review_logs_by_run.get(run.id, []),
        )
        row = projection["summary"]
        summary.append(
            [
                (row[key] or "") if key == "finished_at" else row[key]
                for _, key in summary_spec
            ]
        )
        for item in projection["details"]:
            detail.append([item[key] for _, key in detail_spec])

    _format_export_sheet(
        summary,
        **_layout(summary_headers, _NO_WRAP_SUMMARY),
        date_columns={
            index for index, header in enumerate(summary_headers, start=1)
            if header in _DATE_HEADERS
        },
    )
    _format_export_sheet(
        detail,
        **_layout(detail_headers, _NO_WRAP_DETAIL),
        confidence_columns={
            index for index, header in enumerate(detail_headers, start=1)
            if header in _CONFIDENCE_HEADERS
        },
    )

    settings.exports_dir.mkdir(parents=True, exist_ok=True)
    path = settings.exports_dir / ("batch_%s_scores.xlsx" % batch_id)
    # 先写临时文件 → 落库提交 → 原子改名为最终文件：commit 失败不会留下孤儿导出文件。
    tmp_path = path.with_name(path.name + ".tmp")
    workbook.save(tmp_path)
    try:
        for run in runs:
            db.add(
                SpreadsheetWriteLog(
                    scoring_run_id=run.id,
                    target_type="excel",
                    target_id=str(path),
                    status="success",
                    response={"path": str(path)},
                )
            )
        db.commit()
    except Exception:
        db.rollback()
        tmp_path.unlink(missing_ok=True)
        raise
    tmp_path.replace(path)
    return path


def _format_export_sheet(
    sheet,
    *,
    widths,
    wrap_columns,
    score_columns,
    date_columns=frozenset(),
    confidence_columns=frozenset(),
):
    """Apply a readable, print-friendly presentation to an export sheet."""

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddFooter.center.text = "第 &P 页 / 共 &N 页"
    sheet.oddFooter.center.size = 9
    sheet.oddFooter.center.color = "666666"

    sheet.row_dimensions[1].height = 30
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row_number in range(2, sheet.max_row + 1):
        row = sheet[row_number]
        if row_number % 2 == 0:
            for cell in row:
                cell.fill = _ALT_ROW_FILL
        for cell in row:
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(
                horizontal=(
                    "right"
                    if cell.column in score_columns | confidence_columns
                    else "left"
                ),
                vertical="top",
                wrap_text=cell.column in wrap_columns,
            )
            if cell.column in score_columns:
                cell.number_format = "0.00"
            elif cell.column in confidence_columns:
                cell.number_format = "0.000"
            elif cell.column in date_columns and cell.value:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
        sheet.row_dimensions[row_number].height = _row_height(
            row,
            widths,
            wrap_columns,
        )


def _row_height(row, widths, wrap_columns):
    lines = 1
    for cell in row:
        if cell.column not in wrap_columns or cell.value in (None, ""):
            continue
        width = widths[cell.column - 1]
        text = str(cell.value)
        estimated = sum(
            max(1, (len(part) + max(int(width) - 1, 1)) // max(int(width), 1))
            for part in text.splitlines() or [text]
        )
        lines = max(lines, estimated)
    # Keep enough height for long evidence/recommendation cells. A generous
    # cap prevents pathological exports from becoming unbounded while avoiding
    # the visible clipping that occurred with the previous fixed-height rows.
    return min(240, max(22, 16 * lines))


def _review_logs_by_run(db, runs):
    if not runs:
        return {}
    logs = db.scalars(
        select(ReviewLog)
        .where(ReviewLog.scoring_run_id.in_([run.id for run in runs]))
        .order_by(ReviewLog.created_at)
    ).all()
    result = {}
    for log in logs:
        result.setdefault(log.scoring_run_id, []).append(log)
    return result
