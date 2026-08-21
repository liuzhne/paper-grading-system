"""Readable Excel projection for ``grading-core/run-export@2``."""

from __future__ import annotations

import json

from openpyxl import Workbook

from backend.app.core.config import settings
from backend.app.db.models import SpreadsheetWriteLog
from backend.app.services.report.generic_export import build_run_export_v2
from backend.app.services.spreadsheet.excel import _format_export_sheet


def _json_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def _safe_cell(value):
    """Prevent exported user text from being interpreted as a formula."""

    if value is None or isinstance(value, (int, float, bool)):
        return value
    text = _json_text(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _append(sheet, values) -> None:
    sheet.append([_safe_cell(value) for value in values])


def _flatten(value, path=""):
    if isinstance(value, dict):
        if not value:
            yield path or "$", "{}"
        for key in sorted(value):
            nested_path = "%s.%s" % (path, key) if path else str(key)
            yield from _flatten(value[key], nested_path)
        return
    if isinstance(value, list):
        if not value:
            yield path or "$", "[]"
        for index, nested in enumerate(value):
            yield from _flatten(nested, "%s[%s]" % (path, index))
        return
    yield path or "$", value


def _summary_sheet(workbook, export):
    sheet = workbook.active
    sheet.title = "运行摘要"
    sheet.append(["字段", "值"])
    for prefix in ("run", "submission", "identity"):
        for path, value in _flatten(export[prefix], prefix):
            _append(sheet, [path, value])
    _format_export_sheet(
        sheet,
        widths=[38, 78],
        wrap_columns={1, 2},
        score_columns=set(),
    )
    for field_cell, value_cell in sheet.iter_rows(min_row=2, max_col=2):
        field = str(field_cell.value or "")
        if (
            field.endswith("_hash")
            or field.endswith(".hash")
            or field.endswith(".ref")
            or field.endswith("idempotency_key")
        ):
            value_cell.number_format = "@"
            value_cell.quotePrefix = True


def _criteria_sheet(workbook, export):
    sheet = workbook.create_sheet("评分项")
    sheet.append(
        [
            "评分项编码",
            "评分项名称",
            "满分",
            "AI得分",
            "人工终分",
            "自动状态",
            "证据充分",
            "需要复核",
            "理由",
            "聚合结果",
        ]
    )
    for criterion in export["criteria"]:
        _append(
            sheet,
            [
                criterion["criterion_code"],
                criterion["criterion_name"],
                criterion["max_score"],
                criterion["ai_score"],
                criterion["final_score"],
                criterion["automatic_status"],
                criterion["evidence_sufficient"],
                criterion["need_manual_review"],
                criterion["reason"],
                criterion["aggregation"],
            ],
        )
    _format_export_sheet(
        sheet,
        widths=[20, 28, 11, 11, 11, 16, 13, 13, 54, 48],
        wrap_columns={1, 2, 6, 9, 10},
        score_columns={3, 4, 5},
    )


def _rules_sheet(workbook, export):
    sheet = workbook.create_sheet("规则结果")
    sheet.append(
        [
            "评分项编码",
            "规则编码",
            "方向",
            "效果类型",
            "状态",
            "档位",
            "计算效果",
            "证据数量",
            "Observations",
            "规则版本哈希",
        ]
    )
    for criterion in export["criteria"]:
        for rule in criterion["rules"]:
            _append(
                sheet,
                [
                    criterion["criterion_code"],
                    rule["rule_code"],
                    rule["direction"],
                    rule["effect_type"],
                    rule["status"],
                    rule["level_code"],
                    rule["calculated_effect"],
                    len(rule["evidence"]),
                    rule["observations"],
                    rule["version_hash"],
                ],
            )
    _format_export_sheet(
        sheet,
        widths=[20, 28, 12, 16, 16, 18, 28, 12, 54, 68],
        wrap_columns={1, 2, 4, 5, 6, 7, 9, 10},
        score_columns=set(),
    )


def _evidence_sheet(workbook, export):
    sheet = workbook.create_sheet("证据")
    sheet.append(
        ["评分项编码", "规则编码", "类型", "证据ID", "位置", "详情"]
    )
    for criterion in export["criteria"]:
        sources = [("", value) for value in criterion["evidence"]]
        sources.extend(
            (rule["rule_code"], value)
            for rule in criterion["rules"]
            for value in rule["evidence"]
        )
        for rule_code, evidence in sources:
            value = evidence if isinstance(evidence, dict) else {"value": evidence}
            _append(
                sheet,
                [
                    criterion["criterion_code"],
                    rule_code,
                    value.get("type") or value.get("evidence_type"),
                    value.get("evidence_unit_id") or value.get("id"),
                    value.get("location") or value.get("locator"),
                    value,
                ],
            )
    _format_export_sheet(
        sheet,
        widths=[20, 28, 18, 38, 38, 70],
        wrap_columns={1, 2, 3, 4, 5, 6},
        score_columns=set(),
    )


def _reviews_sheet(workbook, export):
    sheet = workbook.create_sheet("复核记录")
    sheet.append(
        [
            "时间",
            "类型",
            "评分项ID",
            "修改前",
            "修改后",
            "复核人",
            "说明",
            "策略哈希",
        ]
    )
    for review in export["review_logs"]:
        _append(
            sheet,
            [
                review["created_at"],
                review["resolution_type"],
                review["score_item_id"],
                review["before_score"],
                review["after_score"],
                review["reviewer_id"],
                review["reason"],
                review["policy_hash"],
            ],
        )
    _format_export_sheet(
        sheet,
        widths=[24, 22, 38, 12, 12, 38, 54, 68],
        wrap_columns={1, 2, 3, 6, 7, 8},
        score_columns={4, 5},
    )


def _extensions_sheet(workbook, export):
    sheet = workbook.create_sheet("Profile扩展")
    sheet.append(["路径", "值"])
    for path, value in _flatten(export["profile_extensions"]):
        _append(sheet, [path, value])
    _format_export_sheet(
        sheet,
        widths=[52, 78],
        wrap_columns={1, 2},
        score_columns=set(),
    )


def export_run_excel_v2(db, run_id: str):
    export = build_run_export_v2(db, run_id)
    workbook = Workbook()
    _summary_sheet(workbook, export)
    _criteria_sheet(workbook, export)
    _rules_sheet(workbook, export)
    _evidence_sheet(workbook, export)
    _reviews_sheet(workbook, export)
    _extensions_sheet(workbook, export)

    settings.exports_dir.mkdir(parents=True, exist_ok=True)
    path = settings.exports_dir / ("run_%s_export_v2.xlsx" % run_id)
    temporary = path.with_name(path.name + ".tmp")
    workbook.save(temporary)
    try:
        db.add(
            SpreadsheetWriteLog(
                scoring_run_id=run_id,
                target_type="excel_v2",
                target_id=str(path),
                status="success",
                response={
                    "path": str(path),
                    "schema": export["schema"],
                    "sheet_names": workbook.sheetnames,
                },
            )
        )
        db.commit()
    except Exception:
        db.rollback()
        temporary.unlink(missing_ok=True)
        raise
    temporary.replace(path)
    return path


__all__ = ["export_run_excel_v2"]
