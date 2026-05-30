from io import BytesIO

from docx import Document

from backend.app.services.rubric_import.docx_comments import parse_comments


def _docx_with_comment():
    doc = Document()
    doc.add_paragraph("第三章 研究方法")  # 作为最近上文标题（章节）
    paragraph = doc.add_paragraph("")
    run = paragraph.add_run("数据来源描述")
    doc.add_comment(runs=run, text="此处需说明数据来源且不少于100字", author="导师")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_parse_comments_extracts_text_anchor_and_section():
    comments = parse_comments(_docx_with_comment())
    assert len(comments) == 1
    comment = comments[0]
    assert "数据来源" in comment["comment_text"]
    assert comment["anchor_text"] == "数据来源描述"
    assert comment["section_title"] == "第三章 研究方法"
    assert comment["author"] == "导师"


def test_parse_comments_empty_when_no_comments():
    doc = Document()
    doc.add_paragraph("没有批注的文档")
    buffer = BytesIO()
    doc.save(buffer)
    assert parse_comments(buffer.getvalue()) == []
