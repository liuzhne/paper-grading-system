from io import BytesIO

import pytest
from openpyxl import Workbook

from backend.app.services.rubric_import.sources.docx_adapter import load_docx
from backend.app.services.rubric_import.sources.units import SourceLedger
from backend.app.services.rubric_import.sources.xlsx_adapter import load_xlsx
from backend.app.tests import rubric_parse_fixtures as fx


def test_xlsx_registers_every_non_empty_cell_with_stable_ids():
    ledger = SourceLedger()
    sheets = load_xlsx(fx.simple_rules_xlsx(), ledger, doc_id="rules", doc_role="rules")
    assert [sheet.title for sheet in sheets] == ["评分规则"]
    unit = ledger.unit("xlsx:评分规则!R2C3")
    assert unit.text == "20"
    assert unit.kind == "cell"
    assert unit.context["value_type"] == "num"
    assert len(ledger.units()) == 18  # 3 行 × 6 列，均非空


def test_xlsx_merged_ranges_register_only_anchor_and_expand_values_in_rows():
    ledger = SourceLedger()
    sheet = load_xlsx(fx.complex_rules_xlsx(), ledger, doc_id="rules", doc_role="rules")[0]
    assert ledger.has("xlsx:评分表!R1C1")
    assert not ledger.has("xlsx:评分表!R1C2")
    assert ledger.unit("xlsx:评分表!R4C1").context["merged"] == "A4:A5"
    # 行视图与原 _rows_with_merged_values 一致：合并区域展开为同值
    assert sheet.rows[4][0] == "选题与综述"
    assert sheet.unit_at(4, 0) == "xlsx:评分表!R4C1"
    assert sheet.unit_at(1, 3) is None  # 空行
    assert sheet.row_units(2)[0] == "xlsx:评分表!R3C1"


def test_xlsx_reads_all_sheets_and_skips_blank_strings():
    workbook = Workbook()
    workbook.active.append(["  ", None, "有效"])
    workbook.create_sheet("第二张").append(["x"])
    buffer = BytesIO()
    workbook.save(buffer)
    ledger = SourceLedger()
    sheets = load_xlsx(buffer.getvalue(), ledger, doc_id="rules", doc_role="rules")
    assert [s.title for s in sheets] == ["Sheet", "第二张"]
    assert [u.unit_id for u in ledger.units()] == ["xlsx:Sheet!R1C3", "xlsx:第二张!R1C1"]


def test_xlsx_rejects_non_workbook_bytes():
    with pytest.raises(ValueError):
        load_xlsx(b"not a workbook", SourceLedger(), doc_id="rules", doc_role="rules")


def test_docx_registers_headings_paragraphs_table_cells_and_comments():
    ledger = SourceLedger()
    view = load_docx(fx.template_docx_with_comments(), ledger, doc_id="tpl", doc_role="template")
    kinds = {u.unit_id: u.kind for u in ledger.units()}
    heading = next(u for u in ledger.units() if u.text == "第三章 研究方法")
    assert kinds[heading.unit_id] == "heading"
    body = next(u for u in ledger.units() if u.text.startswith("正文不少于800字"))
    assert body.kind == "paragraph"
    assert body.context["heading_path"] == ["本科毕业论文模板", "第三章 研究方法"]
    cell = next(u for u in ledger.units() if u.text == "须有题注")
    assert cell.kind == "table_cell"
    assert cell.unit_id == "docx:tbl[0]/r0/c1"
    assert cell.context["heading_path"] == ["本科毕业论文模板", "第三章 研究方法"]
    comment = next(u for u in ledger.units() if u.kind == "comment")
    assert comment.text == "此处需说明数据来源，缺失扣3分"
    assert comment.context["anchor_text"] == "数据来源描述"
    assert comment.context["section_title"] == "第三章 研究方法"
    assert [c["unit_id"] for c in view.comments] == [comment.unit_id]
    assert view.tables[0][0][1] == ("docx:tbl[0]/r0/c1", "须有题注")


def test_docx_merged_table_cells_register_once():
    from docx import Document

    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合并表头"
    table.cell(1, 0).text = "a"
    buffer = BytesIO()
    document.save(buffer)
    ledger = SourceLedger()
    view = load_docx(buffer.getvalue(), ledger, doc_id="d", doc_role="rules")
    assert [u.text for u in ledger.units()] == ["合并表头", "a"]
    assert view.tables[0][0][1][0] == view.tables[0][0][0][0]


def test_docx_rejects_invalid_bytes():
    with pytest.raises(ValueError):
        load_docx(b"garbage", SourceLedger(), doc_id="d", doc_role="rules")


def test_docx_multi_row_table_registers_every_cell_once():
    ledger = SourceLedger()
    view = load_docx(fx.rules_docx(), ledger, doc_id="d", doc_role="rules")
    cells = [u for u in ledger.units() if u.kind == "table_cell"]
    assert len(cells) == 16
    assert view.tables[0][1][1] == ("docx:tbl[0]/r1/c1", "问题分析")
