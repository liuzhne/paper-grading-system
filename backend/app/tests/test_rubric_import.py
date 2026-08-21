from io import BytesIO

from openpyxl import Workbook
from openpyxl import load_workbook

from backend.app.tests.conftest import make_rules_xlsx
from backend.app.tests.conftest import make_template_docx


def test_download_rubric_import_template(client):
    response = client.get("/api/rubrics/import-template.xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    workbook = load_workbook(BytesIO(response.content), data_only=False)
    assert workbook.sheetnames == ["评分规则", "填写说明"]
    sheet = workbook["评分规则"]
    assert [cell.value for cell in sheet[1]] == ["编号", "评分项", "分值", "评分说明", "证据提示", "扣分规则", "顺序"]
    assert sheet["B2"].value == "选题意义"
    assert sheet["C9"].value == "=SUM(C2:C8)"


def test_import_rubric_from_word_template_and_excel_rules(client):
    response = client.post(
        "/api/rubrics/import-files",
        data={"name": "模板导入评分标准", "version": "v1.0", "description": "导入测试"},
        files={
            "template_file": (
                "template.docx",
                make_template_docx().getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
            "rules_file": (
                "rules.xlsx",
                make_rules_xlsx().getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
    )
    assert response.status_code == 200, response.text
    result = response.json()
    rubric = result["rubric"]

    assert rubric["name"] == "模板导入评分标准"
    assert rubric["status"] == "draft"
    assert rubric["total_score"] == 45
    assert len(rubric["criteria"]) == 3
    assert "研究方法" in rubric["criteria"][0]["evidence_hints"]
    assert "实验设计" in rubric["criteria"][0]["evidence_hints"]
    assert "Word 模板解析提示" in rubric["criteria"][0]["description"]
    assert "参考文献" in result["template_summary"]["hints"]
    # 模板期望格式规格已抽取并持久化（格式检查器的基准前置）。
    assert rubric["format_spec"]["source"] == "template"
    assert "body_font_ascii" in rubric["format_spec"]
    assert "format_spec" in result["template_summary"]

    duplicate_response = client.post(
        "/api/rubrics/import-files",
        data={"name": "模板导入评分标准", "version": "v1.0"},
        files={
            "rules_file": (
                "rules.xlsx",
                make_rules_xlsx().getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert duplicate_response.status_code == 400
    assert duplicate_response.json()["detail"] == "rubric name and version already exist"


def test_imports_school_template_with_embedded_scores_and_merged_names(client):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "指导教师评分项"
    sheet.append(["打分项", "评价内容", "具  体  要  求"])
    rows = [
        ("指导教师成绩项1（20分）", "选题与开题", "选题合理，计划清楚。"),
        ("指导教师成绩项2（20分）", "分析与解决问题", None),
        ("指导教师成绩项3（20分）", None, "检索并运用相关技术。"),
        ("指导教师成绩项4（10分）", None, "合理使用工具并持续优化。"),
        ("指导教师成绩项5（10分）", None, "兼顾维护、安全与可持续性。"),
        ("指导教师成绩项6（20分）", "学习与工作态度", "按期完成任务。"),
    ]
    for row in rows:
        sheet.append(row)
    sheet.merge_cells("B3:B6")
    rules = BytesIO()
    workbook.save(rules)

    response = client.post(
        "/api/rubrics/import-files",
        data={"name": "院校指导教师评分标准", "version": "v1.0"},
        files={
            "rules_file": (
                "打分规则.xlsx",
                rules.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 200, response.text
    rubric = response.json()["rubric"]
    assert rubric["total_score"] == 100
    assert [item["code"] for item in rubric["criteria"]] == [
        "T01",
        "T02",
        "T03",
        "T04",
        "T05",
        "T06",
    ]
    assert [item["max_score"] for item in rubric["criteria"]] == [
        20,
        20,
        20,
        10,
        10,
        20,
    ]
    assert [item["name"] for item in rubric["criteria"]][1:5] == [
        "分析与解决问题",
    ] * 4
