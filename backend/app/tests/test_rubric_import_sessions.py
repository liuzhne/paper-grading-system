"""评分标准临时导入会话：解析不落 Rubric，人工确认后才原子创建。"""

from io import BytesIO

from openpyxl import Workbook
from sqlalchemy import func
from sqlalchemy import select
from datetime import timedelta
from urllib.parse import quote

from backend.app.db import models
from backend.app.tests import rubric_parse_fixtures as fx


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _rules_with_decimal_scores(*scores):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "评分规则"
    sheet.append(["编号", "评分项", "分值", "评分说明"])
    for index, score in enumerate(scores, start=1):
        sheet.append([f"C{index:02d}", f"评分项{index}", score, f"说明{index}"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _rules_with_rows(*rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "评分规则"
    sheet.append(["编号", "评分项", "分值", "评分说明"])
    for row in rows:
        sheet.append(list(row))
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _two_rules_for_one_criterion():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["原子规则编号", "评分项编号", "评分项", "分值", "类型", "评分模式", "评分说明"])
    sheet.append(["rule.one", "METHOD", "研究方法", 10, "semantic", "banded", "规则一"])
    sheet.append(["rule.two", "METHOD", "研究方法", 10, "semantic", "banded", "规则二"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _create_session(client, rules, *, name="待确认标准"):
    return client.post(
        "/api/rubrics/import-sessions",
        data={"name": name, "version": "v1"},
        files={"rules_file": ("rules.xlsx", rules, XLSX)},
    )


def _rubric_count(client):
    with client.session_factory() as session:
        return session.scalar(select(func.count()).select_from(models.Rubric))


def test_parsing_creates_session_without_persisting_rubric(client):
    response = _create_session(client, fx.simple_rules_xlsx())

    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "draft"
    assert body["state_version"] == 1
    assert body["rubric_id"] is None
    assert [item["code"] for item in body["criteria"]] == ["C01", "C02"]
    assert _rubric_count(client) == 0


def test_decimal_scores_are_rounded_half_up_and_reported(client):
    response = _create_session(client, _rules_with_decimal_scores(10.4, 10.5))

    assert response.status_code == 201, response.text
    body = response.json()
    assert [item["max_score"] for item in body["criteria"]] == [10, 11]
    assert body["total_score"] == 21
    adjustments = body["score_adjustments"]
    assert [(item["code"], item["original"], item["rounded"]) for item in adjustments] == [
        ("C01", "10.4", 10),
        ("C02", "10.5", 11),
    ]
    assert all("四舍五入" in item["message"] for item in adjustments)
    assert _rubric_count(client) == 0


def test_multiple_atomic_rules_for_one_criterion_produce_one_draft_item(client):
    created = _create_session(client, _two_rules_for_one_criterion()).json()

    assert [item["code"] for item in created["criteria"]] == ["METHOD"]
    confirmed = client.post(
        f"/api/rubrics/import-sessions/{created['id']}/confirm",
        json={"expected_state_version": 1, "idempotency_key": "deduplicated-criterion"},
    )
    assert confirmed.status_code == 200, confirmed.text
    assert [item["code"] for item in confirmed.json()["rubric"]["criteria"]] == ["METHOD"]


def test_confirm_session_persists_once_and_is_idempotent(client):
    created = _create_session(client, fx.simple_rules_xlsx()).json()
    payload = {
        "expected_state_version": created["state_version"],
        "idempotency_key": "confirm-session-once",
    }

    first = client.post(
        f"/api/rubrics/import-sessions/{created['id']}/confirm",
        json=payload,
    )
    second = client.post(
        f"/api/rubrics/import-sessions/{created['id']}/confirm",
        json=payload,
    )

    assert first.status_code == 200, first.text
    assert second.status_code == 200, second.text
    assert first.json()["rubric"]["id"] == second.json()["rubric"]["id"]
    assert first.json()["status"] == second.json()["status"] == "confirmed"
    assert _rubric_count(client) == 1


def test_confirm_rejects_score_sum_mismatch_without_partial_rubric(client):
    created = _create_session(client, fx.simple_rules_xlsx()).json()
    updated = client.patch(
        f"/api/rubrics/import-sessions/{created['id']}",
        json={
            "expected_state_version": created["state_version"],
            "total_score": 100,
            "criteria": created["criteria"],
        },
    )
    assert updated.status_code == 200, updated.text

    confirmed = client.post(
        f"/api/rubrics/import-sessions/{created['id']}/confirm",
        json={
            "expected_state_version": updated.json()["state_version"],
            "idempotency_key": "sum-mismatch",
        },
    )

    assert confirmed.status_code == 422, confirmed.text
    assert "合计" in confirmed.text
    assert _rubric_count(client) == 0


def test_session_requires_at_least_one_supported_file(client):
    empty = client.post(
        "/api/rubrics/import-sessions",
        data={"name": "空会话", "version": "v1"},
    )
    wrong = client.post(
        "/api/rubrics/import-sessions",
        data={"name": "错误文件", "version": "v1"},
        files={"rules_file": ("rules.csv", b"x", "text/csv")},
    )

    assert empty.status_code == 400
    assert "至少上传一份" in empty.text
    assert wrong.status_code == 400
    assert ".xlsx" in wrong.text


def test_stale_edit_is_rejected_and_does_not_overwrite_newer_draft(client):
    created = _create_session(client, fx.simple_rules_xlsx()).json()
    first = client.patch(
        f"/api/rubrics/import-sessions/{created['id']}",
        json={"expected_state_version": 1, "name": "较新的名称"},
    )
    stale = client.patch(
        f"/api/rubrics/import-sessions/{created['id']}",
        json={"expected_state_version": 1, "name": "过期覆盖"},
    )

    assert first.status_code == 200
    assert stale.status_code == 409
    current = client.get(f"/api/rubrics/import-sessions/{created['id']}").json()
    assert current["name"] == "较新的名称"
    assert current["state_version"] == 2


def test_soft_deleted_and_manual_criteria_are_applied_only_on_confirmation(client):
    created = _create_session(client, fx.simple_rules_xlsx()).json()
    criteria = created["criteria"]
    criteria[0]["deleted"] = True
    criteria.append(
        {
            "code": "C03",
            "name": "人工新增",
            "max_score": 20,
            "description": "待第二步补充评分规则",
            "display_order": 2,
            "source_refs": [],
            "parse_status": "manual",
            "deleted": False,
        }
    )
    updated = client.patch(
        f"/api/rubrics/import-sessions/{created['id']}",
        json={
            "expected_state_version": 1,
            "total_score": 35,
            "criteria": criteria,
        },
    )
    assert updated.status_code == 200, updated.text
    assert _rubric_count(client) == 0

    confirmed = client.post(
        f"/api/rubrics/import-sessions/{created['id']}/confirm",
        json={"expected_state_version": 2, "idempotency_key": "manual-and-delete"},
    )

    assert confirmed.status_code == 200, confirmed.text
    assert [item["code"] for item in confirmed.json()["rubric"]["criteria"]] == ["C02", "C03"]


def test_source_preview_returns_safe_extracted_units_not_original_binary(client):
    created = _create_session(client, fx.simple_rules_xlsx()).json()

    response = client.get(
        f"/api/rubrics/import-sessions/{created['id']}/source-preview",
        params={"document": "excel"},
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["document"] == "excel"
    assert body["items"]
    assert all(isinstance(item["text"], str) for item in body["items"])
    assert "binary" not in body


def test_expired_session_cannot_be_confirmed(client):
    created = _create_session(client, fx.simple_rules_xlsx()).json()
    with client.session_factory() as session:
        row = session.get(models.RubricImportSession, created["id"])
        row.expires_at = models.utcnow() - timedelta(seconds=1)
        session.commit()

    response = client.post(
        f"/api/rubrics/import-sessions/{created['id']}/confirm",
        json={"expected_state_version": 1, "idempotency_key": "expired"},
    )

    assert response.status_code == 410
    assert _rubric_count(client) == 0


def test_reupload_previews_diff_and_changes_session_only_after_confirmation(client):
    created = _create_session(client, fx.simple_rules_xlsx()).json()
    replacement = _rules_with_decimal_scores(20, 5)

    preview = client.post(
        f"/api/rubrics/import-sessions/{created['id']}/reupload-preview",
        data={"expected_state_version": 1},
        files={"rules_file": ("replacement.xlsx", replacement, XLSX)},
    )

    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["fingerprint"]
    assert [(item["code"], item["change_type"]) for item in body["criteria_diff"]] == [
        ("C01", "modified"),
        ("C02", "modified"),
    ]
    unchanged = client.get(f"/api/rubrics/import-sessions/{created['id']}").json()
    assert unchanged["state_version"] == 1
    assert [item["name"] for item in unchanged["criteria"]] == ["研究方法", "文献综述"]

    applied = client.post(
        f"/api/rubrics/import-sessions/{created['id']}/reupload-confirm",
        data={"expected_state_version": 1, "fingerprint": body["fingerprint"]},
        files={"rules_file": ("replacement.xlsx", replacement, XLSX)},
    )

    assert applied.status_code == 200, applied.text
    applied_body = applied.json()
    assert applied_body["state_version"] == 2
    assert [item["max_score"] for item in applied_body["criteria"]] == [20, 5]
    assert all(item["parse_status"] == "changed_requires_confirmation" for item in applied_body["criteria"])
    assert _rubric_count(client) == 0


def test_reupload_retains_unchanged_manual_work_and_marks_added_or_removed_items(client):
    created = _create_session(client, fx.simple_rules_xlsx()).json()
    edited = created["criteria"]
    edited[0]["description"] = "用户补充说明"
    edited[1]["deleted"] = True
    edited.append({
        "code": "M01", "name": "人工项", "max_score": 3, "description": "人工创建",
        "display_order": 2, "source_refs": [], "parse_status": "manual", "deleted": False,
    })
    updated = client.patch(
        f"/api/rubrics/import-sessions/{created['id']}",
        json={"expected_state_version": 1, "criteria": edited, "total_score": 23},
    ).json()
    replacement = _rules_with_rows(
        ("C01", "研究方法", 20, "方法合理，数据来源清楚。"),
        ("C02", "文献综述", 15, "综述覆盖充分。"),
        ("C03", "创新性", 5, "有创新。"),
    )
    preview = client.post(
        f"/api/rubrics/import-sessions/{created['id']}/reupload-preview",
        data={"expected_state_version": updated["state_version"]},
        files={"rules_file": ("replacement.xlsx", replacement, XLSX)},
    ).json()
    applied = client.post(
        f"/api/rubrics/import-sessions/{created['id']}/reupload-confirm",
        data={"expected_state_version": updated["state_version"], "fingerprint": preview["fingerprint"]},
        files={"rules_file": ("replacement.xlsx", replacement, XLSX)},
    )

    assert applied.status_code == 200, applied.text
    by_code = {item["code"]: item for item in applied.json()["criteria"]}
    assert by_code["C01"]["description"] == "用户补充说明"
    assert by_code["C02"]["deleted"] is True
    assert by_code["M01"]["parse_status"] == "manual"
    assert by_code["C03"]["parse_status"] == "added_requires_confirmation"


def test_reupload_soft_deletes_removed_source_item(client):
    created = _create_session(client, fx.simple_rules_xlsx()).json()
    replacement = _rules_with_rows(("C01", "研究方法", 20, "方法合理，数据来源清楚。"))
    preview = client.post(
        f"/api/rubrics/import-sessions/{created['id']}/reupload-preview",
        data={"expected_state_version": 1},
        files={"rules_file": ("replacement.xlsx", replacement, XLSX)},
    ).json()
    applied = client.post(
        f"/api/rubrics/import-sessions/{created['id']}/reupload-confirm",
        data={"expected_state_version": 1, "fingerprint": preview["fingerprint"]},
        files={"rules_file": ("replacement.xlsx", replacement, XLSX)},
    )

    assert applied.status_code == 200, applied.text
    removed = next(item for item in applied.json()["criteria"] if item["code"] == "C02")
    assert removed["deleted"] is True
    assert removed["parse_status"] == "removed"


def test_source_conflicts_require_explicit_resolution_before_confirmation(client):
    response = client.post(
        "/api/rubrics/import-sessions",
        data={"name": "双文件冲突", "version": "v1"},
        files={
            "rules_file": ("rules.xlsx", fx.simple_rules_xlsx(), XLSX),
            "template_file": (
                "template.docx",
                fx.template_docx_with_rule_table(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        },
    )
    assert response.status_code == 201, response.text
    created = response.json()
    assert len(created["conflicts"]) == 2

    blocked = client.post(
        f"/api/rubrics/import-sessions/{created['id']}/confirm",
        json={"expected_state_version": 1, "idempotency_key": "unresolved"},
    )
    assert blocked.status_code == 422
    assert "来源冲突" in blocked.text

    version = created["state_version"]
    for conflict in created["conflicts"]:
        resolved = client.post(
            f"/api/rubrics/import-sessions/{created['id']}/conflicts/"
            f"{quote(conflict['anchor_unit_id'], safe='')}/resolve",
            json={
                "expected_state_version": version,
                "decision": "use_excel",
                "reason": "确认以 Excel 为结构主干",
            },
        )
        assert resolved.status_code == 200, resolved.text
        version = resolved.json()["state_version"]

    confirmed = client.post(
        f"/api/rubrics/import-sessions/{created['id']}/confirm",
        json={"expected_state_version": version, "idempotency_key": "resolved"},
    )
    assert confirmed.status_code == 200, confirmed.text
    assert _rubric_count(client) == 1


def test_cancelled_session_is_persisted_and_cannot_be_confirmed(client):
    created = _create_session(client, fx.simple_rules_xlsx()).json()

    cancelled = client.post(
        f"/api/rubrics/import-sessions/{created['id']}/cancel",
        json={"expected_state_version": created["state_version"]},
    )
    assert cancelled.status_code == 200, cancelled.text
    assert cancelled.json()["status"] == "cancelled"
    assert cancelled.json()["state_version"] == 2

    confirmed = client.post(
        f"/api/rubrics/import-sessions/{created['id']}/confirm",
        json={"expected_state_version": 2, "idempotency_key": "cancelled"},
    )
    assert confirmed.status_code == 409
    assert _rubric_count(client) == 0


def test_confirmed_rubric_reupload_creates_successor_compilation(client):
    created = _create_session(client, fx.simple_rules_xlsx(), name="可重新上传").json()
    confirmed = client.post(
        f"/api/rubrics/import-sessions/{created['id']}/confirm",
        json={"expected_state_version": 1, "idempotency_key": "formal-reupload-base"},
    ).json()
    rubric_id = confirmed["rubric"]["id"]
    c01_id = next(item["id"] for item in confirmed["rubric"]["criteria"] if item["code"] == "C01")
    c02_id = next(item["id"] for item in confirmed["rubric"]["criteria"] if item["code"] == "C02")
    review = client.get(f"/api/rubrics/{rubric_id}/review-workspace").json()
    for criterion_id in (c01_id, c02_id):
        rule = next(item for item in review["rules"] if item["criterion_id"] == criterion_id)
        accepted = client.post(
            f"/api/rubrics/{rubric_id}/rules/{rule['rule_code']}/confirm",
            json={
                "compilation_id": review["compilation_id"],
                "rule_id": rule["id"],
                "content_token": rule["content_token"],
                "reason": "确认用于重新上传继承测试",
            },
        )
        assert accepted.status_code == 200, accepted.text
    with client.session_factory() as session:
        before = session.scalar(
            select(models.RubricCompilation).where(
                models.RubricCompilation.rubric_id == rubric_id,
                models.RubricCompilation.status != "superseded",
            )
        )
        before_id = before.id

    replacement = _rules_with_rows(
        ("C01", "研究方法", 20, "方法合理，数据来源清楚。"),
        ("C02", "文献研究", 15, "综述覆盖充分。"),
    )
    preview = client.post(
        f"/api/rubrics/{rubric_id}/reupload-preview",
        files={"rules_file": ("replacement.xlsx", replacement, XLSX)},
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["modified_count"] == 1
    assert preview.json()["retained_rules_count"] == 1
    assert preview.json()["invalidated_rules_count"] == 1
    with client.session_factory() as session:
        assert session.get(models.RubricCompilation, before_id).status != "superseded"

    applied = client.post(
        f"/api/rubrics/{rubric_id}/reupload-confirm",
        data={"fingerprint": preview.json()["fingerprint"]},
        files={"rules_file": ("replacement.xlsx", replacement, XLSX)},
    )

    assert applied.status_code == 200, applied.text
    assert applied.json()["id"] == rubric_id
    assert [item["max_score"] for item in applied.json()["criteria"]] == [20.0, 15.0]
    after_review = client.get(f"/api/rubrics/{rubric_id}/review-workspace").json()
    assert next(
        rule for rule in after_review["rules"] if rule["criterion_id"] == c01_id
    )["status"] == "approved"
    assert all(
        rule["status"] == "draft"
        for rule in after_review["rules"]
        if rule["criterion_id"] == c02_id
    )
    with client.session_factory() as session:
        assert session.get(models.RubricCompilation, before_id).status == "superseded"
        active = session.scalars(
            select(models.RubricCompilation).where(
                models.RubricCompilation.rubric_id == rubric_id,
                models.RubricCompilation.status != "superseded",
            )
        ).all()
        assert len(active) == 1
        assert active[0].id != before_id


def test_published_rubric_reupload_is_rejected_without_mutation(client):
    created = _create_session(client, fx.simple_rules_xlsx(), name="已发布不可覆盖").json()
    confirmed = client.post(
        f"/api/rubrics/import-sessions/{created['id']}/confirm",
        json={"expected_state_version": 1, "idempotency_key": "published-base"},
    ).json()
    rubric_id = confirmed["rubric"]["id"]
    with client.session_factory() as session:
        session.execute(
            models.Rubric.__table__.update()
            .where(models.Rubric.id == rubric_id)
            .values(status="published")
        )
        session.commit()

    response = client.post(
        f"/api/rubrics/{rubric_id}/reupload-preview",
        files={"rules_file": ("replacement.xlsx", fx.simple_rules_xlsx(), XLSX)},
    )

    assert response.status_code == 409
    assert "不能被重新上传覆盖" in response.text
