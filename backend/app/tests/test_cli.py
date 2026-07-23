"""CLI 端（pgs）冒烟测试：CliRunner + 临时 sqlite + 强制 Mock，自包含不联网。

注意：conftest 的 mock 隔离是普通 fixture（仅 client 测试生效），CLI 测试不走它，
故本文件 autouse 强制 settings.LLM_PROVIDER=mock 并在结束后还原被 configure 改写的全局设置。
"""

import json
from contextlib import contextmanager

import pytest
from openpyxl import Workbook
from sqlalchemy import select
from typer.testing import CliRunner

from backend.app.cli import db as cli_db
from backend.app.cli import main as cli_main
from backend.app.cli.main import app
from backend.app.core.config import settings
from backend.app.db import models
from backend.app.tests.conftest import make_rules_xlsx
from backend.app.tests.conftest import make_sample_docx
from backend.app.tests.m4_import_contract_fixtures import real_rules_xlsx_bytes
from backend.app.tests.m4_import_contract_fixtures import real_template_docx_bytes

runner = CliRunner()

SEED_RUBRIC = "本科毕业论文通用评分标准"


@pytest.fixture(autouse=True)
def _cli_isolation():
    saved = (settings.DATABASE_URL, settings.STORAGE_ROOT, settings.LLM_PROVIDER)
    settings.LLM_PROVIDER = "mock"  # 任何 get_llm_scorer() → mock，绝不联网
    yield
    settings.DATABASE_URL, settings.STORAGE_ROOT, settings.LLM_PROVIDER = saved


@pytest.fixture
def local(tmp_path):
    """每个测试独立的本地 sqlite + storage。"""
    return ["--db", str(tmp_path / "cli.db"), "--storage", str(tmp_path / "storage")]


def test_init_seed_then_rubrics_lists_default(local):
    assert runner.invoke(app, ["init", "--seed", *local]).exit_code == 0
    listing = runner.invoke(app, ["rubrics", "--json", *local])
    assert listing.exit_code == 0
    assert SEED_RUBRIC in listing.output


def test_check_mock_ok():
    result = runner.invoke(app, ["check", "--json"])
    assert result.exit_code == 0
    assert '"stage": "mock"' in result.output


def test_score_mock_writes_report(tmp_path, local):
    runner.invoke(app, ["init", "--seed", *local])
    docx = tmp_path / "thesis.docx"
    docx.write_bytes(make_sample_docx().getvalue())
    report_dir = tmp_path / "reports"
    result = runner.invoke(
        app,
        ["score", str(docx), "--rubric", SEED_RUBRIC, "--mock", "--report-dir", str(report_dir), *local],
    )
    assert result.exit_code == 0, result.output
    assert "thesis.docx" in result.output
    assert (report_dir / "thesis.html").exists()


def test_score_bad_file_nonzero_exit(tmp_path, local):
    runner.invoke(app, ["init", "--seed", *local])
    bad = tmp_path / "broken.docx"
    bad.write_text("this is not a docx")
    result = runner.invoke(app, ["score", str(bad), "--rubric", SEED_RUBRIC, "--mock", "--json", *local])
    assert result.exit_code == 1  # 解析失败 → 非零退出（供脚本/CI 判定）
    assert "broken.docx" in result.output


def test_score_unknown_rubric_errors(tmp_path, local):
    runner.invoke(app, ["init", *local])
    docx = tmp_path / "t.docx"
    docx.write_bytes(make_sample_docx().getvalue())
    result = runner.invoke(app, ["score", str(docx), "--rubric", "不存在的标准", "--mock", *local])
    assert result.exit_code != 0


def test_import_rubric_from_excel(tmp_path, local):
    rules = tmp_path / "rules.xlsx"
    rules.write_bytes(make_rules_xlsx().getvalue())
    result = runner.invoke(app, ["import", str(rules), "--name", "CLI导入标准", "--version", "v1.0", *local])
    assert result.exit_code == 0, result.output
    listing = runner.invoke(app, ["rubrics", "--json", *local])
    assert "CLI导入标准" in listing.output


def test_import_uses_m4_two_phase_pipeline_and_keeps_review_state_empty(
    tmp_path, local, monkeypatch
):
    from backend.app.services.rubric_import import pipeline as rubric_pipeline
    from backend.app.services.rubrics import lifecycle

    rules = tmp_path / "m4-rules.xlsx"
    template = tmp_path / "m4-template.docx"
    rules.write_bytes(real_rules_xlsx_bytes())
    template.write_bytes(real_template_docx_bytes())

    open_sessions = 0
    prepare_calls = []
    original_session = cli_db.cli_session
    original_prepare = rubric_pipeline.prepare_file_import

    @contextmanager
    def tracked_session(*args, **kwargs):
        nonlocal open_sessions
        open_sessions += 1
        try:
            with original_session(*args, **kwargs) as session:
                yield session
        finally:
            open_sessions -= 1

    def tracked_prepare(*args, **kwargs):
        # The expensive parser/compiler phase must not hold a DB transaction.
        assert open_sessions == 0
        prepare_calls.append(kwargs["command"])
        return original_prepare(*args, **kwargs)

    def forbidden_signoff(*_args, **_kwargs):
        raise AssertionError("CLI import must not forge review or publication")

    monkeypatch.setattr(cli_main.clidb, "cli_session", tracked_session)
    monkeypatch.setattr(rubric_pipeline, "prepare_file_import", tracked_prepare)
    monkeypatch.setattr(lifecycle, "approve_atomic_rule", forbidden_signoff)
    monkeypatch.setattr(lifecycle, "publish_rubric", forbidden_signoff)

    result = runner.invoke(
        app,
        [
            "import",
            str(rules),
            "--template",
            str(template),
            "--name",
            "M4 CLI provenance",
            "--version",
            "v1.0",
            *local,
        ],
    )
    assert result.exit_code == 0, result.output
    assert len(prepare_calls) == 1
    assert prepare_calls[0]["schema_version"] == rubric_pipeline.IMPORT_SCHEMA_VERSION

    with original_session() as session:
        rubric = session.scalar(
            select(models.Rubric).where(models.Rubric.name == "M4 CLI provenance")
        )
        compilation = session.scalar(
            select(models.RubricCompilation).where(
                models.RubricCompilation.rubric_id == rubric.id
            )
        )
        version = session.scalar(
            select(models.RubricVersion).where(
                models.RubricVersion.rubric_id == rubric.id
            )
        )
        atomic_rules = session.scalars(
            select(models.AtomicRule).where(
                models.AtomicRule.rubric_version_id == version.id
            )
        ).all()
        artifacts = session.scalars(
            select(models.SourceArtifact).where(
                models.SourceArtifact.compilation_id == compilation.id
            )
        ).all()

        assert rubric.status == "draft"
        assert compilation.human_changes == []
        assert compilation.reviewed_by is None
        assert compilation.reviewed_at is None
        assert compilation.published_at is None
        assert atomic_rules
        assert all(rule.status == "draft" for rule in atomic_rules)
        assert all(rule.reviewed_by is None for rule in atomic_rules)
        assert {artifact.artifact_type for artifact in artifacts} == {"excel", "word"}


def test_runs_and_show_after_score(tmp_path, local):
    runner.invoke(app, ["init", "--seed", *local])
    docx = tmp_path / "t.docx"
    docx.write_bytes(make_sample_docx().getvalue())
    scored = runner.invoke(app, ["score", str(docx), "--rubric", SEED_RUBRIC, "--mock", "--json", *local])
    assert scored.exit_code == 0, scored.output
    run_id = json.loads(scored.output)["results"][0]["run_id"]

    listed = runner.invoke(app, ["runs", "--json", *local])
    assert listed.exit_code == 0
    assert run_id in listed.output

    shown = runner.invoke(app, ["show", run_id, "--json", *local])
    assert shown.exit_code == 0, shown.output
    detail = json.loads(shown.output)
    assert detail["run"]["run_id"] == run_id
    assert len(detail["items"]) == 7  # 种子 7 个评分项


def test_batches_lists_after_score(tmp_path, local):
    runner.invoke(app, ["init", "--seed", *local])
    docx = tmp_path / "t.docx"
    docx.write_bytes(make_sample_docx().getvalue())
    runner.invoke(app, ["score", str(docx), "--rubric", SEED_RUBRIC, "--mock", *local])
    result = runner.invoke(app, ["batches", "--json", *local])
    assert result.exit_code == 0
    assert json.loads(result.output)  # 非空


def test_publish_rejects_unreviewed_import(tmp_path, local):
    rules = tmp_path / "rules.xlsx"
    rules.write_bytes(make_rules_xlsx().getvalue())
    assert runner.invoke(app, ["import", str(rules), "--name", "待发布标准", *local]).exit_code == 0
    rid = next(r["id"] for r in json.loads(runner.invoke(app, ["rubrics", "--json", *local]).output) if r["name"] == "待发布标准")
    assert runner.invoke(app, ["publish", rid, *local]).exit_code == 2
    after = json.loads(runner.invoke(app, ["rubrics", "--json", *local]).output)
    assert next(r["status"] for r in after if r["id"] == rid) == "draft"


def test_cli_graph_discovery_and_explicit_m4_publish_chain(tmp_path, local):
    rules = tmp_path / "strict-rules.xlsx"
    template = tmp_path / "strict-template.docx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "原子规则编号",
            "评分项编号",
            "评分项",
            "分值",
            "类型",
            "评分模式",
            "评分说明",
            "适用范围",
            "分档",
            "evidence_policy",
            "effect_type",
            "strictness",
        ]
    )
    sheet.append(
        [
            "thesis.risk_fit.v1",
            "RISK_FIT",
            "风险方案质量",
            10,
            "semantic",
            "banded",
            "风险控制方案应与论文需求一致。",
            "风险控制",
            "HIGH:10;LOW:5",
            json.dumps(
                {
                    "mode": "source_quote",
                    "requirement": "required",
                    "minimum_coverage": "1",
                }
            ),
            "score",
            "required",
        ]
    )
    workbook.save(rules)
    template.write_bytes(real_template_docx_bytes())
    imported = runner.invoke(
        app,
        [
            "import",
            str(rules),
            "--template",
            str(template),
            "--name",
            "CLI strict lifecycle",
            *local,
        ],
    )
    assert imported.exit_code == 0, imported.output
    rubrics = json.loads(runner.invoke(app, ["rubrics", "--json", *local]).output)
    rubric_id = next(item["id"] for item in rubrics if item["name"] == "CLI strict lifecycle")

    graph_result = runner.invoke(app, ["rubric-graph", rubric_id, "--json", *local])
    assert graph_result.exit_code == 0, graph_result.output
    graph = json.loads(graph_result.output)
    active = graph["active_compilation"]
    assert active["id"]
    assert active["rules"]
    assert "raw_model_output" not in graph_result.output
    assert "rule_text" not in graph_result.output

    for rule in active["rules"]:
        submitted = runner.invoke(
            app,
            [
                "rule-submit",
                rubric_id,
                rule["rule_code"],
                "--reason",
                "CLI 提交",
                *local,
            ],
        )
        assert submitted.exit_code == 0, submitted.output
        approved = runner.invoke(
            app,
            [
                "rule-approve",
                rubric_id,
                rule["rule_code"],
                "--reason",
                "CLI 批准",
                *local,
            ],
        )
        assert approved.exit_code == 0, approved.output
    for link in active["template_links"]:
        reviewed = runner.invoke(
            app,
            [
                "template-link-review",
                rubric_id,
                link["id"],
                "--decision",
                "confirmed",
                "--reason",
                "CLI 确认映射",
                *local,
            ],
        )
        assert reviewed.exit_code == 0, reviewed.output

    submitted_rubric = runner.invoke(
        app, ["rubric-submit-review", rubric_id, *local]
    )
    assert submitted_rubric.exit_code == 0, submitted_rubric.output
    published = runner.invoke(
        app,
        [
            "publish",
            rubric_id,
            "--compilation-id",
            active["id"],
            *local,
        ],
    )
    assert published.exit_code == 0, published.output
    after = json.loads(runner.invoke(app, ["rubrics", "--json", *local]).output)
    assert next(item["status"] for item in after if item["id"] == rubric_id) == "published"


def test_cli_rubric_return_draft_and_explicit_legacy_upgrade(tmp_path, local):
    rules = tmp_path / "return-rules.xlsx"
    rules.write_bytes(make_rules_xlsx().getvalue())
    imported = runner.invoke(
        app,
        ["import", str(rules), "--name", "CLI return lifecycle", *local],
    )
    assert imported.exit_code == 0, imported.output
    rubrics = json.loads(runner.invoke(app, ["rubrics", "--json", *local]).output)
    rubric_id = next(item["id"] for item in rubrics if item["name"] == "CLI return lifecycle")
    assert runner.invoke(app, ["rubric-submit-review", rubric_id, *local]).exit_code == 0
    returned = runner.invoke(app, ["rubric-return-draft", rubric_id, *local])
    assert returned.exit_code == 0, returned.output

    with cli_db.cli_session() as session:
        from backend.app.services.dev_user import ensure_dev_user

        actor = ensure_dev_user(session)
        legacy = models.Rubric(
            name="CLI legacy upgrade",
            version="legacy-v1",
            total_score=10,
            status="draft",
            created_by=actor.id,
            owner_id=actor.id,
        )
        legacy.criteria.append(
            models.RubricCriterion(
                code="LEGACY",
                name="Legacy direct",
                max_score=10,
                criterion_type="llm_judgment",
                scoring_mode="llm_direct",
                applies_to="global",
                display_order=0,
            )
        )
        session.add(legacy)
        session.commit()
        legacy_id = legacy.id

    upgraded = runner.invoke(
        app,
        [
            "rubric-upgrade",
            legacy_id,
            "--reason",
            "CLI 显式升级",
            *local,
        ],
    )
    assert upgraded.exit_code == 0, upgraded.output
    graph = runner.invoke(app, ["rubric-graph", legacy_id, "--json", *local])
    assert graph.exit_code == 0, graph.output
    data = json.loads(graph.output)
    assert data["active_compilation"]["id"]
    assert data["active_compilation"]["blockers"]


def test_database_score_refuses_implicit_unreviewed_file_import(tmp_path, local):
    rules = tmp_path / "rules.xlsx"
    rules.write_bytes(make_rules_xlsx().getvalue())
    docx = tmp_path / "thesis.docx"
    docx.write_bytes(make_sample_docx().getvalue())

    result = runner.invoke(
        app,
        ["score", str(docx), "--rubric-file", str(rules), "--mock", *local],
    )

    assert result.exit_code != 0
    assert "不再隐式导入未审核规则" in result.output

    imported = runner.invoke(
        app,
        ["import", str(rules), "--name", "未审核 CLI 标准", *local],
    )
    assert imported.exit_code == 0, imported.output
    selected_draft = runner.invoke(
        app,
        ["score", str(docx), "--rubric", "未审核 CLI 标准", "--mock", *local],
    )
    assert selected_draft.exit_code != 0
    assert "尚未发布" in selected_draft.output


def test_score_multiple_with_workers(tmp_path, local):
    runner.invoke(app, ["init", "--seed", *local])
    docs = []
    for i in range(3):
        d = tmp_path / ("p%d.docx" % i)
        d.write_bytes(make_sample_docx().getvalue())
        docs.append(str(d))
    result = runner.invoke(
        app, ["score", *docs, "--rubric", SEED_RUBRIC, "--mock", "--workers", "3", "--json", *local]
    )
    assert result.exit_code == 0, result.output
    data = json.loads(result.output)
    assert len(data["results"]) == 3
    assert all(r["status"] == "ok" for r in data["results"])


def test_score_no_db_is_stateless(tmp_path):
    rules = tmp_path / "rules.xlsx"
    rules.write_bytes(make_rules_xlsx().getvalue())
    docx = tmp_path / "t.docx"
    docx.write_bytes(make_sample_docx().getvalue())
    ghost_db = tmp_path / "should_not_exist.db"  # --no-db 即便给了 --db 也不应建库
    result = runner.invoke(
        app,
        ["score", str(docx), "--no-db", "--rubric-file", str(rules), "--mock", "--json", "--db", str(ghost_db)],
    )
    assert result.exit_code == 0, result.output
    data = json.loads(result.output)
    assert data["stateless"] is True
    assert len(data["results"]) == 1
    assert data["results"][0]["status"] == "ok"
    assert data["results"][0]["items"]  # 含逐项明细
    assert not ghost_db.exists()  # 关键：无状态不建任何 sqlite


def test_score_no_db_requires_rubric_file(tmp_path):
    docx = tmp_path / "t.docx"
    docx.write_bytes(make_sample_docx().getvalue())
    result = runner.invoke(app, ["score", str(docx), "--no-db", "--mock"])
    assert result.exit_code != 0  # 缺 --rubric-file → 报错退出


def test_review_overrides_item_and_submits(tmp_path, local):
    runner.invoke(app, ["init", "--seed", *local])
    docx = tmp_path / "t.docx"
    docx.write_bytes(make_sample_docx().getvalue())
    run_id = json.loads(
        runner.invoke(app, ["score", str(docx), "--rubric", SEED_RUBRIC, "--mock", "--json", *local]).output
    )["results"][0]["run_id"]

    reviewed = runner.invoke(app, ["review", run_id, "--set", "C01=9", "--note", "人工调整", "--submit", *local])
    assert reviewed.exit_code == 0, reviewed.output

    shown = json.loads(runner.invoke(app, ["show", run_id, "--json", *local]).output)
    assert shown["run"]["status"] == "reviewed"
    assert next(i["score"] for i in shown["items"] if i["code"] == "C01") == 9.0


def test_review_rejects_unknown_code(tmp_path, local):
    runner.invoke(app, ["init", "--seed", *local])
    docx = tmp_path / "t.docx"
    docx.write_bytes(make_sample_docx().getvalue())
    run_id = json.loads(
        runner.invoke(app, ["score", str(docx), "--rubric", SEED_RUBRIC, "--mock", "--json", *local]).output
    )["results"][0]["run_id"]
    assert runner.invoke(app, ["review", run_id, "--set", "ZZ=5", *local]).exit_code != 0


def test_scores_template_columns_match_rubric(tmp_path, local):
    runner.invoke(app, ["init", "--seed", *local])
    rid = next(
        r["id"]
        for r in json.loads(runner.invoke(app, ["rubrics", "--json", *local]).output)
        if r["name"] == SEED_RUBRIC
    )
    out = tmp_path / "tmpl.xlsx"
    result = runner.invoke(app, ["scores-template", rid, "-o", str(out), *local])
    assert result.exit_code == 0, result.output

    from openpyxl import load_workbook

    headers = [cell.value for cell in load_workbook(out)["教师评分"][1]]
    assert headers[:2] == ["文件名", "总分"]
    assert "C01" in headers and "C07" in headers


def test_check_mock_overrides_real_provider(monkeypatch):
    monkeypatch.setattr(settings, "LLM_PROVIDER", "openai_compatible")  # 即使配了真实 provider
    result = runner.invoke(app, ["check", "--mock", "--json"])
    assert result.exit_code == 0
    assert '"stage": "mock"' in result.output


def test_apply_llm_overrides_local(monkeypatch):
    from backend.app.cli.main import _apply_llm_overrides

    monkeypatch.setattr(settings, "LLM_PROVIDER", "mock")
    monkeypatch.setattr(settings, "LOCAL_LLM_BASE_URL", "placeholder")
    monkeypatch.setattr(settings, "LOCAL_LLM_MODEL", "placeholder")
    _apply_llm_overrides("local", "qwen3-30b", "http://localhost:1234/v1")
    assert settings.LLM_PROVIDER == "local"
    assert settings.LOCAL_LLM_BASE_URL == "http://localhost:1234/v1"
    assert settings.LOCAL_LLM_MODEL == "qwen3-30b"


def test_doctor_reports_offline_ready(monkeypatch):
    # mock LLM + mock 表格写入 → 纯本地、零外呼。
    monkeypatch.setattr(settings, "LLM_PROVIDER", "mock")
    monkeypatch.setattr(settings, "SHEET_WRITER_PROVIDER", "mock")
    result = runner.invoke(app, ["doctor", "--json"])
    assert result.exit_code == 0
    assert '"offline_ready": true' in result.output


def test_doctor_flags_cloud_touchpoint(monkeypatch):
    monkeypatch.setattr(settings, "LLM_PROVIDER", "mock")
    monkeypatch.setattr(settings, "SHEET_WRITER_PROVIDER", "google_sheets")  # 外呼环节
    result = runner.invoke(app, ["doctor", "--json"])
    assert '"offline_ready": false' in result.output
