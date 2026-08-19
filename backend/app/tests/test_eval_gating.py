import json
from contextlib import nullcontext
from copy import deepcopy
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from sqlalchemy import select

from backend.app.db.models import RubricCompilation
from backend.app.eval.gating import APPROVAL_SCHEMA
from backend.app.eval.gating import GateValidationError
from backend.app.eval.gating import approved_regression_issues
from backend.app.eval.gating import build_gate_candidate
from backend.app.eval.gating import build_holdout_manifest_from_rows
from backend.app.eval.gating import build_regression_record
from backend.app.eval.gating import ensure_external_artifact_dir
from backend.app.eval.gating import evaluation_sha256
from backend.app.eval.gating import finalize_gate
from backend.app.eval.gating import summarize_run_identities
from backend.app.eval.release_preflight import PGS8_CRITERION_CODES
from backend.app.eval.release_preflight import validate_release_rubric_and_scores
from backend.app.tests.conftest import publish_rubric_via_api


def _sha(character):
    return character * 64


def _report(provider="openai"):
    run_identity = {
        "schema_version": "paper-grading/thesis-eval-run-identity@1",
        "rubric_version_id": "rubric-version-1",
        "rubric_version_hash": _sha("1"),
        "rubric_hash_scheme": "rubric-content-v1",
        "rubric_snapshot_hash": _sha("2"),
        "policy_hash": _sha("3"),
        "policy_snapshot_sha256": _sha("a"),
        "grade_scale_sha256": _sha("b"),
        "execution_plan_hash": _sha("4"),
        "plan_schema_version": "rule-execution-plan@2",
        "checker_manifest_sha256": _sha("5"),
        "business_profile_key": "thesis",
        "business_profile_version": "thesis-legacy-profile@1",
        "workflow_profile": "thesis",
        "prompt_version": "2026-08-02-9",
        "runtime_identity_sha256": _sha("c"),
        "model_provider": provider,
        "model_name": "fixed-model",
        "model_version": "2026-07-01",
        "engine_version": "scoring-core@1",
        "source_artifact_hash": _sha("6"),
        "normalized_content_hash": _sha("7"),
        "document_snapshot_hash": _sha("8"),
    }
    return {
        "n": 1,
        "dataset_size": 1,
        "completed_runs": 1,
        "errors": [],
        "unmatched_predictions": [],
        "qwk": 1.0,
        "mae": 1.0,
        "rmse": 1.0,
        "exact_grade_agreement": 1.0,
        "adjacent_grade_agreement": 1.0,
        "review_rate": 0.0,
        "blocked_rate": 0.0,
        "invalid_evidence_rate": 0.0,
        "per_criterion": {"C01": {"n": 1, "mae": 1.0, "bias": 1.0}},
        "grade_confusion": [[1]],
        "grade_labels": ["优秀"],
        "per_sample": [
            {
                "sample_id": _sha("9"),
                "human_total": 90.0,
                "system_total": 91.0,
                "human_items": {"C01": 9.0},
                "system_items": {"C01": 10.0},
                "run_identity": run_identity,
            }
        ],
    }


def _candidate(provider="openai"):
    report = _report(provider)
    run_identity, issues = summarize_run_identities(report)
    return build_gate_candidate(
        evaluation_id="pgs-8-test",
        report=report,
        dataset_identity={
            "dataset_id": "teacher-holdout",
            "dataset_version": "2026-07-30-v1",
            "sample_count": 1,
            "dataset_manifest_sha256": _sha("a"),
            "sample_ids_sha256": evaluation_sha256([_sha("9")]),
            "paper_manifest_sha256": _sha("b"),
            "truth_sha256": _sha("c"),
        },
        run_identity=run_identity,
        run_identity_issues=issues,
        anchors_identity={
            "schema": "paper-grading/anchor-manifest@1",
            "count": 0,
            "manifest_sha256": _sha("d"),
            "holdout_exclusion_proven": False,
        },
        model_identity={
            "provider": provider,
            "model_name": "fixed-model",
            "model_version": "2026-07-01",
            "identity_method": "provider_immutable_revision",
            "immutable_revision": "fixed-model-2026-07-01",
            "artifact_identity_sha256": _sha("e"),
        },
        prompt_identity={
            "prompt_version": "2026-07-20-7",
            "prompt_source_manifest_sha256": _sha("f"),
        },
        repository_identity={
            "revision": "a" * 40,
            "branch": "codex/pgs-8",
            "dirty": False,
            "lockfile": "uv.lock",
            "lockfile_sha256": _sha("0"),
        },
        private_artifacts={
            "location": "external_private_not_in_repository",
            "private_manifest_sha256": _sha("1"),
            "private_report_sha256": _sha("2"),
        },
        artifact_directory_external=True,
        generated_at="2026-07-30T00:00:00+00:00",
    )


def _approval(candidate):
    metrics = deepcopy(candidate["evaluation"]["metrics"])
    return {
        "schema": APPROVAL_SCHEMA,
        "candidate_sha256": candidate["candidate_sha256"],
        "approved_by": "authorized-maintainer",
        "approved_at": "2026-07-30T12:00:00+08:00",
        "privacy_review": {
            "dataset_deidentified": True,
            "repository_scan_passed": True,
            "anchors_exclude_holdout": True,
            "teacher_truth_external_only": True,
        },
        "accepted_baseline": {
            "metrics": metrics,
            "acceptance_thresholds": {
                "minimum_qwk": 0.5,
                "maximum_mae": 2.0,
                "maximum_rmse": 2.0,
                "minimum_exact_grade_agreement": 0.8,
                "minimum_adjacent_grade_agreement": 0.9,
                "maximum_review_rate": 0.2,
                "maximum_blocked_rate": 0.0,
                "maximum_invalid_evidence_rate": 0.0,
            },
            "regression_tolerances": {
                "qwk_drop": 0.02,
                "mae_rise": 1.0,
                "rmse_rise": 1.0,
                "exact_grade_agreement_drop": 0.05,
                "adjacent_grade_agreement_drop": 0.05,
                "review_rate_rise": 0.1,
                "invalid_evidence_rate_rise": 0.02,
            },
        },
    }


def test_private_holdout_manifest_is_stable_and_omits_original_filenames(tmp_path):
    papers = tmp_path / "private-papers"
    papers.mkdir()
    (papers / "张三.docx").write_bytes(b"paper-a")
    (papers / "李四.pdf").write_bytes(b"paper-b")
    rows = [
        {"filename": "张三.docx", "total": 90, "items": {"C01": 9}},
        {"filename": "李四.pdf", "total": 70, "items": {"C01": 7}},
    ]

    first, mapping = build_holdout_manifest_from_rows(
        papers_dir=papers,
        rows=rows,
        dataset_id="teacher-holdout",
        dataset_version="v1",
    )
    second, _ = build_holdout_manifest_from_rows(
        papers_dir=papers,
        rows=list(reversed(rows)),
        dataset_id="teacher-holdout",
        dataset_version="v1",
    )

    serialized = json.dumps(first, ensure_ascii=False)
    assert "张三" not in serialized
    assert "李四" not in serialized
    assert set(mapping) == {"张三.docx", "李四.pdf"}
    assert (
        first["public_identity"]["dataset_manifest_sha256"]
        == second["public_identity"]["dataset_manifest_sha256"]
    )
    assert first["truth"] == second["truth"]


def test_holdout_manifest_rejects_path_escape_and_duplicate_content(tmp_path):
    papers = tmp_path / "papers"
    papers.mkdir()
    outside = tmp_path / "outside.docx"
    outside.write_bytes(b"outside")
    with pytest.raises(GateValidationError, match="escapes"):
        build_holdout_manifest_from_rows(
            papers_dir=papers,
            rows=[{"filename": "../outside.docx", "total": 80, "items": {}}],
            dataset_id="d",
            dataset_version="v1",
        )

    (papers / "a.docx").write_bytes(b"same")
    (papers / "b.docx").write_bytes(b"same")
    with pytest.raises(GateValidationError, match="duplicate paper content"):
        build_holdout_manifest_from_rows(
            papers_dir=papers,
            rows=[
                {"filename": "a.docx", "total": 80, "items": {}},
                {"filename": "b.docx", "total": 81, "items": {}},
            ],
            dataset_id="d",
            dataset_version="v1",
        )


def test_holdout_manifest_rejects_supported_papers_missing_from_truth(tmp_path):
    papers = tmp_path / "papers"
    papers.mkdir()
    (papers / "scored.docx").write_bytes(b"scored-paper")
    (papers / "silently-omitted.pdf").write_bytes(b"omitted-paper")

    with pytest.raises(GateValidationError, match="missing from teacher scores"):
        build_holdout_manifest_from_rows(
            papers_dir=papers,
            rows=[
                {
                    "filename": "scored.docx",
                    "total": 80,
                    "items": {"T01": 16},
                }
            ],
            dataset_id="d",
            dataset_version="v1",
        )


def test_holdout_manifest_ignores_non_paper_sidecar_files(tmp_path):
    papers = tmp_path / "papers"
    papers.mkdir()
    (papers / "scored.docx").write_bytes(b"scored-paper")
    (papers / "README.txt").write_text("private operator notes", encoding="utf-8")

    manifest, _mapping = build_holdout_manifest_from_rows(
        papers_dir=papers,
        rows=[
            {
                "filename": "scored.docx",
                "total": 80,
                "items": {"T01": 16},
            }
        ],
        dataset_id="d",
        dataset_version="v1",
    )

    assert manifest["public_identity"]["sample_count"] == 1


def test_holdout_manifest_rejects_unsupported_scored_file(tmp_path):
    papers = tmp_path / "papers"
    papers.mkdir()
    (papers / "not-a-paper.txt").write_text("plain text", encoding="utf-8")

    with pytest.raises(GateValidationError, match="unsupported paper format"):
        build_holdout_manifest_from_rows(
            papers_dir=papers,
            rows=[
                {
                    "filename": "not-a-paper.txt",
                    "total": 80,
                    "items": {"T01": 16},
                }
            ],
            dataset_id="d",
            dataset_version="v1",
        )


def test_complete_candidate_is_eligible_but_waits_for_human_approval():
    candidate = _candidate()

    assert candidate["reproducible"] is True
    assert candidate["gating_eligible"] is True
    assert candidate["gate_passed"] is False
    assert candidate["status"] == "candidate_awaiting_approval"
    serialized = json.dumps(candidate)
    assert "human_total" not in serialized
    assert "per_sample" not in serialized


def test_mock_or_dirty_revision_can_never_be_gating_eligible():
    mock = _candidate("mock")
    assert mock["gating_eligible"] is False
    assert any("mock provider" in issue for issue in mock["missing_or_invalid"])

    report = _report()
    run_identity, issues = summarize_run_identities(report)
    dirty = build_gate_candidate(
        evaluation_id="dirty",
        report=report,
        dataset_identity=_candidate()["dataset"],
        run_identity=run_identity,
        run_identity_issues=issues,
        anchors_identity=_candidate()["anchors"],
        model_identity=_candidate()["model"],
        prompt_identity={
            "prompt_version": "v",
            "prompt_source_manifest_sha256": _sha("f"),
        },
        repository_identity={
            "revision": "a" * 40,
            "branch": "main",
            "dirty": True,
            "lockfile": "uv.lock",
            "lockfile_sha256": _sha("0"),
        },
        private_artifacts=_candidate()["private_artifacts"],
        artifact_directory_external=True,
    )
    assert dirty["gating_eligible"] is False
    assert "code worktree must be clean" in dirty["missing_or_invalid"]


def test_approval_is_content_bound_and_passes_only_after_privacy_and_thresholds():
    candidate = _candidate()
    approval = _approval(candidate)

    final = finalize_gate(candidate, approval)

    assert final["gate_passed"] is True
    assert final["status"] == "passed"
    assert final["anchors"]["holdout_exclusion_proven"] is True
    assert final["manual_confirmations_pending"] == []

    wrong = deepcopy(approval)
    wrong["candidate_sha256"] = _sha("0")
    with pytest.raises(GateValidationError, match="exact candidate"):
        finalize_gate(candidate, wrong)

    privacy_missing = deepcopy(approval)
    privacy_missing["privacy_review"]["repository_scan_passed"] = False
    with pytest.raises(GateValidationError, match="privacy review"):
        finalize_gate(candidate, privacy_missing)


def test_approved_threshold_failure_is_not_a_pass():
    candidate = _candidate()
    approval = _approval(candidate)
    approval["accepted_baseline"]["acceptance_thresholds"]["minimum_qwk"] = 1.1

    final = finalize_gate(candidate, approval)

    assert final["gate_passed"] is False
    assert final["status"] == "failed_thresholds"
    assert any("qwk" in issue for issue in final["missing_or_invalid"])


def test_invalid_evidence_metric_is_required_thresholded_and_regression_checked():
    candidate = _candidate()
    missing = deepcopy(candidate)
    missing["evaluation"]["metrics"].pop("invalid_evidence_rate")
    missing.pop("candidate_sha256")
    missing["candidate_sha256"] = evaluation_sha256(missing)
    missing_final = finalize_gate(missing, _approval(missing))
    assert missing_final["gate_passed"] is False
    assert any(
        "invalid_evidence_rate" in issue
        for issue in missing_final["missing_or_invalid"]
    )

    approval = _approval(candidate)
    approval["accepted_baseline"]["acceptance_thresholds"][
        "maximum_invalid_evidence_rate"
    ] = -0.01
    with pytest.raises(GateValidationError, match="non-negative"):
        finalize_gate(candidate, approval)

    approved = finalize_gate(candidate, _approval(candidate))
    current = deepcopy(candidate["evaluation"]["metrics"])
    current["invalid_evidence_rate"] = 0.03
    issues = approved_regression_issues(current, approved)
    assert any("invalid_evidence_rate regressed" in issue for issue in issues)


def test_approved_regression_uses_frozen_tolerances():
    candidate = _candidate()
    approved = finalize_gate(candidate, _approval(candidate))
    current = deepcopy(candidate["evaluation"]["metrics"])
    current["qwk"] = 0.97
    current["mae"] = 2.1

    issues = approved_regression_issues(current, approved)

    assert any("qwk regressed" in issue for issue in issues)
    assert any("mae regressed" in issue for issue in issues)


def test_candidate_regression_reuses_only_same_dataset_rubric_and_anchors():
    candidate = _candidate()
    approved = finalize_gate(candidate, _approval(candidate))

    passed = build_regression_record(
        candidate=candidate,
        approved_record=approved,
    )
    assert passed["regression_passed"] is True

    changed = deepcopy(candidate)
    changed["anchors"]["manifest_sha256"] = _sha("0")
    changed.pop("candidate_sha256")
    changed["candidate_sha256"] = evaluation_sha256(changed)
    failed = build_regression_record(
        candidate=changed,
        approved_record=approved,
    )
    assert failed["regression_passed"] is False
    assert any("anchors.manifest_sha256 changed" in issue for issue in failed["issues"])


def test_parity_regression_archives_deltas_dimensions_and_explanations():
    baseline_candidate = _candidate()
    approved = finalize_gate(
        baseline_candidate,
        _approval(baseline_candidate),
    )
    candidate = deepcopy(baseline_candidate)
    candidate["evaluation_id"] = "gate-02-core-candidate"
    candidate["evaluation"]["metrics"]["mae"] = 1.25
    candidate["evaluation"]["per_criterion"]["C01"] = {
        "n": 1,
        "mae": 1.25,
        "bias": 0.25,
    }
    candidate["difference_explanations"] = [
        {
            "criterion_code": "C01",
            "reason_code": "EVIDENCE_CORRECTNESS_FIX",
            "summary": "Core rejects unsupported evidence and uses the authorized band.",
        }
    ]
    candidate.pop("candidate_sha256")
    candidate["candidate_sha256"] = evaluation_sha256(candidate)

    record = build_regression_record(
        candidate=candidate,
        approved_record=approved,
    )

    assert record["baseline_metrics"]["mae"] == 1.0
    assert record["metric_deltas"]["mae"] == pytest.approx(0.25)
    assert record["per_criterion"]["C01"]["bias_delta"] == pytest.approx(-0.75)
    assert record["difference_explanations"] == candidate[
        "difference_explanations"
    ]


def test_candidate_and_approved_record_hashes_are_revalidated():
    candidate = _candidate()
    approved = finalize_gate(candidate, _approval(candidate))

    tampered_candidate = deepcopy(candidate)
    tampered_candidate["evaluation"]["metrics"]["qwk"] = -1.0
    with pytest.raises(GateValidationError, match="candidate content"):
        build_regression_record(
            candidate=tampered_candidate,
            approved_record=approved,
        )

    tampered_approved = deepcopy(approved)
    tampered_approved["approval"]["approved_by"] = "someone-else"
    with pytest.raises(GateValidationError, match="approved reference content"):
        build_regression_record(
            candidate=candidate,
            approved_record=tampered_approved,
        )


def test_private_artifact_directory_must_be_outside_repository(tmp_path):
    repository = tmp_path / "repository"
    repository.mkdir()
    with pytest.raises(GateValidationError, match="outside"):
        ensure_external_artifact_dir(repository / "private", repository)

    external = tmp_path / "external"
    assert ensure_external_artifact_dir(external, repository) == external.resolve()


def test_cli_finalizes_exact_candidate_without_rerunning_evaluation(tmp_path):
    from backend.app.scripts import run_qwk_eval

    candidate = _candidate()
    candidate_path = tmp_path / "gate_candidate.json"
    approval_path = tmp_path / "approval.json"
    public_path = tmp_path / "approved.json"
    candidate_path.write_text(json.dumps(candidate), encoding="utf-8")
    approval_path.write_text(
        json.dumps(_approval(candidate)),
        encoding="utf-8",
    )

    exit_code = run_qwk_eval.main(
        [
            "--finalize-candidate",
            str(candidate_path),
            "--approval",
            str(approval_path),
            "--public-record-output",
            str(public_path),
        ]
    )

    assert exit_code == 0
    assert json.loads(public_path.read_text(encoding="utf-8"))["gate_passed"] is True
    assert (tmp_path / "gate_final.json").is_file()


def test_cli_release_preflight_does_not_require_or_invoke_model(
    monkeypatch, capsys
):
    from backend.app.scripts import run_qwk_eval

    rubric_version = SimpleNamespace(
        id="rubric-version-1",
        version_hash=_sha("a"),
        hash_scheme="rubric-content-v2",
    )
    public_identity = {
        "dataset_id": "teacher-holdout",
        "dataset_version": "2026-07-31-v2",
        "sample_count": 40,
        "dataset_manifest_sha256": _sha("b"),
        "sample_ids_sha256": _sha("c"),
        "paper_manifest_sha256": _sha("d"),
        "truth_sha256": _sha("e"),
    }
    monkeypatch.setattr(
        run_qwk_eval,
        "_validate_release_input_paths",
        lambda _args: None,
        raising=False,
    )
    monkeypatch.setattr(
        run_qwk_eval,
        "SessionLocal",
        lambda: nullcontext(object()),
    )
    monkeypatch.setattr(
        run_qwk_eval,
        "validate_release_rubric_and_scores",
        lambda _db, _rubric_id, _scores: rubric_version,
    )
    monkeypatch.setattr(
        run_qwk_eval,
        "build_holdout_manifest",
        lambda **_kwargs: ({"public_identity": public_identity}, {}),
    )
    monkeypatch.setattr(
        run_qwk_eval,
        "_run_release_gate",
        lambda _args: pytest.fail("preflight must not run the model gate"),
    )
    monkeypatch.setattr(
        run_qwk_eval,
        "build_labeled_eval",
        lambda *_args, **_kwargs: pytest.fail(
            "preflight must not invoke candidate scoring"
        ),
    )

    exit_code = run_qwk_eval.main(
        [
            "--release-gate",
            "--preflight-only",
            "--rubric-id",
            "rubric-1",
            "--papers-dir",
            "/external/private-papers",
            "--scores",
            "/external/teacher-scores.xlsx",
            "--dataset-id",
            "teacher-holdout",
            "--dataset-version",
            "2026-07-31-v2",
        ]
    )

    output = capsys.readouterr().out
    assert exit_code == 0
    assert '"status": "ready_for_real_model_run"' in output
    assert '"gate_passed": false' in output
    assert '"gating_eligible": false' in output
    assert rubric_version.version_hash in output
    assert public_identity["dataset_manifest_sha256"] in output
    assert "private-papers" not in output
    assert "teacher-scores.xlsx" not in output


def test_cli_preflight_only_requires_release_gate(capsys):
    from backend.app.scripts import run_qwk_eval

    with pytest.raises(SystemExit) as exc_info:
        run_qwk_eval.main(["--preflight-only"])

    assert exc_info.value.code == 2
    assert "--preflight-only is only valid with --release-gate" in (
        capsys.readouterr().err
    )


def test_cli_preflight_resolves_database_gate_profile_before_required_args(
    monkeypatch,
):
    from backend.app.scripts import run_qwk_eval

    seen = []

    def apply_profile(args):
        seen.append(args.gate_profile_id)
        args.rubric_id = "rubric-from-profile"
        args.dataset_id = "dataset-from-profile"
        args.dataset_version = "v1"

    monkeypatch.setattr(
        run_qwk_eval,
        "_apply_release_gate_profile",
        apply_profile,
        raising=False,
    )
    monkeypatch.setattr(
        run_qwk_eval,
        "_run_release_preflight",
        lambda args: 0,
    )

    exit_code = run_qwk_eval.main(
        [
            "--release-gate",
            "--preflight-only",
            "--gate-profile-id",
            "profile-1",
            "--papers-dir",
            "/external/papers",
            "--scores",
            "/external/scores.xlsx",
        ]
    )

    assert exit_code == 0
    assert seen == ["profile-1"]


def test_database_gate_profile_populates_identity_and_rejects_overrides(
    monkeypatch,
):
    from backend.app.scripts import run_qwk_eval

    profile = SimpleNamespace(
        id="profile-1",
        status="active",
        gate_key="GATE-02",
        rubric_id="rubric-1",
        dataset_identity={
            "dataset_id": "dataset-1",
            "dataset_version": "v1",
        },
        model_identity={
            "identity_method": "provider_immutable_revision",
            "immutable_revision": "fixed-model@2026-08-02",
        },
    )
    monkeypatch.setattr(
        run_qwk_eval,
        "SessionLocal",
        lambda: nullcontext(object()),
    )
    monkeypatch.setattr(
        run_qwk_eval.release_gates,
        "resolve_profile",
        lambda _db, _profile_id: profile,
    )
    args = SimpleNamespace(
        gate_profile_id="profile-1",
        rubric_id=None,
        dataset_id=None,
        dataset_version=None,
        model_artifact=None,
        immutable_model_revision=None,
    )

    run_qwk_eval._apply_release_gate_profile(args)

    assert args.rubric_id == "rubric-1"
    assert args.dataset_id == "dataset-1"
    assert args.dataset_version == "v1"
    assert args.immutable_model_revision == "fixed-model@2026-08-02"
    assert args.gate_key == "GATE-02"

    args.rubric_id = "rubric-override"
    with pytest.raises(GateValidationError, match="rubric-id does not match"):
        run_qwk_eval._apply_release_gate_profile(args)


def test_cli_loads_strict_public_difference_explanations(tmp_path):
    from backend.app.scripts import run_qwk_eval

    path = tmp_path / "differences.json"
    expected = [
        {
            "criterion_code": "T01",
            "reason_code": "EVIDENCE_CORRECTNESS_FIX",
            "summary": "Core rejects unsupported evidence.",
        }
    ]
    path.write_text(json.dumps(expected), encoding="utf-8")
    assert run_qwk_eval._load_difference_explanations(path) == expected

    path.write_text(json.dumps([{"criterion_code": "T01"}]), encoding="utf-8")
    with pytest.raises(GateValidationError, match="difference_explanations"):
        run_qwk_eval._load_difference_explanations(path)


def test_database_gate_candidate_is_registered_against_exact_profile(
    monkeypatch,
):
    from backend.app.scripts import run_qwk_eval

    expected_run = SimpleNamespace(id="run-1")
    captured = {}
    monkeypatch.setattr(
        run_qwk_eval,
        "SessionLocal",
        lambda: nullcontext(object()),
    )

    def register_run(db, profile_id, candidate):
        captured.update(
            db=db,
            profile_id=profile_id,
            candidate=candidate,
        )
        return expected_run

    monkeypatch.setattr(
        run_qwk_eval.release_gates,
        "register_run",
        register_run,
    )
    args = SimpleNamespace(gate_profile_id="profile-1")
    candidate = {"candidate_sha256": _sha("f")}

    result = run_qwk_eval._register_db_gate_candidate(args, candidate)

    assert result is expected_run
    assert captured["profile_id"] == "profile-1"
    assert captured["candidate"] is candidate


def test_release_preflight_requires_published_version_and_complete_truth(
    client, tmp_path
):
    payload = {
        "name": "发布门禁评分标准",
        "version": "v1.0",
        "total_score": 10,
        "criteria": [
            {
                "code": "C01",
                "name": "研究质量",
                "max_score": 10,
                "criterion_type": "deterministic",
                "scoring_mode": "deductive",
                "deduction_rules_structured": [
                    {
                        "match": "研究质量不足",
                        "points": 10,
                        "reason": "研究质量不足扣分",
                        "checker_key": "thesis.legacy_required_fields.v1",
                        "checker_params": {
                            "criterion_code": "C01",
                            "applies_to": "global",
                        },
                    }
                ],
                "display_order": 1,
            }
        ],
    }
    created = client.post("/api/rubrics", json=payload)
    assert created.status_code == 200, created.text
    rubric_id = created.json()["id"]

    workbook = Workbook()
    workbook.active.append(["文件名", "总分", "C01"])
    workbook.active.append(["sample.docx", 8, 8])
    scores_path = tmp_path / "scores.xlsx"
    workbook.save(scores_path)

    with client.session_factory() as db:
        with pytest.raises(GateValidationError, match="must be published"):
            validate_release_rubric_and_scores(
                db,
                rubric_id,
                scores_path,
                expected_codes=("C01",),
            )

    _published, identity = publish_rubric_via_api(client, rubric_id)
    with client.session_factory() as db:
        version = validate_release_rubric_and_scores(
            db,
            rubric_id,
            scores_path,
            expected_codes=("C01",),
        )
        assert version.id == identity["rubric_version_id"]

    incomplete = Workbook()
    incomplete.active.append(["文件名", "总分"])
    incomplete.active.append(["sample.docx", 8])
    incomplete_path = tmp_path / "incomplete.xlsx"
    incomplete.save(incomplete_path)
    with client.session_factory() as db:
        with pytest.raises(GateValidationError, match="criterion mismatch"):
            validate_release_rubric_and_scores(
                db,
                rubric_id,
                incomplete_path,
                expected_codes=("C01",),
            )

    assert PGS8_CRITERION_CODES == (
        "T01",
        "T02",
        "T03",
        "T04",
        "T05",
        "T06",
    )


def test_release_preflight_reports_only_safe_draft_blocker_identifiers(
    client, tmp_path
):
    payload = {
        "name": "待修复发布标准",
        "version": "v1.0",
        "total_score": 10,
        "criteria": [
            {
                "code": "C01",
                "name": "研究质量",
                "max_score": 10,
                "criterion_type": "deterministic",
                "scoring_mode": "deductive",
                "deduction_rules_structured": [
                    {
                        "match": "不足",
                        "points": 10,
                        "reason": "不足扣分",
                        "checker_key": "thesis.legacy_required_fields.v1",
                        "checker_params": {
                            "criterion_code": "C01",
                            "applies_to": "global",
                        },
                    }
                ],
                "display_order": 1,
            }
        ],
    }
    created = client.post("/api/rubrics", json=payload)
    assert created.status_code == 200, created.text
    rubric_id = created.json()["id"]
    with client.session_factory() as db:
        compilation = db.scalar(
            select(RubricCompilation).where(
                RubricCompilation.rubric_id == rubric_id
            )
        )
        assert compilation is not None
        compilation.status = "blocked"
        compilation.blockers = [
            {
                "code": "MISSING_CRITERION_DESCRIPTION",
                "criterion_code": "C01",
                "message": "private source text must not be echoed",
            }
        ]
        db.commit()

    scores = Workbook()
    scores.active.append(["文件名", "总分", "C01"])
    scores.active.append(["sample.docx", 8, 8])
    scores_path = tmp_path / "scores.xlsx"
    scores.save(scores_path)

    with client.session_factory() as db:
        with pytest.raises(GateValidationError) as exc_info:
            validate_release_rubric_and_scores(
                db,
                rubric_id,
                scores_path,
                expected_codes=("C01",),
            )

    message = str(exc_info.value)
    assert "MISSING_CRITERION_DESCRIPTION(C01)" in message
    assert "private source text" not in message
