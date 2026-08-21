import json
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select

from backend.app.core.config import settings
from backend.app.db import models
from backend.app.services.spreadsheet import writer as writer_module
from backend.app.services.scoring.profiles.thesis import ThesisProfile
from backend.app.tests.conftest import make_sample_docx
from backend.app.tests.conftest import publish_rubric_via_api


def test_core_api_flow(client, monkeypatch):
    artifact_projections = []
    spreadsheet_projections = []
    item_review_guards = []
    run_review_guards = []
    original_artifact = ThesisProfile.build_artifact_projection
    original_spreadsheet = ThesisProfile.build_spreadsheet_projection
    original_item_guard = ThesisProfile.assert_ordinary_item_override_allowed
    original_run_guard = ThesisProfile.assert_review_submission_allowed

    def record_artifact(profile, **kwargs):
        projection = original_artifact(profile, **kwargs)
        artifact_projections.append(projection)
        return projection

    def record_spreadsheet(profile, **kwargs):
        projection = original_spreadsheet(profile, **kwargs)
        spreadsheet_projections.append(projection)
        return projection

    def record_item_guard(profile, **kwargs):
        item_review_guards.append(kwargs["item"].id)
        return original_item_guard(profile, **kwargs)

    def record_run_guard(profile, **kwargs):
        run_review_guards.append(kwargs["run"].id)
        return original_run_guard(profile, **kwargs)

    monkeypatch.setattr(ThesisProfile, "build_artifact_projection", record_artifact)
    monkeypatch.setattr(ThesisProfile, "build_spreadsheet_projection", record_spreadsheet)
    monkeypatch.setattr(
        ThesisProfile,
        "assert_ordinary_item_override_allowed",
        record_item_guard,
    )
    monkeypatch.setattr(
        ThesisProfile,
        "assert_review_submission_allowed",
        record_run_guard,
    )
    rubric_payload = {
        "name": "测试评分标准",
        "version": "v1.0",
        "total_score": 30,
        "criteria": [
            {
                "code": "C01",
                "name": "研究方法",
                "max_score": 15,
                "criterion_type": "deterministic",
                "evidence_hints": ["研究方法", "实验设计", "数据来源"],
                "deduction_rules": ["方法说明不足扣分"],
                "scoring_mode": "deductive",
                "deduction_rules_structured": [
                    {
                        "match": "方法说明不足",
                        "points": 15,
                        "reason": "方法说明不足扣分",
                        "checker_key": "thesis.legacy_required_fields.v1",
                        "checker_params": {
                            "criterion_code": "C01",
                            "applies_to": "global",
                        },
                    }
                ],
                "display_order": 1,
            },
            {
                "code": "C02",
                "name": "参考文献",
                "max_score": 15,
                "criterion_type": "deterministic",
                "evidence_hints": ["参考文献", "引用"],
                "deduction_rules": ["参考文献不足扣分"],
                "scoring_mode": "deductive",
                "deduction_rules_structured": [
                    {
                        "match": "参考文献不足",
                        "points": 15,
                        "reason": "参考文献不足扣分",
                        "checker_key": "thesis.legacy_required_fields.v1",
                        "checker_params": {
                            "criterion_code": "C02",
                            "applies_to": "global",
                        },
                    }
                ],
                "display_order": 2,
            },
        ],
    }
    rubric_response = client.post("/api/rubrics", json=rubric_payload)
    assert rubric_response.status_code == 200, rubric_response.text
    rubric_id = rubric_response.json()["id"]

    duplicate_rubric_response = client.post("/api/rubrics", json=rubric_payload)
    assert duplicate_rubric_response.status_code == 400
    assert duplicate_rubric_response.json()["detail"] == "rubric name and version already exist"

    rubric_update_payload = {
        "name": "测试评分标准-草稿",
        "version": "v1.0-draft",
        "description": "调整后的草稿标准",
        "total_score": 30,
        "criteria": [
            {
                "code": "C01",
                "name": "研究方法",
                "max_score": 12,
                "evidence_hints": ["研究方法", "实验设计", "数据来源"],
                "deduction_rules": ["方法说明不足扣分"],
                "display_order": 1,
            },
            {
                "code": "C02",
                "name": "参考文献",
                "max_score": 18,
                "evidence_hints": ["参考文献", "引用"],
                "deduction_rules": ["参考文献不足扣分"],
                "display_order": 2,
            },
        ],
    }
    # M4 create immediately establishes an immutable provenance graph.  Its
    # content is edited through AtomicRule/recompilation (or clone-for-edit),
    # never through the legacy whole-rubric PATCH facade.
    rubric_update_response = client.patch(
        "/api/rubrics/%s" % rubric_id,
        json=rubric_update_payload,
    )
    assert rubric_update_response.status_code == 400, rubric_update_response.text
    assert "AtomicRule" in rubric_update_response.json()["detail"]
    unchanged = client.get("/api/rubrics/%s" % rubric_id).json()
    assert unchanged["name"] == rubric_payload["name"]
    assert unchanged["criteria"][0]["max_score"] == 15

    publish_response, _identity = publish_rubric_via_api(client, rubric_id)
    assert publish_response["status"] == "published"

    update_published_response = client.patch("/api/rubrics/%s" % rubric_id, json={"description": "发布后修改"})
    assert update_published_response.status_code == 400

    clone_response = client.post("/api/rubrics/%s/clone" % rubric_id, json={"new_version": "v1.1"})
    assert clone_response.status_code == 200, clone_response.text
    cloned = clone_response.json()
    assert cloned["version"] == "v1.1"
    assert cloned["status"] == "draft"
    assert len(cloned["criteria"]) == len(rubric_payload["criteria"])

    duplicate_clone_response = client.post("/api/rubrics/%s/clone" % rubric_id, json={"new_version": "v1.1"})
    assert duplicate_clone_response.status_code == 400

    batch_response = client.post(
        "/api/batches",
        json={"name": "测试批次", "rubric_id": rubric_id, "department": "计算机学院", "major": "软件工程"},
    )
    assert batch_response.status_code == 200, batch_response.text
    batch_id = batch_response.json()["id"]

    update_batch_response = client.patch(
        "/api/batches/%s" % batch_id,
        json={"name": "测试批次-更新", "department": "人工智能学院", "paper_type": "本科毕业论文"},
    )
    assert update_batch_response.status_code == 200, update_batch_response.text
    updated_batch = update_batch_response.json()
    assert updated_batch["name"] == "测试批次-更新"
    assert updated_batch["department"] == "人工智能学院"

    docx = make_sample_docx()
    upload_response = client.post(
        "/api/papers/upload",
        data={"batch_id": batch_id},
        files={"file": ("sample.docx", docx.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert upload_response.status_code == 200, upload_response.text
    paper = upload_response.json()
    assert paper["status"] == "parsed"

    update_paper_response = client.patch(
        "/api/papers/%s" % paper["id"],
        json={
            "student_id": "20260001",
            "student_name": "张三",
            "title": "校正后的论文题目",
            "department": "人工智能学院",
            "major": "智能科学与技术",
            "advisor": "李老师",
        },
    )
    assert update_paper_response.status_code == 200, update_paper_response.text
    updated_paper = update_paper_response.json()
    assert updated_paper["student_name"] == "张三"
    assert updated_paper["title"] == "校正后的论文题目"

    change_rubric_response = client.patch("/api/batches/%s" % batch_id, json={"rubric_id": cloned["id"]})
    assert change_rubric_response.status_code == 400

    papers_response = client.get("/api/papers?batch_id=%s" % batch_id)
    assert papers_response.status_code == 200
    assert papers_response.json()[0]["id"] == paper["id"]
    assert papers_response.json()[0]["student_id"] == "20260001"

    parsed_response = client.get("/api/papers/%s/parsed" % paper["id"])
    assert parsed_response.status_code == 200
    assert parsed_response.json()["parsed"]["sections"]

    chunks_response = client.get("/api/papers/%s/chunks" % paper["id"])
    assert chunks_response.status_code == 200
    chunks = chunks_response.json()
    assert chunks
    assert chunks[0]["paper_id"] == paper["id"]

    reparse_response = client.post("/api/papers/%s/parse" % paper["id"])
    assert reparse_response.status_code == 200, reparse_response.text
    reparsed_paper = reparse_response.json()
    assert reparsed_paper["status"] == "parsed"
    assert reparsed_paper["parse_quality"] is not None

    chunks_after_reparse_response = client.get("/api/papers/%s/chunks" % paper["id"])
    assert chunks_after_reparse_response.status_code == 200
    chunk_ids = {chunk["id"] for chunk in chunks_after_reparse_response.json()}
    assert chunk_ids

    scoring_response = client.post("/api/papers/%s/score" % paper["id"])
    assert scoring_response.status_code == 200, scoring_response.text
    run = scoring_response.json()
    assert run["final_total_score"] is not None

    runs_response = client.get("/api/scoring-runs?paper_id=%s" % paper["id"])
    assert runs_response.status_code == 200
    assert runs_response.json()[0]["id"] == run["id"]

    retry_response = client.post("/api/scoring-runs/%s/retry" % run["id"])
    assert retry_response.status_code == 200, retry_response.text
    retry_run = retry_response.json()
    assert retry_run["id"] != run["id"]
    assert retry_run["paper_id"] == paper["id"]

    runs_after_retry_response = client.get("/api/scoring-runs?paper_id=%s" % paper["id"])
    assert runs_after_retry_response.status_code == 200
    assert len(runs_after_retry_response.json()) >= 2

    items_response = client.get("/api/scoring-runs/%s/items" % run["id"])
    assert items_response.status_code == 200
    items = items_response.json()
    assert len(items) == 2
    assert items[0]["criterion_name"] in {"研究方法", "参考文献"}
    with client.session_factory() as db:
        stored_items = [db.get(models.ScoreItem, item["id"]) for item in items]
        assert all(
            item.rule_results_schema_version == "rule-results@2"
            for item in stored_items
        )
        rule_results = [
            result
            for item in stored_items
            for result in item.rule_results
        ]
    evidence_refs = [
        evidence_ref
        for result in rule_results
        for evidence_ref in result["evidence_refs"]
    ]
    assert evidence_refs
    assert all(
        evidence_ref["evidence_type"] == "deterministic_observation"
        for evidence_ref in evidence_refs
    )
    assert all(
        evidence_ref["locator"]["kind"] == "document_structure"
        and evidence_ref["locator"]["structure_code"].startswith("required_owner:")
        for evidence_ref in evidence_refs
    )

    second_docx = make_sample_docx()
    second_upload_response = client.post(
        "/api/papers/bulk-upload",
        data={"batch_id": batch_id},
        files=[
            (
                "files",
                (
                    "sample-2.docx",
                    second_docx.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
            )
        ],
    )
    assert second_upload_response.status_code == 200, second_upload_response.text
    assert second_upload_response.json()[0]["status"] == "parsed"
    second_paper_id = second_upload_response.json()[0]["id"]

    batch_score_response = client.post("/api/batches/%s/score" % batch_id)
    assert batch_score_response.status_code == 200, batch_score_response.text
    batch_score = batch_score_response.json()
    assert batch_score["total_papers"] == 2
    assert batch_score["scored_count"] == 1
    assert batch_score["skipped_count"] == 1
    assert batch_score["failed_count"] == 0
    assert len(batch_score["run_ids"]) == 1

    start_response = client.post("/api/batches/%s/start" % batch_id)
    assert start_response.status_code == 200, start_response.text
    start_result = start_response.json()
    assert start_result["scored_count"] == 0
    assert start_result["skipped_count"] == 2

    before_rescore = client.get(
        "/api/scoring-runs", params={"batch_id": batch_id}
    ).json()
    batch_rescore_response = client.post(
        "/api/batches/%s/score?rescore=true" % batch_id
    )
    assert batch_rescore_response.status_code == 200, batch_rescore_response.text
    batch_rescore = batch_rescore_response.json()
    assert batch_rescore["scored_count"] == 2
    assert batch_rescore["skipped_count"] == 0
    assert batch_rescore["failed_count"] == 0
    assert len(set(batch_rescore["run_ids"])) == 2
    assert set(batch_rescore["run_ids"]).isdisjoint(
        {item["id"] for item in before_rescore}
    )
    with client.session_factory() as db:
        generations = {
            target: sorted(
                run.rescore_generation
                for run in db.scalars(
                    select(models.ScoringRun).where(
                        models.ScoringRun.paper_id == target
                    )
                ).all()
            )
            for target in (paper["id"], second_paper_id)
        }
    assert generations[paper["id"]] == [0, 1, 2]
    assert generations[second_paper_id] == [0, 1]

    summary_response = client.get("/api/batches/%s/summary" % batch_id)
    assert summary_response.status_code == 200, summary_response.text
    summary = summary_response.json()
    assert summary["batch"]["id"] == batch_id
    assert len(summary["papers"]) == 2
    assert all(item["latest_run_id"] for item in summary["papers"])
    assert summary["run_stats"]["scored"] + summary["run_stats"]["reviewing"] + summary["run_stats"]["reviewed"] == 2

    patch_response = client.patch(
        "/api/score-items/%s" % items[0]["id"],
        json={"final_score": 12, "reason": "教师复核后调整"},
    )
    assert patch_response.status_code == 200, patch_response.text

    review_response = client.post("/api/scoring-runs/%s/review" % run["id"], json={"reason": "复核完成"})
    assert review_response.status_code == 200
    assert review_response.json()["status"] == "reviewed"

    logs_response = client.get("/api/scoring-runs/%s/review-logs" % run["id"])
    assert logs_response.status_code == 200
    assert len(logs_response.json()) >= 2

    monkeypatch.setattr(settings, "SHEET_WRITER_PROVIDER", "google_sheets")
    monkeypatch.setattr(settings, "GOOGLE_SHEETS_WEBAPP_URL", "https://script.google.test/exec")
    monkeypatch.setattr(settings, "GOOGLE_SHEETS_WEBAPP_SECRET", "test-secret")
    monkeypatch.setattr(writer_module.httpx, "Client", lambda timeout: FakeGoogleSheetsClient())
    write_sheet_response = client.post("/api/scoring-runs/%s/write-sheet" % run["id"], json={"target_id": "dev-sheet"})
    assert write_sheet_response.status_code == 200, write_sheet_response.text
    write_sheet_log = write_sheet_response.json()
    assert write_sheet_log["target_type"] == "google_sheets"
    assert write_sheet_log["target_id"] == "dev-sheet"
    assert write_sheet_log["response"]["provider_response"]["ok"]
    assert write_sheet_log["response"]["request"]["summary_row"]["学号"] == "20260001"
    assert write_sheet_log["response"]["request"]["summary_row"]["复核意见"] == "复核完成"
    assert len(write_sheet_log["response"]["request"]["detail_rows"]) == 2
    monkeypatch.setattr(settings, "SHEET_FALLBACK_TO_MOCK", False)
    monkeypatch.setattr(writer_module.httpx, "Client", lambda timeout: FakeBadGoogleSheetsClient())
    failed_write_response = client.post("/api/scoring-runs/%s/write-sheet" % run["id"], json={"target_id": "dev-sheet"})
    assert failed_write_response.status_code == 502
    assert "google sheets write failed" in failed_write_response.json()["detail"]
    monkeypatch.setattr(settings, "SHEET_FALLBACK_TO_MOCK", True)
    monkeypatch.setattr(settings, "SHEET_WRITER_PROVIDER", "mock")

    export_response = client.get("/api/batches/%s/export.xlsx" % batch_id)
    assert export_response.status_code == 200
    assert export_response.content[:2] == b"PK"
    workbook = load_workbook(BytesIO(export_response.content))
    summary_sheet = workbook["总分表"]
    summary_headers = [cell.value for cell in summary_sheet[1]]
    summary_rows = [[cell.value for cell in row] for row in summary_sheet.iter_rows(min_row=2)]
    assert summary_sheet.freeze_panes == "A2"
    assert summary_sheet.auto_filter.ref == summary_sheet.dimensions
    assert summary_sheet.column_dimensions["F"].width >= 38
    assert summary_sheet["K2"].alignment.wrap_text is True
    assert summary_sheet["A1"].font.bold is True
    detail_sheet = workbook["评分明细表"]
    assert detail_sheet.freeze_panes == "A2"
    assert detail_sheet.column_dimensions["I"].width >= 64
    assert detail_sheet["I2"].alignment.wrap_text is True
    assert "复核意见" in summary_headers
    reviewed_row = next(row for row in summary_rows if "20260001" in row and "校正后的论文题目" in row)
    assert reviewed_row[summary_headers.index("复核意见")] == "复核完成"

    export_logs_response = client.get("/api/export-logs?batch_id=%s" % batch_id)
    assert export_logs_response.status_code == 200
    export_logs = export_logs_response.json()
    assert len(export_logs) >= 3
    assert export_logs[0]["target_type"] == "excel"

    run_export_logs_response = client.get("/api/export-logs?run_id=%s" % run["id"])
    assert run_export_logs_response.status_code == 200
    assert run_export_logs_response.json()

    report_response = client.get("/api/scoring-runs/%s/report" % run["id"])
    assert report_response.status_code == 200
    assert "毕业论文智能评分报告" in report_response.text
    assert "校正后的论文题目" in report_response.text
    assert "人工复核记录" in report_response.text
    assert "教师复核后调整" in report_response.text
    assert "复核完成" in report_response.text
    assert artifact_projections and artifact_projections[-1]["profile_key"] == "thesis"
    assert spreadsheet_projections
    assert all(item["profile_key"] == "thesis" for item in spreadsheet_projections)
    assert item_review_guards == [items[0]["id"]]
    assert run_review_guards == [run["id"]]


class FakeGoogleSheetsClient:
    def post(self, url, json):
        assert url == "https://script.google.test/exec"
        assert json["secret"] == "test-secret"
        assert json["target_id"] == "dev-sheet"
        assert json["payload"]["summary_row"]["学号"] == "20260001"
        return FakeGoogleSheetsResponse()


class FakeGoogleSheetsResponse:
    def raise_for_status(self):
        return None

    def json(self):
        return {
            "ok": True,
            "spreadsheet_id": "dev-sheet",
            "spreadsheet_url": "https://docs.google.com/spreadsheets/d/dev-sheet/edit",
            "summary_rows_written": 1,
            "detail_rows_written": 2,
        }


class FakeBadGoogleSheetsClient:
    def post(self, url, json):
        return FakeBadGoogleSheetsResponse()


class FakeBadGoogleSheetsResponse:
    def raise_for_status(self):
        return None

    def json(self):
        return {"ok": False, "error": "permission denied"}
