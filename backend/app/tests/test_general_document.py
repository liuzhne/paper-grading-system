"""通用文档评分：类型的选择、流转、评分、解析、导出与报告（通用文档评分改造方案）。

全部离线：Mock 模型，内存 SQLite。论文路径的行为由既有测试覆盖，这里只补“论文不变”的关键断言。
"""

from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import Workbook
from openpyxl import load_workbook

from backend.app.api import actions
from backend.app.db import models
from backend.app.services.assistant import views
from backend.app.services.rubrics.executable_validator import DEFAULT_BUSINESS_PROFILE_REGISTRY
from backend.app.services.rubrics.executable_validator import DEFAULT_CHECKER_REGISTRY
from backend.app.services.scoring import engine
from backend.app.services.scoring.profiles.general_document import GeneralDocumentProfile
from backend.app.services.scoring.profiles.registry import get_profile_by_key
from backend.app.services.scoring.profiles.registry import get_workbench_profile
from backend.app.services.scoring.profiles.registry import profile_key_for_batch
from backend.app.services.scoring.profiles.thesis import THESIS_CHECKER_KEYS
from backend.app.services.scoring.profiles.thesis import ThesisProfile
from backend.app.services.spreadsheet import excel
from backend.app.tests.conftest import make_sample_docx
from backend.app.tests.conftest import publish_rubric_via_api

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _rules_xlsx():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "评分规则"
    sheet.append(["编号", "评分项", "分值", "评分说明", "扣分规则"])
    sheet.append(["C01", "方案完整性", 60, "方案覆盖需求、实施计划与验收标准。", "缺少实施计划扣3分"])
    sheet.append(["C02", "表达规范", 40, "结构清晰，术语前后一致。", "术语不一致扣2分"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _import(client, *, profile=None, name=None):
    data = {"name": name or "方案评审标准-%s" % uuid4().hex[:6], "version": "v1"}
    if profile is not None:
        data["business_profile_key"] = profile
    return client.post(
        "/api/rubrics/import-sessions",
        data=data,
        files={"rules_file": ("rules.xlsx", _rules_xlsx(), XLSX)},
    )


def _published_rubric(client, *, profile):
    created = _import(client, profile=profile)
    assert created.status_code == 201, created.text
    confirmed = client.post(
        f"/api/rubrics/import-sessions/{created.json()['id']}/confirm",
        json={"expected_state_version": created.json()["state_version"], "idempotency_key": uuid4().hex},
    )
    assert confirmed.status_code == 200, confirmed.text
    rubric_id = confirmed.json()["rubric"]["id"]
    publish_rubric_via_api(client, rubric_id)
    return rubric_id


def _batch_with_paper(client, rubric_id, *, file_name="数据中心扩容方案.docx"):
    batch = client.post("/api/batches", json={
        "name": "第三批采购", "rubric_id": rubric_id, "department": "信息中心", "major": "基础设施",
    })
    assert batch.status_code == 200, batch.text
    upload = client.post(
        "/api/papers/upload",
        data={"batch_id": batch.json()["id"]},
        files={"file": (file_name, make_sample_docx(), DOCX)},
    )
    assert upload.status_code == 200, upload.text
    return batch.json(), upload.json()


# --- 类型注册 ------------------------------------------------------------------------------

def test_general_document_profile_reuses_thesis_machinery_with_its_own_identity():
    profile = get_profile_by_key("general_document")
    assert isinstance(profile, GeneralDocumentProfile)
    assert profile.profile_version == "general-document-profile@1"
    assert profile.prompt_version != ThesisProfile.prompt_version  # 缓存键分开
    assert profile.semantic_coherence is False and ThesisProfile.semantic_coherence is True
    assert profile.select_prompt_metadata(
        metadata={"title": "方案", "student_id": "x", "advisor": "y"}
    ) == {"title": "方案"}

    registry = profile.build_checker_registry()
    manifest = registry.manifest(profile_key="general_document", document_schema_version="document-snapshot@1")
    assert THESIS_CHECKER_KEYS["research_conclusion"] not in manifest
    assert set(manifest) == {
        *(key for name, key in THESIS_CHECKER_KEYS.items() if name != "research_conclusion"),
        "thesis.legacy_required_fields.v1",
    }
    registry.resolve(
        checker_key=THESIS_CHECKER_KEYS["text_length"], checker_version="1.0.0",
        checker_params={"minimum_chars": 1, "maximum_chars": 100000}, profile_key="general_document",
        document_schema_version="document-snapshot@1",
    )


def test_thesis_checker_registrations_are_unchanged():
    """论文的注册内容不变：仍是 7 个检查器、只支持 thesis。"""

    registry = ThesisProfile().build_checker_registry()
    schema = "document-snapshot@1"
    assert len(registry.manifest(profile_key="thesis", document_schema_version=schema)) == 7
    assert registry.manifest(profile_key="general_document", document_schema_version=schema) == {}


def test_publication_registries_accept_general_document():
    assert DEFAULT_BUSINESS_PROFILE_REGISTRY["general_document"]["profile_key"] == "general_document"
    reused = DEFAULT_CHECKER_REGISTRY.publication_registration(THESIS_CHECKER_KEYS["structure"])
    assert reused["supported_profiles"] == {"thesis", "general_document"}
    thesis_only = DEFAULT_CHECKER_REGISTRY.publication_registration(THESIS_CHECKER_KEYS["research_conclusion"])
    assert thesis_only["supported_profiles"] == {"thesis"}


def test_workbench_profile_rejects_technical_proposal():
    assert isinstance(get_workbench_profile("thesis"), ThesisProfile)
    assert isinstance(get_workbench_profile("general_document"), GeneralDocumentProfile)
    with pytest.raises(ValueError):
        get_workbench_profile("technical_proposal")


# --- 导入 ----------------------------------------------------------------------------------

def test_import_session_records_the_chosen_type(client):
    default = _import(client)
    assert default.status_code == 201, default.text
    assert default.json()["business_profile_key"] == "thesis"

    general = _import(client, profile="general_document")
    assert general.status_code == 201, general.text
    assert general.json()["business_profile_key"] == "general_document"
    reread = client.get(f"/api/rubrics/import-sessions/{general.json()['id']}")
    assert reread.json()["business_profile_key"] == "general_document"

    rejected = _import(client, profile="technical_proposal")
    assert rejected.status_code == 422
    assert "不支持的文档类型" in rejected.json()["detail"]


# --- 全链路 --------------------------------------------------------------------------------

def test_general_document_end_to_end(client, monkeypatch):
    semantic_calls = []
    monkeypatch.setattr(engine, "analyze_semantic_coherence", lambda parsed, scorer: semantic_calls.append(1) or [])

    rubric_id = _published_rubric(client, profile="general_document")
    listed = {item["id"]: item for item in client.get("/api/rubrics").json()}
    assert listed[rubric_id]["business_profile_key"] == "general_document"
    assert client.get(f"/api/rubrics/{rubric_id}").json()["business_profile_key"] == "general_document"

    batch, paper = _batch_with_paper(client, rubric_id)
    assert batch["business_profile_key"] == "general_document"
    assert client.get(f"/api/batches/{batch['id']}").json()["business_profile_key"] == "general_document"
    listed_batches = {item["id"]: item for item in client.get("/api/batches").json()}
    assert listed_batches[batch["id"]]["business_profile_key"] == "general_document"

    # 通用文档不从封面猜学号、姓名（样例封面上有“姓名：张三 / 学号：20260001”）。
    parsed = client.get(f"/api/papers/{paper['id']}").json()
    assert parsed["student_id"] is None and parsed["student_name"] is None
    assert parsed["title"]

    scored = client.post(f"/api/papers/{paper['id']}/score")
    assert scored.status_code == 200, scored.text
    run_id = scored.json()["id"]
    with client.session_factory() as session:
        run = session.get(models.ScoringRun, run_id)
        assert run.business_profile_key == "general_document"
        assert run.business_profile_version == "general-document-profile@1"
        assert run.prompt_version == GeneralDocumentProfile.prompt_version
    assert semantic_calls == []  # 不跑按论文写的语义一致性

    exported = client.get(f"/api/batches/{batch['id']}/export.xlsx")
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(BytesIO(exported.content))
    summary = [cell.value for cell in workbook["总分表"][1]]
    assert summary[:6] == ["批次名称", "文件名", "标题", "归属单位", "类别", "评分标准版本"]
    assert "学号" not in summary and "论文题目" not in summary
    row = [cell.value for cell in workbook["总分表"][2]]
    assert row[1] == "数据中心扩容方案.docx" and row[3] == "信息中心" and row[4] == "基础设施"
    detail = [cell.value for cell in workbook["评分明细表"][1]]
    assert detail[:3] == ["文件名", "标题", "评分项"]

    report = client.get(f"/api/scoring-runs/{run_id}/report")
    assert report.status_code == 200, report.text
    assert "材料评分报告" in report.text and "数据中心扩容方案.docx" in report.text
    assert "毕业论文" not in report.text


def test_thesis_batch_keeps_student_fields_and_export_columns(client, monkeypatch):
    semantic_calls = []
    monkeypatch.setattr(engine, "analyze_semantic_coherence", lambda parsed, scorer: semantic_calls.append(1) or [])

    rubric_id = _published_rubric(client, profile="thesis")
    batch, paper = _batch_with_paper(client, rubric_id, file_name="thesis.docx")
    assert batch["business_profile_key"] == "thesis"
    parsed = client.get(f"/api/papers/{paper['id']}").json()
    assert parsed["student_id"] == "20260001" and parsed["student_name"] == "张三"

    scored = client.post(f"/api/papers/{paper['id']}/score")
    assert scored.status_code == 200, scored.text
    with client.session_factory() as session:
        assert session.get(models.ScoringRun, scored.json()["id"]).business_profile_key == "thesis"
    assert semantic_calls == [1]

    exported = client.get(f"/api/batches/{batch['id']}/export.xlsx")
    summary = [cell.value for cell in load_workbook(BytesIO(exported.content))["总分表"][1]]
    assert summary == excel.SUMMARY_HEADERS
    report = client.get(f"/api/scoring-runs/{scored.json()['id']}/report")
    assert "毕业论文智能评分报告" in report.text and "张三（20260001）" in report.text


def test_create_batch_rejects_types_the_workbench_cannot_score(client, monkeypatch):
    rubric_id = _published_rubric(client, profile="general_document")
    # 已发布版本不可改，这里只替换“版本 → 类型”的读取，模拟一份技术方案标准。
    monkeypatch.setattr(actions, "profile_key_for_version", lambda version: "technical_proposal")
    response = client.post("/api/batches", json={"name": "x", "rubric_id": rubric_id})
    assert response.status_code == 400
    assert "technical_proposal" in response.json()["detail"]


def test_unversioned_batches_are_treated_as_thesis():
    batch = models.GradingBatch(name="旧批次", rubric_id="r")
    assert profile_key_for_batch(batch) == "thesis"
    assert batch.business_profile_key == "thesis"


# --- 评分助手 ------------------------------------------------------------------------------

def test_assistant_names_general_documents_by_title():
    papers = [
        {"paper_id": "p1", "file_name": "智慧园区建设方案.docx", "title": "智慧园区建设方案", "student_name": None},
        {"paper_id": "p2", "file_name": "b.docx", "title": "数据中心扩容方案", "student_name": None},
    ]
    assert views.paper_label(papers[1], "general_document") == "数据中心扩容方案"
    assert views.paper_label({"file_name": "a.docx", "student_name": "张三"}, "thesis") == "张三"
    match, _ = views.resolve_paper(papers, name="数据中心", profile_key="general_document")
    assert match["paper_id"] == "p2"
    match, _ = views.resolve_paper(papers, text="看看智慧园区建设方案的扣分", profile_key="general_document")
    assert match["paper_id"] == "p1"
    # 论文仍按姓名定位，不因标题里出现的词误命中。
    match, _ = views.resolve_paper(papers, name="数据中心", profile_key="thesis")
    assert match is None
    assert views.document_type_label("general_document") == "通用文档"
    assert views.document_type_label(None) == "毕业论文"
