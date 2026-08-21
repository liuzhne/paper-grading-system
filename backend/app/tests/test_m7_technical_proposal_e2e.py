from copy import deepcopy
from datetime import timedelta
from io import BytesIO
import json
from types import SimpleNamespace

from docx import Document
from sqlalchemy import select

from backend.app.db import models
from backend.app.services.cache.llm_cache import key_of
from backend.app.services.dev_user import ensure_dev_user
from backend.app.services.document_parser.extractor import extract_document
from backend.app.services.llm.mock import MockLLMScorer
from backend.app.services.rubric_import.pipeline import prepare_file_import
from backend.app.services.rubric_import.pipeline import persist_prepared_import
from backend.app.services.rubrics import lifecycle
from backend.app.services.scoring.profiles.registry import get_profile_by_key
from backend.app.services.scoring.profiles.technical_proposal import (
    TECHNICAL_PROPOSAL_PROFILE_VERSION,
)
from backend.app.services.scoring.profiles.thesis import THESIS_CHECKER_KEYS
from backend.app.services.scoring.profiles.thesis import THESIS_PROFILE_VERSION
from backend.app.tests.m2_contract_fixtures import technical_policy_snapshot_payload
from backend.app.tests.test_atomic_rule_models import _make_p1_graph
from backend.app.tests.test_m7_technical_proposal_profile import _import_command
from backend.app.tests.test_m7_technical_proposal_profile import _rules_xlsx


PROFILE_KEY = "technical_proposal"
METADATA = {
    "proposal_id": "TP-2026-001",
    "vendor_name": "Acme Solutions",
    "project_name": "智能交易平台",
}
EXPECTED_SCORES = {
    "complete": {
        "IMPLEMENTATION.IMPLEMENTATION_FEASIBILITY": 15.0,
        "IMPLEMENTATION.MILESTONE_COMPLETENESS": 10.0,
        "LENGTH": 10.0,
        "RISK_CONTROL": 15.0,
        "SOLUTION_ALIGNMENT": 20.0,
        "STRUCTURE": 10.0,
    },
    "low": {
        "IMPLEMENTATION.IMPLEMENTATION_FEASIBILITY": 7.5,
        "IMPLEMENTATION.MILESTONE_COMPLETENESS": 0.0,
        "LENGTH": 5.0,
        "RISK_CONTROL": 12.0,
        "SOLUTION_ALIGNMENT": 5.0,
        "STRUCTURE": 5.0,
    },
}


def _publish_fixture(client):
    with client.session_factory() as db:
        actor = ensure_dev_user(db)
        db.commit()
        prepared = prepare_file_import(
            command=_import_command(),
            rules_bytes=_rules_xlsx(),
        )
        identity = persist_prepared_import(
            session=db,
            prepared=prepared,
            actor_id=actor.id,
        )
        rules = db.scalars(
            select(models.AtomicRule)
            .where(
                models.AtomicRule.rubric_version_id
                == identity.rubric_version_id
            )
            .order_by(models.AtomicRule.rule_code)
        ).all()
        start = max(rule.created_at for rule in rules) + timedelta(minutes=1)
        for index, rule in enumerate(rules):
            submitted_at = start + timedelta(minutes=index * 2)
            lifecycle.submit_atomic_rule_for_review(
                db,
                identity.rubric_id,
                rule.rule_code,
                actor.id,
                "E2E fixture rule source verified.",
                now=submitted_at,
            )
            db.commit()
            lifecycle.approve_atomic_rule(
                db,
                identity.rubric_id,
                rule.rule_code,
                actor.id,
                "E2E fixture rule approved.",
                now=submitted_at + timedelta(minutes=1),
            )
            db.commit()
        lifecycle.submit_for_review(db, identity.rubric_id)
        db.commit()
        published = lifecycle.publish_rubric(
            db,
            identity.rubric_id,
            identity.compilation_id,
            actor.id,
            now=start + timedelta(hours=2),
        )
        db.commit()
        return identity.rubric_id, published.id


def _publish_thesis_isolation_fixture(client):
    with client.session_factory() as db:
        graph = _make_p1_graph(db, "m7-profile-isolation")
        graph.rubric.total_score = 100
        graph.criterion.code = "THESIS_STRUCTURE"
        graph.criterion.name = "Thesis structure"
        graph.criterion.max_score = 100
        graph.criterion.criterion_type = "deterministic"
        graph.criterion.scoring_mode = "deductive"
        graph.criterion.applies_to = "global"
        graph.version.business_profile_key = "thesis"
        graph.version.hash_scheme = "rubric-content-v2"
        graph.version.workflow_profile = "manual_json"
        graph.version.global_policy = technical_policy_snapshot_payload(
            total_score="100",
            rounding_digits=1,
        )
        reviewed_at = graph.compilation.created_at + timedelta(minutes=1)
        db.add(
            models.AtomicRule(
                rubric_version_id=graph.version.id,
                criterion_id=graph.criterion.id,
                rule_code="thesis.structure.isolation.v1",
                name="Thesis structure isolation rule",
                rule_text="The configured section must be present.",
                direction="deduct",
                effect_type="score",
                max_points=5,
                repeat_policy="once",
                cap_points=None,
                judge_type="deterministic",
                checker_key=THESIS_CHECKER_KEYS["structure"],
                checker_params={"required_sections": ["需求理解"]},
                evidence_policy={
                    "mode": "deterministic_observation",
                    "requirement": "required",
                    "minimum_coverage": "1",
                },
                positive_example="需求理解章节存在。",
                negative_example="需求理解章节缺失。",
                boundary_example=None,
                strictness="required",
                applies_to="global",
                mutex_group=None,
                depends_on_rule_codes=[],
                status="approved",
                creation_method="compiler",
                reviewed_by=graph.user.id,
                reviewed_at=reviewed_at,
            )
        )
        db.commit()
        lifecycle.submit_for_review(db, graph.rubric.id)
        db.commit()
        published = lifecycle.publish_rubric(
            db,
            graph.rubric.id,
            graph.compilation.id,
            graph.user.id,
            now=reviewed_at + timedelta(minutes=5),
        )
        db.commit()
        return published.id


def _proposal_docx(*, complete: bool) -> bytes:
    document = Document()

    def heading(text):
        document.add_heading(text, level=1)

    heading("需求理解")
    document.add_paragraph(
        "系统需要支撑高并发交易、十五分钟故障恢复和全年可用性目标。"
        if complete
        else "目标待确认。"
    )
    heading("总体方案")
    document.add_paragraph(
        "采用双可用区服务和异步消息架构，逐项对应吞吐、恢复和可用性需求。"
        if complete
        else "采用普通架构。"
    )
    heading("实施计划")
    if complete:
        table = document.add_table(rows=2, cols=4)
        for cell, value in zip(
            table.rows[0].cells,
            ("里程碑", "截止日期", "交付物", "负责人"),
        ):
            cell.text = value
        for cell, value in zip(
            table.rows[1].cells,
            ("架构评审", "2026-09-01", "评审纪要", "项目经理"),
        ):
            cell.text = value
    else:
        document.add_paragraph("待定。")
    heading("风险控制")
    risk = document.add_table(rows=2, cols=5)
    for cell, value in zip(
        risk.rows[0].cells,
        ("风险", "概率", "影响", "措施", "负责人"),
    ):
        cell.text = value
    values = (
        ("容量不足", "中", "延期", "压测并扩容", "技术经理")
        if complete
        else ("容量不足", "中", "延期", "扩容", "")
    )
    for cell, value in zip(risk.rows[1].cells, values):
        cell.text = value
    if complete:
        heading("服务承诺")
        document.add_paragraph(
            "提供 7x24 支持，重大故障十五分钟响应并持续通报。"
        )
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _create_batch(client, version_id):
    response = client.post(
        "/api/v2/evaluation-batches",
        json={
            "name": "M7 TechnicalProposal E2E",
            "rubric_version_id": version_id,
            "business_profile_key": PROFILE_KEY,
            "business_profile_version": TECHNICAL_PROPOSAL_PROFILE_VERSION,
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def _upload(client, batch_id, *, variant):
    metadata = {**METADATA, "proposal_id": "TP-2026-" + variant.upper()}
    raw = _proposal_docx(complete=variant == "complete")
    response = client.post(
        "/api/v2/submissions",
        data={
            "evaluation_batch_id": batch_id,
            "metadata_json": json.dumps(metadata, ensure_ascii=False),
        },
        files={
            "file": (
                variant + ".docx",
                raw,
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 201, response.text
    return response.json(), metadata, raw


def _upload_raw(client, batch_id, *, raw, file_name, metadata):
    response = client.post(
        "/api/v2/submissions",
        data={
            "evaluation_batch_id": batch_id,
            "metadata_json": json.dumps(metadata, ensure_ascii=False),
        },
        files={
            "file": (
                file_name,
                raw,
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def _score(client, submission_id):
    response = client.post(
        f"/api/v2/submissions/{submission_id}/score",
        json={"rescore_generation": 0},
    )
    assert response.status_code == 200, response.text
    return response.json()


def _items(run):
    return {item["criterion_code"]: item for item in run["items"]}


def _assert_exact_rule_audit(run, *, variant):
    items = _items(run)
    assert {code: item["ai_score"] for code, item in items.items()} == (
        EXPECTED_SCORES[variant]
    )
    assert {code: item["final_score"] for code, item in items.items()} == (
        EXPECTED_SCORES[variant]
    )
    assert all(len(item["rule_results"]) == 1 for item in items.values())

    expected = {
        "STRUCTURE": ("not_triggered", "0", 0),
        "LENGTH": ("not_triggered", "0", 0),
        "SOLUTION_ALIGNMENT": ("triggered", "20", 1),
        "RISK_CONTROL": ("not_triggered", "0", 0),
        "IMPLEMENTATION.MILESTONE_COMPLETENESS": (
            "not_triggered",
            "0",
            0,
        ),
        "IMPLEMENTATION.IMPLEMENTATION_FEASIBILITY": (
            "triggered",
            "15",
            1,
        ),
    }
    if variant == "low":
        expected = {
            "STRUCTURE": ("triggered", "-5", 1),
            "LENGTH": ("triggered", "-5", 1),
            "SOLUTION_ALIGNMENT": ("triggered", "5", 1),
            "RISK_CONTROL": ("triggered", "-3", 1),
            "IMPLEMENTATION.MILESTONE_COMPLETENESS": (
                "triggered",
                "-10",
                1,
            ),
            "IMPLEMENTATION.IMPLEMENTATION_FEASIBILITY": (
                "triggered",
                "7.5",
                1,
            ),
        }

    for code, (status, effect, occurrence_count) in expected.items():
        result = items[code]["rule_results"][0]
        assert result["schema_version"] == "rule-result@2"
        assert result["criterion_code"] == code
        assert result["status"] == status
        assert result["calculated_effect"] == effect
        assert len(result["occurrences"]) == occurrence_count
        assert len(result["evidence_refs"]) == occurrence_count
        for occurrence in result["occurrences"]:
            assert len(occurrence["occurrence_id"]) == 64
            assert occurrence["calculated_effect"] == effect
            assert occurrence["evidence_refs"] == result["evidence_refs"]
        for evidence in result["evidence_refs"]:
            assert evidence["evidence_type"] in {
                "deterministic_observation",
                "source_quote",
            }
            assert len(evidence["payload_hash"]) == 64

    expected_contributions = {
        code: sorted(
            (
                entry["kind"],
                None if entry["amount"] is None else float(entry["amount"]),
            )
            for entry in item["aggregation"]["contributions"]
        )
        for code, item in items.items()
    }
    assert expected_contributions == {
        "IMPLEMENTATION.IMPLEMENTATION_FEASIBILITY": [
            ("band", 15.0 if variant == "complete" else 7.5)
        ],
        "IMPLEMENTATION.MILESTONE_COMPLETENESS": (
            [("base", 10.0)]
            if variant == "complete"
            else [("base", 10.0), ("deduction", -10.0)]
        ),
        "LENGTH": (
            [("base", 10.0)]
            if variant == "complete"
            else [("base", 10.0), ("deduction", -5.0)]
        ),
        "RISK_CONTROL": (
            [("base", 15.0)]
            if variant == "complete"
            else [("base", 15.0), ("deduction", -3.0)]
        ),
        "SOLUTION_ALIGNMENT": [
            ("band", 20.0 if variant == "complete" else 5.0)
        ],
        "STRUCTURE": (
            [("base", 10.0)]
            if variant == "complete"
            else [("base", 10.0), ("deduction", -5.0)]
        ),
    }


def test_technical_proposal_complete_and_low_fixture_full_e2e(client):
    rubric_id, version_id = _publish_fixture(client)
    batch = _create_batch(client, version_id)

    complete_submission, complete_metadata, _ = _upload(
        client,
        batch["id"],
        variant="complete",
    )
    low_submission, low_metadata, _ = _upload(
        client,
        batch["id"],
        variant="low",
    )

    assert batch["rubric_id"] == rubric_id
    assert batch["rubric_version_id"] == version_id
    assert complete_submission["metadata"] == complete_metadata
    assert low_submission["metadata"] == low_metadata
    complete = _score(client, complete_submission["id"])
    low = _score(client, low_submission["id"])

    for submission, expected_sections in (
        (complete_submission, 5),
        (low_submission, 4),
    ):
        snapshot = client.get(
            f"/api/v2/submissions/{submission['id']}/document-snapshot"
        )
        assert snapshot.status_code == 200, snapshot.text
        assert snapshot.json()["schema_version"] == "document-snapshot@1"
        assert snapshot.json()["business_profile_key"] == PROFILE_KEY
        assert snapshot.json()["section_count"] == expected_sections
        assert snapshot.json()["evidence_unit_count"] > expected_sections

    assert (
        complete["ai_total_score"],
        complete["final_total_score"],
        complete["grade"],
        complete["need_manual_review"],
    ) == (80.0, 80.0, "Gold", False)
    assert (
        low["ai_total_score"],
        low["final_total_score"],
        low["grade"],
        low["need_manual_review"],
    ) == (34.5, 34.5, "Rework", True)
    _assert_exact_rule_audit(complete, variant="complete")
    _assert_exact_rule_audit(low, variant="low")
    with client.session_factory() as db:
        versions = db.scalars(
            select(models.RubricVersion).where(
                models.RubricVersion.rubric_id == rubric_id
            )
        ).all()
        assert len(versions) == 1
        assert versions[0].id == version_id
        rubric = db.get(models.Rubric, rubric_id)
        compilation = db.get(
            models.RubricCompilation,
            versions[0].compilation_id,
        )
        assert rubric.status == "published"
        assert rubric.published_at is not None
        assert compilation.status == "validated"
        assert compilation.published_at == rubric.published_at
        assert compilation.final_version_hash == versions[0].version_hash
        for run_id in (complete["id"], low["id"]):
            stored = db.get(models.ScoringRun, run_id)
            assert stored.plan_schema_version == "rule-execution-plan@2"
            assert stored.execution_plan_snapshot["rubric_version_id"] == version_id
            assert set(stored.checker_manifest) == {
                "core.required_sections.v1",
                "core.text_length_range.v1",
                "generic.hybrid.required.v1",
            }

    reviewed = client.post(
        f"/api/v2/scoring-runs/{low['id']}/review",
        json={"reason": "Low fixture policy review approved."},
    )
    assert reviewed.status_code == 200, reviewed.text
    assert reviewed.json()["status"] == "reviewed"
    assert reviewed.json()["need_manual_review"] is False
    logs = client.get(f"/api/v2/scoring-runs/{low['id']}/review-logs")
    assert logs.status_code == 200, logs.text
    assert [(item["score_item_id"], item["resolution_type"]) for item in logs.json()] == [
        (None, "ordinary_override")
    ]

    exported_response = client.get(
        f"/api/v2/scoring-runs/{low['id']}/export.json"
    )
    assert exported_response.status_code == 200, exported_response.text
    exported = exported_response.json()
    assert exported["schema"] == "grading-core/run-export@2"
    assert exported["run"]["final_total_score"] == 34.5
    assert exported["run"]["grade"] == "Rework"
    assert exported["profile_extensions"]["metadata"] == low_metadata
    assert exported["profile_extensions"]["findings"] == {
        "risk_items": [
            {
                "item_ordinal": 0,
                "fields_present": [
                    "impact",
                    "mitigation",
                    "probability",
                ],
                "source_text": "容量不足 | 中 | 延期 | 扩容",
            }
        ],
        "milestones": [],
        "milestone_complete": False,
    }
    serialized = json.dumps(exported, ensure_ascii=False).casefold()
    for forbidden in (
        "student_id",
        "student_name",
        "advisor",
        "coherence_findings",
        "format_findings",
        '"thesis"',
    ):
        assert forbidden not in serialized


def test_same_content_profile_identity_forces_cache_miss_and_no_scope_leakage(
    tmp_path,
):
    raw = _proposal_docx(complete=True)
    source = tmp_path / "same-content.docx"
    source.write_bytes(raw)
    extracted = extract_document(source)
    scorer = MockLLMScorer()
    submission = SimpleNamespace(
        id="same-submission",
        source_artifact_hash="a" * 64,
        submission_metadata=deepcopy(METADATA),
    )
    proposal = get_profile_by_key(PROFILE_KEY)
    thesis = get_profile_by_key("thesis")

    proposal_document = proposal.interpret_document(
        extracted_document=extracted,
        submission=submission,
    ).to_mapping()
    thesis_document = thesis.interpret_document(
        extracted_document=extracted,
        submission=submission,
    )
    if callable(getattr(thesis_document, "to_mapping", None)):
        thesis_document = thesis_document.to_mapping()

    assert proposal.profile_version == TECHNICAL_PROPOSAL_PROFILE_VERSION
    assert thesis.profile_version == THESIS_PROFILE_VERSION
    assert proposal_document["profile_key"] == PROFILE_KEY
    assert thesis_document["profile_key"] == "thesis"
    assert proposal_document["content_hash"] != thesis_document["content_hash"]
    assert (
        proposal_document["document_snapshot_hash"]
        != thesis_document["document_snapshot_hash"]
    )

    proposal_manifest = proposal.build_checker_registry().manifest(
        profile_key=PROFILE_KEY,
        document_schema_version="document-snapshot@1",
    )
    thesis_manifest = thesis.build_checker_registry().manifest(
        profile_key="thesis",
        document_schema_version="document-snapshot@1",
    )
    assert set(proposal_manifest).isdisjoint(thesis_manifest)
    assert all(
        PROFILE_KEY in item["supported_profiles"]
        for item in proposal_manifest.values()
    )
    assert all(
        "thesis" in item["supported_profiles"]
        for item in thesis_manifest.values()
    )

    submission_snapshot = {
        "metadata": deepcopy(METADATA),
        "profile_key": PROFILE_KEY,
    }
    proposal_extensions = proposal.build_prompt_extensions(
        submission_snapshot=submission_snapshot,
        document_snapshot=proposal_document,
    )
    thesis_extensions = thesis.build_prompt_extensions(
        submission_snapshot={**submission_snapshot, "profile_key": "thesis"},
        document_snapshot=thesis_document,
    )
    assert proposal_extensions["metadata"] == METADATA
    assert thesis_extensions["metadata"] == {}
    assert "thesis" not in proposal_extensions["profile_extensions"]
    assert PROFILE_KEY not in thesis_extensions["profile_extensions"]

    proposal_cache_key = key_of(
        {
            "prompt_version": proposal.prompt_version,
            "runtime_identity": proposal.build_runtime_identity(scorer),
            "document_snapshot_hash": proposal_document[
                "document_snapshot_hash"
            ],
            "profile_prompt_extensions": proposal_extensions,
        }
    )
    thesis_cache_key = key_of(
        {
            "prompt_version": thesis.prompt_version,
            "runtime_identity": thesis.build_runtime_identity(scorer),
            "document_snapshot_hash": thesis_document[
                "document_snapshot_hash"
            ],
            "profile_prompt_extensions": thesis_extensions,
        }
    )
    assert proposal_cache_key != thesis_cache_key
    assert len(proposal_cache_key) == len(thesis_cache_key) == 64


def test_same_docx_scores_through_both_production_profiles_without_cache_reuse(
    client,
):
    _rubric_id, proposal_version_id = _publish_fixture(client)
    thesis_version_id = _publish_thesis_isolation_fixture(client)
    proposal_batch = _create_batch(client, proposal_version_id)
    thesis_batch_response = client.post(
        "/api/v2/evaluation-batches",
        json={
            "name": "M7 Thesis isolation E2E",
            "rubric_version_id": thesis_version_id,
            "business_profile_key": "thesis",
            "business_profile_version": THESIS_PROFILE_VERSION,
        },
    )
    assert thesis_batch_response.status_code == 201, thesis_batch_response.text
    raw = _proposal_docx(complete=True)
    proposal_submission = _upload_raw(
        client,
        proposal_batch["id"],
        raw=raw,
        file_name="same-content.docx",
        metadata=METADATA,
    )
    thesis_submission = _upload_raw(
        client,
        thesis_batch_response.json()["id"],
        raw=raw,
        file_name="same-content.docx",
        metadata=METADATA,
    )

    proposal_run = _score(client, proposal_submission["id"])
    thesis_run = _score(client, thesis_submission["id"])

    assert proposal_submission["source_artifact_hash"] == thesis_submission[
        "source_artifact_hash"
    ]
    assert proposal_run["source_artifact_hash"] == thesis_run[
        "source_artifact_hash"
    ]
    assert proposal_run["normalized_content_hash"] != thesis_run[
        "normalized_content_hash"
    ]
    assert proposal_run["document_snapshot_hash"] != thesis_run[
        "document_snapshot_hash"
    ]
    assert proposal_run["idempotency_key"] != thesis_run["idempotency_key"]
    assert proposal_run["execution_plan_hash"] != thesis_run[
        "execution_plan_hash"
    ]
    assert proposal_run["prompt_version"] != thesis_run["prompt_version"]
    assert proposal_run["runtime_identity"]["profile_key"] == PROFILE_KEY
    assert thesis_run["runtime_identity"]["profile_key"] == "thesis"
    assert proposal_run["ai_total_score"] == 80.0
    assert thesis_run["ai_total_score"] == 100.0

    proposal_export = client.get(
        f"/api/v2/scoring-runs/{proposal_run['id']}/export.json"
    ).json()
    thesis_export = client.get(
        f"/api/v2/scoring-runs/{thesis_run['id']}/export.json"
    ).json()
    assert proposal_export["profile_extensions"]["metadata"] == METADATA
    assert thesis_export["profile_extensions"]["metadata"] == {}
    assert set(proposal_export["profile_extensions"]["findings"]) == {
        "risk_items",
        "milestones",
        "milestone_complete",
    }
    assert set(thesis_export["profile_extensions"]["findings"]) == {
        "coherence",
        "format",
        "references",
    }
    assert "student_id" not in json.dumps(
        proposal_export,
        ensure_ascii=False,
    )
