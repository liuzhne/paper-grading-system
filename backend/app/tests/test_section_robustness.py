from docx import Document

from backend.app.services.document_parser.parser import parse_document


def _save(doc, tmp_path, name="paper.docx"):
    path = tmp_path / name
    doc.save(str(path))
    return str(path)


def test_style_headings_detected_when_regex_would_miss(tmp_path):
    doc = Document()
    # 标题文本既不含"第X章"也不在已知章节名里，仅靠样式才能识别。
    doc.add_heading("研究背景与意义", level=1)
    doc.add_paragraph("本文围绕教学质量评价展开，说明研究背景与价值。" * 3)
    doc.add_heading("模型构建与求解", level=1)
    doc.add_paragraph("给出建模过程、求解方法与实验设置。" * 3)
    parsed = parse_document(_save(doc, tmp_path))

    titles = [section.title for section in parsed.sections]
    assert any("研究背景" in title for title in titles)
    assert any("模型构建" in title for title in titles)
    # 有样式标题 → 置信度不应过低。
    assert parsed.structure_confidence >= 0.5


def test_low_structure_confidence_is_flagged(tmp_path):
    doc = Document()
    doc.add_paragraph("这是一段没有任何标题、编号或预期章节的连续正文。" * 50)
    parsed = parse_document(_save(doc, tmp_path))

    assert parsed.structure_confidence < 0.5
    checks = {check.code: check for check in parsed.structure_checks}
    assert "SECTION_DETECTION" in checks
    assert checks["SECTION_DETECTION"].passed is False
    assert "人工确认" in (checks["SECTION_DETECTION"].location or "")
