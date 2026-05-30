from docx import Document

from backend.app.services.coherence import analyze_coherence
from backend.app.services.document_parser.parser import parse_document


def _kinds(findings):
    return {finding["kind"] for finding in findings}


def test_figure_referenced_but_not_declared():
    findings = analyze_coherence("实验结果见图3，说明方法有效。", [])
    missing = [f for f in findings if f["kind"] == "figure_ref_missing"]
    assert missing and missing[0]["refs"] == [3]


def test_figure_declared_but_not_referenced():
    findings = analyze_coherence("图1 系统架构\n本文介绍研究方法与流程。", [])
    unref = [f for f in findings if f["kind"] == "figure_unreferenced"]
    assert unref and unref[0]["refs"] == [1]


def test_figure_numbering_gap():
    text = "图1 架构\n图3 流程\n如图1所示，又如图3所示。"
    findings = analyze_coherence(text, [])
    gaps = [f for f in findings if f["kind"] == "figure_numbering_gap"]
    assert gaps and gaps[0]["refs"] == [2]


def test_citation_missing_entry_and_orphan():
    findings = analyze_coherence("研究[2]表明结论成立。", ["[1] 张三. 研究. 2024."])
    kinds = _kinds(findings)
    assert "citation_missing_entry" in kinds  # [2] 无条目
    assert "reference_uncited" in kinds  # [1] 未被引用


def test_citation_consistent_has_no_warnings():
    findings = analyze_coherence("方法见[1]，结果见[2]。", ["[1] A. 2023.", "[2] B. 2024."])
    warnings = [f for f in findings if f["severity"] == "warning"]
    assert warnings == []


def test_reference_numbering_gap():
    findings = analyze_coherence("见[1]与[3]。", ["[1] A.", "[3] C."])
    gaps = [f for f in findings if f["kind"] == "reference_numbering_gap"]
    assert gaps and gaps[0]["refs"] == [2]


def test_author_year_flagged_as_info():
    findings = analyze_coherence("已有研究（张三, 2020）指出……", [])
    author_year = [f for f in findings if f["kind"] == "citation_author_year"]
    assert author_year and author_year[0]["severity"] == "info"


def test_parse_document_populates_coherence_findings(tmp_path):
    doc = Document()
    doc.add_heading("第一章 绪论", level=1)
    doc.add_paragraph("实验结果见图3，说明方法有效。")  # 引用图3但无题注
    doc.add_paragraph("研究[2]表明结论成立。")  # 引用[2]但条目只有[1]
    doc.add_heading("参考文献", level=1)
    doc.add_paragraph("[1] 张三. 研究. 2024.")
    path = tmp_path / "coherence.docx"
    doc.save(str(path))

    parsed = parse_document(str(path))
    kinds = {finding["kind"] for finding in parsed.coherence_findings}
    assert "figure_ref_missing" in kinds
    assert "citation_missing_entry" in kinds
