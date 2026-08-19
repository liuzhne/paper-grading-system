from copy import deepcopy
from datetime import timedelta
from io import BytesIO
import json

from docx import Document
from sqlalchemy import select

from backend.app.db import models
from backend.app.services.rubrics import lifecycle
from backend.app.services.scoring.profiles.registry import temporary_profile_registration
from backend.app.tests.m2_contract_fixtures import (
    PROFILE_KEY,
    PROFILE_VERSION,
    runtime_identity_payload,
    technical_document_payload,
    technical_policy_snapshot_payload,
    technical_submission_payload,
)
from backend.app.tests.m3_contract_fixtures import (
    CHECKER_KEY,
    DETERMINISTIC_RULE_CODE,
    SEMANTIC_RULE_CODE,
)
from backend.app.tests.test_atomic_rule_models import _make_p1_graph
from backend.app.tests.test_m3_loader_registry_plan import _TestProfile
from backend.app.tests.test_m3_loader_registry_plan import _registry
from backend.app.tests.test_m3_score_submission import _CapturingRuntime


class _ApiTechnicalProposalProfile(_TestProfile):
    def build_checker_registry(self):
        return _registry()

    def build_llm_runtime(self, _scorer):
        return _CapturingRuntime()

    def build_runtime_identity(self, _scorer):
        return runtime_identity_payload()

    def interpret_document(self, *, extracted_document, submission):
        assert extracted_document.schema_version == "extracted-document@1"
        assert submission.source_artifact_hash
        return technical_document_payload(profile_version=self.profile_version)

    def build_export_extensions(self, *, submission, document_snapshot, run):
        del document_snapshot, run
        metadata = submission.submission_metadata
        return {
            "metadata": {
                "project_name": metadata["project_name"],
            },
            "findings": [],
        }


def _published_technical_version(client, suffix):
    with client.session_factory() as db:
        graph = _make_p1_graph(db, suffix)
        graph.rubric.total_score = 100
        graph.criterion.code = "SOLUTION_FIT"
        graph.criterion.name = "Solution fit"
        graph.criterion.max_score = 80
        graph.criterion.criterion_type = "llm_judgment"
        graph.criterion.scoring_mode = "banded"
        graph.criterion.applies_to = "requirements_understanding"
        graph.criterion.display_order = 1
        risk_criterion = models.RubricCriterion(
            rubric_id=graph.rubric.id,
            code="RISK_CONTROL",
            name="Risk control completeness",
            max_score=20,
            criterion_type="deterministic",
            scoring_mode="deductive",
            applies_to="risk_control",
            dimension="content",
            display_order=0,
        )
        db.add(risk_criterion)
        db.flush()
        graph.version.business_profile_key = PROFILE_KEY
        graph.version.hash_scheme = "rubric-content-v2"
        graph.version.global_policy = technical_policy_snapshot_payload(
            total_score="100",
            rounding_digits=2,
        )

        reviewed_at = graph.compilation.created_at + timedelta(minutes=1)
        deterministic_rule = models.AtomicRule(
            rubric_version_id=graph.version.id,
            criterion_id=risk_criterion.id,
            rule_code=DETERMINISTIC_RULE_CODE,
            name="Risk owner required",
            rule_text="Risk control must identify an accountable owner.",
            direction="deduct",
            effect_type="score",
            max_points=10,
            repeat_policy="once",
            cap_points=None,
            judge_type="deterministic",
            checker_key=CHECKER_KEY,
            checker_params={"required_fields": ["owner"]},
            evidence_policy={
                "mode": "scoped_absence",
                "requirement": "required",
                "minimum_coverage": "1",
            },
            positive_example="Each risk names an owner.",
            negative_example="No accountable owner is named.",
            boundary_example=None,
            strictness="required",
            applies_to="risk_control",
            mutex_group=None,
            depends_on_rule_codes=[],
            status="approved",
            creation_method="compiler",
            reviewed_by=graph.user.id,
            reviewed_at=reviewed_at,
        )
        semantic_rule = models.AtomicRule(
            rubric_version_id=graph.version.id,
            criterion_id=graph.criterion.id,
            rule_code=SEMANTIC_RULE_CODE,
            name="Solution fit",
            rule_text="Requirements and solution must be explicitly aligned.",
            direction="band",
            effect_type="score",
            max_points=None,
            repeat_policy=None,
            cap_points=None,
            judge_type="semantic",
            checker_key=None,
            checker_params={},
            evidence_policy={
                "mode": "source_quote",
                "requirement": "required",
                "minimum_coverage": "1",
            },
            positive_example="The mapping is explicit.",
            negative_example="Requirements are only listed.",
            boundary_example=None,
            strictness="required",
            applies_to="requirements_understanding",
            mutex_group=None,
            depends_on_rule_codes=[],
            status="approved",
            creation_method="compiler",
            reviewed_by=graph.user.id,
            reviewed_at=reviewed_at,
        )
        db.add_all([deterministic_rule, semantic_rule])
        db.flush()
        db.add_all(
            [
                models.RuleLevel(
                    atomic_rule_id=semantic_rule.id,
                    level_code="FIT_HIGH",
                    points=80,
                    descriptor="Requirements and solution are explicitly aligned.",
                    display_order=0,
                ),
                models.RuleLevel(
                    atomic_rule_id=semantic_rule.id,
                    level_code="FIT_LOW",
                    points=40,
                    descriptor="The solution only partially addresses requirements.",
                    display_order=1,
                ),
            ]
        )
        db.commit()
        lifecycle.submit_for_review(db, graph.rubric.id)
        db.commit()
        published = lifecycle.publish_rubric(
            db,
            graph.rubric.id,
            graph.compilation.id,
            graph.user.id,
            now=reviewed_at + timedelta(minutes=5),
        )
        db.commit()
        return graph.rubric.id, published.id


def _proposal_docx():
    document = Document()
    document.add_heading("需求理解", level=1)
    document.add_paragraph("本方案覆盖吞吐、时延与可用性目标。")
    document.add_heading("总体方案", level=1)
    document.add_paragraph("风险负责人为项目经理，升级窗口为十五分钟。")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _create_batch_and_submission(client, suffix):
    rubric_id, version_id = _published_technical_version(client, suffix)
    batch = client.post(
        "/api/v2/evaluation-batches",
        json={
            "name": "Technical proposal API batch " + suffix,
            "rubric_version_id": version_id,
            "business_profile_key": PROFILE_KEY,
            "business_profile_version": PROFILE_VERSION,
        },
    )
    assert batch.status_code == 201, batch.text
    assert batch.json()["rubric_id"] == rubric_id
    assert batch.json()["status"] == "active"

    upload = client.post(
        "/api/v2/submissions",
        data={
            "evaluation_batch_id": batch.json()["id"],
            "metadata_json": json.dumps(
                technical_submission_payload()["metadata"],
                ensure_ascii=False,
            ),
        },
        files={
            "file": (
                "proposal.docx",
                _proposal_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload.status_code == 201, upload.text
    assert upload.json()["status"] == "parsed"
    assert upload.json()["business_profile_key"] == PROFILE_KEY
    return batch.json(), upload.json()


def _score(client, submission_id, generation=0):
    response = client.post(
        f"/api/v2/submissions/{submission_id}/score",
        json={"rescore_generation": generation},
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_v2_technical_submission_upload_score_query_review_and_rescore(client):
    with temporary_profile_registration(_ApiTechnicalProposalProfile()):
        _batch, submission = _create_batch_and_submission(client, "flow")

        summary = client.get(
            f"/api/v2/submissions/{submission['id']}/document-snapshot"
        )
        assert summary.status_code == 200, summary.text
        snapshot = summary.json()
        assert snapshot["schema_version"] == "document-snapshot@1"
        assert snapshot["business_profile_key"] == PROFILE_KEY
        assert snapshot["section_count"] == 2
        assert snapshot["evidence_unit_count"] == 3
        assert "snapshot_payload" not in snapshot
        assert "full_text" not in snapshot

        run = _score(client, submission["id"])
        duplicate = _score(client, submission["id"])
        assert duplicate["id"] == run["id"]
        assert run["schema_version"] == "run-read@2"
        assert run["paper_id"] is None
        assert run["submission_id"] == submission["id"]
        assert run["business_profile_key"] == PROFILE_KEY
        assert run["rescore_generation"] == 0
        assert {item["criterion_code"] for item in run["items"]} == {
            "RISK_CONTROL",
            "SOLUTION_FIT",
        }
        assert all(item["rule_results"] for item in run["items"])

        risk = next(
            item for item in run["items"] if item["criterion_code"] == "RISK_CONTROL"
        )
        automatic_facts = {
            key: deepcopy(risk[key])
            for key in ("ai_score", "auto_score_status", "rule_results", "evidence")
        }
        override = client.patch(
            f"/api/v2/score-items/{risk['id']}",
            json={
                "final_score": 9,
                "reason": "Reviewer confirmed a partial exception.",
                "resolution_type": "ordinary_override",
            },
        )
        assert override.status_code == 200, override.text
        assert override.json()["final_score"] == 9
        for key, value in automatic_facts.items():
            assert override.json()[key] == value

        reviewed = client.post(
            f"/api/v2/scoring-runs/{run['id']}/review",
            json={"reason": "Technical proposal review complete."},
        )
        assert reviewed.status_code == 200, reviewed.text
        assert reviewed.json()["status"] == "reviewed"

        logs = client.get(f"/api/v2/scoring-runs/{run['id']}/review-logs")
        assert logs.status_code == 200
        assert [item["resolution_type"] for item in logs.json()] == [
            "ordinary_override",
            "ordinary_override",
        ]

        rescored = _score(client, submission["id"], generation=1)
        replayed = _score(client, submission["id"], generation=1)
        assert rescored["id"] == replayed["id"]
        assert rescored["id"] != run["id"]
        assert rescored["rescore_generation"] == 1


def test_v2_resolution_preserves_blocked_auto_facts_and_gates_review(client):
    with temporary_profile_registration(_ApiTechnicalProposalProfile()):
        _batch, submission = _create_batch_and_submission(client, "resolution")
        run = _score(client, submission["id"])
        with client.session_factory() as db:
            item = db.scalar(
                select(models.ScoreItem).where(
                    models.ScoreItem.scoring_run_id == run["id"],
                    models.ScoreItem.auto_score_status == "calculated",
                )
            )
            item.auto_score_status = "blocked"
            item.ai_score = None
            item.final_score = None
            item.need_manual_review = True
            item.scoring_run.ai_total_score = None
            item.scoring_run.final_total_score = None
            item.scoring_run.grade = None
            item.scoring_run.need_manual_review = True
            item.scoring_run.submission.status = "pending_review"
            immutable_rule_results = deepcopy(item.rule_results)
            blocked_item_id = item.id
            db.commit()

        unresolved = client.post(
            f"/api/v2/scoring-runs/{run['id']}/review",
            json={"reason": "Should remain blocked."},
        )
        assert unresolved.status_code == 409
        assert "unresolved" in unresolved.json()["detail"]

        ordinary = client.patch(
            f"/api/v2/score-items/{blocked_item_id}",
            json={
                "final_score": 10,
                "reason": "Wrong capability.",
                "resolution_type": "ordinary_override",
            },
        )
        assert ordinary.status_code == 409

        wrong_resolution = client.patch(
            f"/api/v2/score-items/{blocked_item_id}",
            json={
                "final_score": 10,
                "reason": "Wrong validation capability.",
                "resolution_type": "resolve_validation",
            },
        )
        assert wrong_resolution.status_code == 409

        resolved = client.patch(
            f"/api/v2/score-items/{blocked_item_id}",
            json={
                "final_score": 10,
                "reason": "Authorized reviewer resolved the execution block.",
                "resolution_type": "resolve_block",
            },
        )
        assert resolved.status_code == 200, resolved.text
        assert resolved.json()["auto_score_status"] == "blocked"
        assert resolved.json()["ai_score"] is None
        assert resolved.json()["final_score"] == 10
        assert resolved.json()["rule_results"] == immutable_rule_results

        reviewed = client.post(
            f"/api/v2/scoring-runs/{run['id']}/review",
            json={"reason": "Block resolved with an auditable human score."},
        )
        assert reviewed.status_code == 200, reviewed.text
        assert reviewed.json()["status"] == "reviewed"
        assert reviewed.json()["final_total_score"] is not None

        fetched_submission = client.get(
            f"/api/v2/submissions/{submission['id']}"
        )
        assert fetched_submission.status_code == 200
        assert fetched_submission.json()["status"] == "reviewed"


def test_v2_submission_errors_are_explicit_and_do_not_fall_back_to_thesis(client):
    missing_batch = client.post(
        "/api/v2/submissions",
        data={"evaluation_batch_id": "missing", "metadata_json": "{}"},
        files={"file": ("proposal.docx", _proposal_docx())},
    )
    assert missing_batch.status_code == 404

    with temporary_profile_registration(_ApiTechnicalProposalProfile()):
        _rubric_id, version_id = _published_technical_version(client, "errors")
        wrong_version = client.post(
            "/api/v2/evaluation-batches",
            json={
                "name": "Wrong profile version",
                "rubric_version_id": version_id,
                "business_profile_key": PROFILE_KEY,
                "business_profile_version": "future-profile@99",
            },
        )
        assert wrong_version.status_code == 409
        assert "profile version" in wrong_version.json()["detail"]
