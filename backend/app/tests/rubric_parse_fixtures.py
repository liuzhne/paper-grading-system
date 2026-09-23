"""评分规则解析重构用的合成样本（无真实 PII）。

每个样本都刻意包含一种现有解析器会静默丢弃或解析错误的形态，
供快照基线、单元台账与覆盖率测试共用。
"""

from io import BytesIO

from docx import Document
from openpyxl import Workbook


def _save(workbook):
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def complex_rules_xlsx():
    """标题行 + 表头 + 合并维度列 + 未映射列 + 多判断单元格 + 表尾全局规则 + 合计 + 第二张规则表。"""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "评分表"
    sheet.append(["XX大学本科毕业论文评分表"])  # R1 标题
    sheet.merge_cells("A1:F1")
    sheet.append([])  # R2 空行
    sheet.append(["评价维度", "评分项", "分值", "评分说明", "扣分规则", "备注"])  # R3 表头
    sheet.append(["选题与综述", "选题意义", 10, "选题具有理论或实践价值。", "意义笼统扣1-3分", None])  # R4
    sheet.append([None, "文献综述", 10, "综述覆盖充分。", "文献少于20篇扣2分", "需结合答辩"])  # R5
    sheet.merge_cells("A4:A5")
    sheet.append(["研究方法", "方法设计", 30, "方法合理。", "格式错误、图表不清、引用不规范各扣2分", None])  # R6
    sheet.append(["研究方法", "写作规范", None, "语言通顺", "表述要严谨", None])  # R7 无满分，被丢弃
    sheet.append(["注：全文错别字每处扣0.5分，最多扣5分，从总分中扣除。"])  # R8 全局规则
    sheet.merge_cells("A8:F8")
    sheet.append(["合计", None, 50])  # R9
    extra = workbook.create_sheet("附加扣分")
    extra.append(["评分项", "分值", "扣分规则"])
    extra.append(["学术诚信", 0, "查重率超过30%扣10分"])
    return _save(workbook)


def simple_rules_xlsx():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "评分规则"
    sheet.append(["编号", "评分项", "分值", "评分说明", "证据提示", "扣分规则"])
    sheet.append(["C01", "研究方法", 20, "方法合理，数据来源清楚。", "研究方法；实验设计", "方法说明不足扣3分"])
    sheet.append(["C02", "文献综述", 15, "综述覆盖充分。", "文献综述", "文献覆盖不足扣分"])
    return _save(workbook)


def merged_parent_dimension_xlsx(*, with_item_labels=True, independently_repeated=False):
    """A real-world layout where one parent category spans several scored rows."""

    workbook = Workbook()
    sheet = workbook.active
    headers = ["打分项", "评价项目", "具体要求"] if with_item_labels else ["编号", "评价项目", "分值", "具体要求"]
    sheet.append(headers)
    rows = [
        ("指导教师成绩项2（20分）", "分析与解决问题", "能够检索并分析相关研究现状。"),
        ("指导教师成绩项3（20分）", "分析与解决问题", "能够运用工程方法形成解决方案。"),
        ("指导教师成绩项4（10分）", "分析与解决问题", "能够使用开发与测试工具。"),
        ("指导教师成绩项5（10分）", "分析与解决问题", "能够分析系统运行与可持续性。"),
    ]
    for index, (label, parent, description) in enumerate(rows, start=2):
        if with_item_labels:
            sheet.append([label, parent, description])
        else:
            score = 20 if index < 4 else 10
            sheet.append([f"T{index:02d}", parent, score, description])
    parent_column = "B"
    if not independently_repeated:
        sheet.merge_cells(f"{parent_column}2:{parent_column}5")
    return _save(workbook)


def unmapped_header_xlsx():
    """列名全部不在别名表中：触发 E1。"""

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["考核内容", "配分", "细则"])
    sheet.append(["选题", 10, "选题新颖"])
    return _save(workbook)


def template_docx_with_comments():
    document = Document()
    document.add_heading("本科毕业论文模板", level=1)
    document.add_heading("第一章 绪论", level=2)
    document.add_paragraph("说明研究背景、研究意义和论文结构。")
    document.add_heading("第三章 研究方法", level=2)
    paragraph = document.add_paragraph("")
    run = paragraph.add_run("数据来源描述")
    document.add_comment(runs=run, text="此处需说明数据来源，缺失扣3分", author="导师")
    document.add_paragraph("正文不少于800字，否则扣2分。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "图表"
    table.cell(0, 1).text = "须有题注"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def rules_docx():
    """规则写在 Word 正文与表格里（仅 Word 导入）。"""

    document = Document()
    document.add_heading("课程报告评分标准", level=1)
    document.add_paragraph("总分100分，各评分项如下。")
    table = document.add_table(rows=4, cols=4)
    for row, values in enumerate(
        [
            ("编号", "评分项", "分值", "评分说明"),
            ("K01", "问题分析", "40", "问题界定清楚，分析有依据。"),
            ("K02", "方案设计", "40", "方案可行，论证充分。"),
            ("K03", "报告规范", "20", "格式规范，引用完整。"),
        ]
    ):
        for col, value in enumerate(values):
            table.cell(row, col).text = value
    document.add_paragraph("迟交一天扣5分。")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def rules_docx_paragraphs():
    """没有表格，评分项写在段落里。"""

    document = Document()
    document.add_heading("课程设计评分标准", level=1)
    document.add_paragraph("一、选题意义（10分）：选题具有实际价值。")
    document.add_paragraph("二、方案设计（60分）：方案完整可行。")
    document.add_paragraph("三、报告规范（30分）：格式规范。")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def template_docx_with_rule_table():
    """与 simple_rules_xlsx 同时上传：C01 一致、C02 分值冲突、C03 为 Word 独有。"""

    document = Document()
    document.add_heading("评分说明", level=1)
    table = document.add_table(rows=4, cols=3)
    for row, values in enumerate(
        [("编号", "评分项", "分值"), ("C01", "研究方法", "20"), ("C02", "文献综述", "10"), ("C03", "创新性", "5")]
    ):
        for col, value in enumerate(values):
            table.cell(row, col).text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
