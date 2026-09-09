from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy import select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.engine import make_url
from sqlalchemy.pool import StaticPool

from backend.app.core.config import settings
from backend.app.db import models
from backend.app.db.models import Base
from backend.app.db.session import get_db
from backend.app.main import app
from backend.app.schemas.rubric import RubricCriterionCreate
from backend.app.services.dev_user import ensure_dev_user
from backend.app.services.rubric_import.persist import build_criterion
from backend.app.services.storage.local import ensure_storage_dirs


#: Hosts a test database may legitimately live on. CI uses 127.0.0.1; local
#: development uses SQLite files. Anything else is someone's real deployment.
_LOCAL_DB_HOSTS = frozenset({"", "localhost", "127.0.0.1", "::1"})


@pytest.fixture(autouse=True)
def tests_bring_their_own_byok_master_key(monkeypatch):
    """测试自带 BYOK 主密钥，不借用环境里的那把。

    `.env.local` 在开发机上有 `BYOK_MASTER_KEY`，于是加密相关的用例在本地一路绿，
    **只有在没有那个文件的机器（CI）上才暴露**。这与「测试打到生产库」是同一个根：
    本地通过是因为悄悄用上了环境里的东西。

    这里给一把固定的测试密钥，让本地与 CI 跑在同一条件下。需要验证「没有密钥时
    应当失败」的用例自己 monkeypatch 成空值。
    """
    from backend.app.core.config import settings

    monkeypatch.setattr(settings, "BYOK_MASTER_KEY", "test-only-master-key")


@pytest.fixture(autouse=True)
def never_let_tests_reach_a_remote_database(tmp_path, monkeypatch):
    """Neutralise an ambient DATABASE_URL that points at a real deployment.

    `alembic/env.py` unconditionally overrides `sqlalchemy.url` with
    `settings.DATABASE_URL`, and `settings` loads `.env.local` — which in this
    repository holds production Supabase credentials. A migration test that
    forgets to monkeypatch `settings.DATABASE_URL` therefore runs DDL against
    production. Transactional DDL and the least-privilege `pgs_app` role both
    happen to block it today, but neither is a guarantee worth relying on.

    Redirecting rather than failing keeps the whole suite runnable on a
    developer machine that has production credentials in `.env.local`. Tests
    needing a specific database still monkeypatch it in the test body, which
    runs after this fixture.
    """

    url = make_url(settings.DATABASE_URL)
    if url.get_backend_name() == "sqlite" or (url.host or "") in _LOCAL_DB_HOSTS:
        return
    monkeypatch.setattr(
        settings, "DATABASE_URL", "sqlite+pysqlite:///%s" % (tmp_path / "neutralised.db")
    )


@pytest.fixture(autouse=True)
def isolate_tests_from_deployment_environment(monkeypatch):
    """Keep the documented test profile independent from local production env."""

    defaults = {
        "AUTH_ENABLED": False,
        "AUTH_PASSWORD": None,
        "STORAGE_PROVIDER": "local",
        "LLM_PROVIDER": "mock",
        "LLM_FALLBACK_TO_MOCK": True,
        "SHEET_WRITER_PROVIDER": "mock",
        "SHEET_FALLBACK_TO_MOCK": True,
    }
    for name, value in defaults.items():
        monkeypatch.setattr(settings, name, value)


@pytest.fixture()
def client(tmp_path):
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    TestingSessionLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False)
    Base.metadata.create_all(engine)

    original_storage_root = settings.STORAGE_ROOT
    original_runtime_settings = {
        "LLM_PROVIDER": settings.LLM_PROVIDER,
        "LLM_FALLBACK_TO_MOCK": settings.LLM_FALLBACK_TO_MOCK,
        "SHEET_WRITER_PROVIDER": settings.SHEET_WRITER_PROVIDER,
        "SHEET_FALLBACK_TO_MOCK": settings.SHEET_FALLBACK_TO_MOCK,
    }
    settings.STORAGE_ROOT = tmp_path / "storage"
    settings.LLM_PROVIDER = "mock"
    settings.LLM_FALLBACK_TO_MOCK = True
    settings.SHEET_WRITER_PROVIDER = "mock"
    settings.SHEET_FALLBACK_TO_MOCK = True
    ensure_storage_dirs()

    def override_get_db():
        db = TestingSessionLocal()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        # Lifecycle-aware integration helpers need a read-only way to discover
        # the immutable version/compilation and AtomicRule identifiers created
        # by the public import endpoint.  State transitions still go through
        # the real API below; exposing the test session factory avoids forging
        # review metadata in fixtures.
        test_client.session_factory = TestingSessionLocal
        yield test_client
    app.dependency_overrides.clear()
    engine.dispose()  # 释放该用例的引擎连接，避免跨用例连接堆积
    settings.STORAGE_ROOT = original_storage_root
    for key, value in original_runtime_settings.items():
        setattr(settings, key, value)


def review_rubric_via_api(client, rubric_id, *, session_factory=None):
    """Complete the real M4 rule/link/rubric review chain for a test rubric.

    The helper only reads identifiers from the fixture database.  Every audit
    event, reviewer identity and timestamp is produced by the public lifecycle
    endpoints, keeping older integration tests on the same strict path as
    production callers.
    """

    session_factory = session_factory or getattr(client, "session_factory", None)
    assert session_factory is not None, "client fixture must expose session_factory"
    with session_factory() as session:
        versions = session.scalars(
            select(models.RubricVersion).where(
                models.RubricVersion.rubric_id == rubric_id
            )
        ).all()
        assert len(versions) == 1, "draft rubric must have exactly one version"
        version = versions[0]
        rules = session.scalars(
            select(models.AtomicRule)
            .where(models.AtomicRule.rubric_version_id == version.id)
            .order_by(models.AtomicRule.rule_code)
        ).all()
        rule_states = [(rule.rule_code, rule.status) for rule in rules]
        rule_ids = [rule.id for rule in rules]
        links = (
            session.scalars(
                select(models.RuleTemplateLink)
                .where(models.RuleTemplateLink.rule_id.in_(rule_ids))
                .order_by(models.RuleTemplateLink.id)
            ).all()
            if rule_ids
            else []
        )
        link_states = [(link.id, link.review_status) for link in links]
        compilation_id = version.compilation_id
        version_id = version.id

    for rule_code, status in rule_states:
        if status == "draft":
            response = client.post(
                f"/api/rubrics/{rubric_id}/rules/{rule_code}/submit-review",
                json={"reason": "测试按 M4 生命周期提交规则审核"},
            )
            assert response.status_code == 200, response.text
            status = "review"
        if status == "review":
            response = client.post(
                f"/api/rubrics/{rubric_id}/rules/{rule_code}/approve",
                json={"reason": "测试按 M4 生命周期确认规则可执行"},
            )
            assert response.status_code == 200, response.text
        else:
            assert status == "approved", f"unsupported AtomicRule status: {status}"

    for link_id, status in link_states:
        if status == "pending":
            response = client.post(
                f"/api/rubrics/{rubric_id}/template-links/{link_id}/review",
                json={
                    "decision": "confirmed",
                    "reason": "测试按 M4 生命周期核对模板映射",
                },
            )
            assert response.status_code == 200, response.text
        else:
            assert status == "confirmed", f"unsupported template-link status: {status}"

    rubric = client.get(f"/api/rubrics/{rubric_id}")
    assert rubric.status_code == 200, rubric.text
    if rubric.json()["status"] == "draft":
        submitted = client.post(f"/api/rubrics/{rubric_id}/submit-review")
        assert submitted.status_code == 200, submitted.text
        assert submitted.json()["status"] == "review"
    else:
        assert rubric.json()["status"] == "review"

    return {
        "compilation_id": compilation_id,
        "rubric_version_id": version_id,
    }


def publish_rubric_via_api(client, rubric_id, *, session_factory=None):
    """Review and publish one imported rubric through the strict M4 API."""

    identity = review_rubric_via_api(
        client,
        rubric_id,
        session_factory=session_factory,
    )
    response = client.post(
        f"/api/rubrics/{rubric_id}/publish",
        json={
            "compilation_id": identity["compilation_id"],
            "reason": "测试按 M4 生命周期发布冻结版本",
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "published"
    return response.json(), identity


def create_legacy_unversioned_rubric_fixture(
    client,
    payload,
    *,
    session_factory=None,
):
    """Seed an explicitly historical, unversioned draft rubric.

    Only M0 characterization and M3 legacy/compare tests should use this
    fixture.  It deliberately creates no compilation/version/review metadata,
    so it cannot be mistaken for a formally published M4 rubric or bypass the
    strict lifecycle.  Current API-facing tests must use
    :func:`publish_rubric_via_api` instead.
    """

    session_factory = session_factory or getattr(client, "session_factory", None)
    assert session_factory is not None, "client fixture must expose session_factory"
    with session_factory() as session:
        user = ensure_dev_user(session)
        rubric = models.Rubric(
            name=payload["name"],
            version=payload.get("version") or "legacy-v1",
            description=payload.get("description"),
            total_score=payload["total_score"],
            status="draft",
            created_by=user.id,
            owner_id=user.id,
        )
        for index, raw in enumerate(payload.get("criteria") or []):
            criterion = RubricCriterionCreate.model_validate(raw)
            rubric.criteria.append(build_criterion(criterion, index))
        session.add(rubric)
        session.commit()
        return rubric.id


def make_sample_docx():
    document = Document()
    document.add_paragraph("基于机器学习的教学质量评价研究")
    document.add_paragraph("姓名：张三")
    document.add_paragraph("学号：20260001")
    document.add_paragraph("中文摘要")
    document.add_paragraph("本文围绕教学质量评价问题展开研究，说明研究背景、研究意义和应用价值。")
    document.add_paragraph("关键词：教学质量；机器学习；评价模型")
    document.add_paragraph("目录")
    document.add_paragraph("第一章 绪论")
    document.add_paragraph("本章介绍研究背景、研究意义、国内外相关工作和论文结构。")
    document.add_paragraph("第二章 文献综述")
    document.add_paragraph("国内外研究现状表明，教学质量评价需要结合多源数据和可解释分析。")
    document.add_paragraph("第三章 研究方法")
    document.add_paragraph("本文采用问卷调查、数据清洗、回归分析和分类模型对教学质量进行建模。")
    document.add_paragraph("第四章 实验结果与分析")
    document.add_paragraph("实验结果显示，模型能够识别关键影响因素，并通过对比实验支撑结论。")
    document.add_paragraph("第五章 创新点")
    document.add_paragraph("本文提出改进的指标加权方法，并增强评价结果解释能力。")
    document.add_paragraph("结论")
    document.add_paragraph("研究结论表明，该方法具有一定应用价值，后续可扩展到更多专业。")
    document.add_paragraph("参考文献")
    document.add_paragraph("[1] 张某某. 教学评价研究综述[J]. 教育研究, 2024.")
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def make_sample_docx_with_required_owner():
    """Return the sample thesis with the deterministic owner field present."""

    document = Document(make_sample_docx())
    document.add_paragraph("负责人：张三")
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def make_dalian_neusoft_cover_docx():
    document = Document()
    document.add_paragraph("大连东软信息学院")
    document.add_paragraph("毕业设计（论文）")
    document.add_paragraph("论文题目： 基于 Spring Boot 的网上图书商城管理系统的")
    document.add_paragraph("设计与实现")
    table = document.add_table(rows=7, cols=2)
    rows = [
        ("学    院：", "软件学院"),
        ("专    业：", "软件工程（专升本）"),
        ("学生姓名：", "王子铭"),
        ("学生学号：", "24201023601"),
        ("指导教师：", "刘真    周绍斌"),
        ("导师职称：", "助教    副教授"),
        ("完成日期：", "2026 年 4 月 13 日"),
    ]
    for row, (label, value) in zip(table.rows, rows):
        row.cells[0].text = label
        row.cells[1].text = value
    document.add_paragraph("中文摘要")
    document.add_paragraph("本文围绕网上图书商城管理系统展开研究，说明研究背景、研究意义和应用价值。")
    document.add_paragraph("关键词：Spring Boot；图书商城；管理系统")
    document.add_paragraph("目录")
    document.add_paragraph("第一章 绪论")
    document.add_paragraph("本章介绍研究背景、研究意义、国内外相关工作和论文结构。")
    document.add_paragraph("第二章 需求分析")
    document.add_paragraph("本章说明用户管理、图书管理、订单管理和后台管理等功能需求。")
    document.add_paragraph("第三章 系统设计")
    document.add_paragraph("本章说明系统架构、数据库设计、接口设计和核心模块划分。")
    document.add_paragraph("第四章 系统实现")
    document.add_paragraph("本章说明基于 Spring Boot 的业务模块实现、页面交互和测试结果。")
    document.add_paragraph("结论")
    document.add_paragraph("研究结论表明，该系统能够支撑网上图书商城管理的基本业务流程。")
    document.add_paragraph("参考文献")
    document.add_paragraph("[1] 王某某. Java Web 应用开发研究[J]. 软件工程, 2024.")
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def make_template_docx():
    document = Document()
    document.add_heading("本科毕业论文模板", level=1)
    document.add_heading("中文摘要", level=2)
    document.add_paragraph("摘要应说明研究背景、研究意义、研究方法和主要结论。")
    document.add_heading("第一章 绪论", level=2)
    document.add_paragraph("说明研究背景、研究意义和论文结构。")
    document.add_heading("第二章 文献综述", level=2)
    document.add_paragraph("梳理国内外研究现状和相关工作。")
    document.add_heading("第三章 研究方法", level=2)
    document.add_paragraph("说明数据来源、实验设计和模型方法。")
    document.add_heading("第四章 结果分析", level=2)
    document.add_heading("第五章 创新点", level=2)
    document.add_heading("结论", level=2)
    document.add_heading("参考文献", level=2)
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def make_rules_xlsx():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "评分规则"
    sheet.append(["编号", "评分项", "分值", "评分说明", "证据提示", "扣分规则"])
    sheet.append(["C01", "研究方法", 20, "方法合理，数据来源清楚。", "研究方法；实验设计", "方法说明不足扣分"])
    sheet.append(["C02", "文献综述", 15, "综述覆盖充分，能归纳研究现状。", "文献综述", "文献覆盖不足扣分"])
    sheet.append(["C03", "参考文献", 10, "引用规范，格式完整。", "参考文献", "参考文献不足扣分"])
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
