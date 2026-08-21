from docx import Document
from types import SimpleNamespace

from backend.app.services.document_parser.extractor import extract_document
from backend.app.services.document_parser.parser import interpret_thesis_document
from backend.app.services.document_parser.types import ParsedPaper
from backend.app.services.scoring.profiles.thesis import ThesisProfile


def test_generic_docx_extraction_is_layout_only_and_thesis_interpretation_is_explicit(
    tmp_path,
):
    path = tmp_path / "layered.docx"
    document = Document()
    document.add_heading("摘要", level=1)
    document.add_paragraph("本文研究通用评分持久化。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "学生姓名"
    table.cell(0, 1).text = "张三"
    document.add_heading("参考文献", level=1)
    document.add_paragraph("[1] Generic scoring architecture.")
    document.save(path)

    extracted = extract_document(path)

    assert extracted.schema_version == "extracted-document@1"
    assert extracted.media_type.endswith("wordprocessingml.document")
    assert [block.kind for block in extracted.blocks].count("table_row") == 1
    assert any(block.style_name.startswith("Heading") for block in extracted.blocks)
    assert "摘要" in extracted.heading_candidates
    assert not hasattr(extracted, "student_name")
    assert not hasattr(extracted, "references")

    parsed = interpret_thesis_document(extracted)

    assert isinstance(parsed, ParsedPaper)
    assert parsed.student_name == "张三"
    assert parsed.references == ["[1] Generic scoring architecture."]


def test_generic_extractor_rejects_unknown_media_without_profile_fallback(tmp_path):
    path = tmp_path / "unsupported.txt"
    path.write_text("not a supported document", encoding="utf-8")

    try:
        extract_document(path)
    except ValueError as exc:
        assert "unsupported file type" in str(exc)
    else:
        raise AssertionError("unsupported input must not be interpreted as a thesis")


def test_thesis_profile_explicitly_interprets_generic_extraction(tmp_path):
    path = tmp_path / "thesis-v2.docx"
    document = Document()
    document.add_heading("摘要", level=1)
    document.add_paragraph("本文验证通用提取层与论文解释层分离。")
    document.save(path)
    extracted = extract_document(path)
    submission = SimpleNamespace(
        id="submission-thesis-v2",
        source_artifact_hash="a" * 64,
        source_artifact_ref="blob:sha256:" + "a" * 64,
        submission_metadata={"student_name": "张三"},
    )

    snapshot = ThesisProfile().interpret_document(
        extracted_document=extracted,
        submission=submission,
    )

    assert snapshot["profile_key"] == "thesis"
    assert snapshot["profile_version"] == ThesisProfile.profile_version
    assert snapshot["parser_version"] == "generic-document-extractor@1"
    assert snapshot["normalizer_version"] == "thesis-document-interpreter@1"
    assert "thesis" in snapshot["profile_extensions"]
