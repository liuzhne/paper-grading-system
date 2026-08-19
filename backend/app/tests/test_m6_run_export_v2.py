from copy import deepcopy
import json
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from backend.app.db import models
from backend.app.services.scoring.profiles.registry import temporary_profile_registration
from backend.app.services.scoring.profiles.thesis import ThesisProfile
from backend.app.tests.test_m6_v2_submissions_api import (
    _ApiTechnicalProposalProfile,
    _create_batch_and_submission,
    _score,
)


GOLDEN = Path(__file__).with_name("golden") / "m6" / "run-export-v2-schema.json"


def _schema_projection(export, *, sheet_names, html):
    criterion = export["criteria"][0]
    rule = criterion["rules"][0]
    review = export["review_logs"][0]
    return {
        "criterion_keys": sorted(criterion),
        "html_markers": [
            marker
            for marker in (
                "评分运行报告",
                "运行身份",
                "评分项与规则",
                "证据",
                "人工复核记录",
                "Profile 扩展",
            )
            if marker in html
        ],
        "identity_keys": sorted(export["identity"]),
        "profile_extension_keys": sorted(export["profile_extensions"]),
        "review_log_keys": sorted(review),
        "rule_keys": sorted(rule),
        "schema": export["schema"],
        "sheet_names": sheet_names,
        "submission_keys": sorted(export["submission"]),
        "top_level_keys": sorted(export),
    }


def _keys(value):
    if isinstance(value, dict):
        for key, nested in value.items():
            yield key
            yield from _keys(nested)
    elif isinstance(value, list):
        for nested in value:
            yield from _keys(nested)


def test_run_export_v2_json_html_excel_schema_golden_and_profile_isolation(
    client,
    tmp_path,
):
    with temporary_profile_registration(_ApiTechnicalProposalProfile()):
        _batch, submission = _create_batch_and_submission(client, "export")
        with client.session_factory() as db:
            stored = db.get(models.Submission, submission["id"])
            stored.submission_metadata["project_name"] = (
                '<script>alert("not trusted")</script>'
            )
            db.commit()
        run = _score(client, submission["id"])
        risk = next(
            item for item in run["items"] if item["criterion_code"] == "RISK_CONTROL"
        )
        changed = client.patch(
            f"/api/v2/score-items/{risk['id']}",
            json={
                "final_score": 9,
                "reason": "A long reviewer explanation " + "x" * 180,
                "resolution_type": "ordinary_override",
            },
        )
        assert changed.status_code == 200, changed.text
        reviewed = client.post(
            f"/api/v2/scoring-runs/{run['id']}/review",
            json={"reason": "Export fixture reviewed."},
        )
        assert reviewed.status_code == 200, reviewed.text

        json_response = client.get(
            f"/api/v2/scoring-runs/{run['id']}/export.json"
        )
        assert json_response.status_code == 200, json_response.text
        exported = json_response.json()
        assert exported["schema"] == "grading-core/run-export@2"
        assert exported["identity"]["rubric_version_id"]
        assert exported["identity"]["runtime_identity"]
        assert exported["criteria"] == sorted(
            exported["criteria"],
            key=lambda item: item["criterion_code"],
        )
        assert all(
            criterion["rules"]
            == sorted(criterion["rules"], key=lambda item: item["rule_code"])
            for criterion in exported["criteria"]
        )
        assert exported["profile_extensions"]["metadata"] == {
            "project_name": '<script>alert("not trusted")</script>'
        }
        forbidden = {
            "student_id",
            "student_name",
            "paper",
            "title",
            "department",
            "major",
            "advisor",
            "references",
        }
        generic_without_extensions = deepcopy(exported)
        generic_without_extensions.pop("profile_extensions")
        assert forbidden.isdisjoint(set(_keys(generic_without_extensions)))
        serialized_extensions = json.dumps(
            exported["profile_extensions"],
            ensure_ascii=False,
        )
        for secret in ("vendor_name", "uploader_email", "student_id"):
            assert secret not in serialized_extensions

        html_response = client.get(
            f"/api/v2/scoring-runs/{run['id']}/report"
        )
        assert html_response.status_code == 200, html_response.text
        html = html_response.text
        assert "毕业论文" not in html
        assert "学生" not in html
        assert '<script>alert("not trusted")</script>' not in html
        assert "&lt;script&gt;alert" in html

        excel_response = client.get(
            f"/api/v2/scoring-runs/{run['id']}/export.xlsx"
        )
        assert excel_response.status_code == 200, excel_response.text
        workbook_path = tmp_path / "run-export-v2.xlsx"
        workbook_path.write_bytes(excel_response.content)
        workbook = load_workbook(workbook_path)
        expected_sheets = [
            "运行摘要",
            "评分项",
            "规则结果",
            "证据",
            "复核记录",
            "Profile扩展",
        ]
        assert workbook.sheetnames == expected_sheets
        for sheet in workbook.worksheets:
            assert sheet.freeze_panes == "A2"
            assert sheet.auto_filter.ref
            assert sheet.sheet_view.showGridLines is False
            assert all(cell.alignment.wrap_text for cell in sheet[1])
            assert all(
                sheet.column_dimensions[cell.column_letter].width > 0
                for cell in sheet[1]
            )
        assert workbook["复核记录"].column_dimensions["G"].width >= 40
        assert workbook["证据"].column_dimensions["F"].width >= 50
        summary_rows = {
            row[0].value: row[1]
            for row in workbook["运行摘要"].iter_rows(min_row=2)
        }
        anchor_hash = summary_rows[
            "identity.runtime_identity.calibration_anchors_hash"
        ]
        assert anchor_hash.value == "9" * 64
        assert anchor_hash.number_format == "@"
        assert anchor_hash.quotePrefix is True
        assert workbook["证据"]["C2"].value == "source_quote"
        assert "text_span" in workbook["证据"]["E2"].value

        actual_schema = _schema_projection(
            exported,
            sheet_names=workbook.sheetnames,
            html=html,
        )
        assert actual_schema == json.loads(GOLDEN.read_text(encoding="utf-8"))

        legacy = client.get(f"/api/scoring-runs/{run['id']}/export.json")
        assert legacy.status_code == 400
        assert "/api/v2" in legacy.json()["detail"]


def test_thesis_export_extensions_are_explicit_and_domain_scoped():
    submission = SimpleNamespace(
        submission_metadata={
            "student_id": "20260001",
            "student_name": "张三",
            "title": "论文标题",
            "project_name": "must-not-cross-profile-boundary",
        }
    )
    document_snapshot = SimpleNamespace(
        snapshot_payload={
            "profile_extensions": {
                "thesis": {
                    "references": ["[1] Reference"],
                    "coherence_findings": [{"kind": "citation"}],
                    "format_findings": [{"field": "font"}],
                }
            }
        }
    )
    run = SimpleNamespace(coherence_findings=[], format_findings=[])

    extension = ThesisProfile().build_export_extensions(
        submission=submission,
        document_snapshot=document_snapshot,
        run=run,
    )

    assert extension["metadata"] == {
        "student_id": "20260001",
        "student_name": "张三",
        "title": "论文标题",
    }
    assert "project_name" not in extension["metadata"]
    assert extension["findings"] == {
        "coherence": [{"kind": "citation"}],
        "format": [{"field": "font"}],
        "references": ["[1] Reference"],
    }
