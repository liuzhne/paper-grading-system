import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy import select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.app.core.config import settings
from backend.app.db.models import Base
from backend.app.db.models import Rubric
from backend.app.db.models import RubricCriterion
from backend.app.db.models import ScoreItem
from backend.app.eval.labeled_dataset import build_labeled_eval
from backend.app.eval.labeled_dataset import load_scores_table
from backend.app.eval.gating import summarize_run_identities
from backend.app.services.llm.mock import MockLLMScorer
from backend.app.tests.conftest import make_sample_docx
from backend.app.tests.conftest import publish_rubric_via_api


def test_load_scores_table_xlsx(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["文件名", "总分", "C01", "C02"])
    sheet.append(["a.docx", 88, 18, 17])
    sheet.append(["b.docx", 72, 12, 14])
    path = tmp_path / "s.xlsx"
    workbook.save(path)

    rows = load_scores_table(str(path))
    assert rows[0] == {"filename": "a.docx", "total": 88.0, "items": {"C01": 18.0, "C02": 17.0}}
    assert len(rows) == 2


def test_load_scores_table_csv(tmp_path):
    path = tmp_path / "s.csv"
    path.write_text("filename,total,C01\nx.docx,80,40\n", encoding="utf-8")
    assert load_scores_table(str(path)) == [{"filename": "x.docx", "total": 80.0, "items": {"C01": 40.0}}]


def test_load_scores_table_requires_filename_and_total(tmp_path):
    workbook = Workbook()
    workbook.active.append(["无关列", "另一列"])
    workbook.active.append([1, 2])
    path = tmp_path / "bad.xlsx"
    workbook.save(path)
    with pytest.raises(ValueError):
        load_scores_table(str(path))


def test_build_labeled_eval_pipeline_runs_end_to_end(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "STORAGE_ROOT", tmp_path / "storage")
    engine = create_engine(
        "sqlite+pysqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine)()
    try:
        rubric = Rubric(name="毕设标准", version="v1.0", total_score=20)
        rubric.criteria.append(RubricCriterion(code="C01", name="研究方法", max_score=10))
        rubric.criteria.append(RubricCriterion(code="C02", name="参考文献", max_score=10))
        db.add(rubric)
        db.commit()

        papers = tmp_path / "papers"
        papers.mkdir()
        for name in ("p0.docx", "p1.docx"):
            (papers / name).write_bytes(make_sample_docx().getvalue())

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["文件名", "总分", "C01", "C02"])
        sheet.append(["p0.docx", 18, 9, 9])
        sheet.append(["p1.docx", 14, 7, 7])
        scores = tmp_path / "scores.xlsx"
        workbook.save(scores)

        report = build_labeled_eval(db, rubric.id, str(papers), str(scores), scorer=MockLLMScorer())
        # 验证管线连通（分数为合成，QWK 数值无意义；真实验收需真实教师分）。
        assert report["n"] == 2
        assert report["mae"] is not None
        assert report["qwk"] is not None
        assert "C01" in report["per_criterion"]
        assert report["dataset_size"] == 2
        assert report["errors"] == []
        assert report["invalid_evidence_rate"] == 0.0
        assert report["invalid_evidence_item_count"] == 0
        assert report["evaluated_item_count"] == 4
    finally:
        db.close()


def test_build_labeled_eval_chunks_visible_without_autoflush(tmp_path, monkeypatch):
    """回归：autoflush=False 会话（cli/db.py 同配置）下"导入后立即评分"必须看得到 chunk，
    否则检索证据为空、模型全给 0 分（QWK 假塌方）。"""
    monkeypatch.setattr(settings, "STORAGE_ROOT", tmp_path / "storage")
    engine = create_engine(
        "sqlite+pysqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine, autoflush=False)()
    try:
        rubric = Rubric(name="毕设标准", version="v1.0", total_score=20)
        rubric.criteria.append(RubricCriterion(code="C01", name="研究方法", max_score=10))
        rubric.criteria.append(RubricCriterion(code="C02", name="参考文献", max_score=10))
        db.add(rubric)
        db.commit()

        papers = tmp_path / "papers"
        papers.mkdir()
        (papers / "p0.docx").write_bytes(make_sample_docx().getvalue())

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["文件名", "总分", "C01", "C02"])
        sheet.append(["p0.docx", 18, 9, 9])
        scores = tmp_path / "scores.xlsx"
        workbook.save(scores)

        report = build_labeled_eval(db, rubric.id, str(papers), str(scores), scorer=MockLLMScorer())
        assert report["errors"] == []
        modes = [
            (item.raw_model_output or {}).get("chunk_evaluation_mode")
            for item in db.scalars(select(ScoreItem)).all()
        ]
        assert modes and "single-empty-evidence" not in modes
    finally:
        db.close()


def test_formal_labeled_eval_retains_exact_thesis_runtime_identity(client, tmp_path):
    payload = {
        "name": "正式 QWK 身份标准",
        "version": "v1",
        "total_score": 10,
        "criteria": [
            {
                "code": "T01",
                "name": "研究质量",
                "max_score": 10,
                "criterion_type": "deterministic",
                "scoring_mode": "deductive",
                "deduction_rules_structured": [
                    {
                        "match": "研究质量不足",
                        "points": 10,
                        "reason": "研究质量不足",
                        "checker_key": "thesis.legacy_required_fields.v1",
                        "checker_params": {
                            "criterion_code": "T01",
                            "applies_to": "global",
                        },
                    }
                ],
                "display_order": 1,
            }
        ],
    }
    rubric_id = client.post("/api/rubrics", json=payload).json()["id"]
    _published, version_identity = publish_rubric_via_api(client, rubric_id)

    papers = tmp_path / "formal-papers"
    papers.mkdir()
    (papers / "p0.docx").write_bytes(make_sample_docx().getvalue())
    workbook = Workbook()
    workbook.active.append(["文件名", "总分", "T01"])
    workbook.active.append(["p0.docx", 10, 10])
    scores = tmp_path / "formal-scores.xlsx"
    workbook.save(scores)

    with client.session_factory() as db:
        report = build_labeled_eval(
            db,
            rubric_id,
            papers,
            scores,
            scorer=MockLLMScorer(),
            sample_ids_by_filename={"p0.docx": "9" * 64},
        )

    identity = report["per_sample"][0]["run_identity"]
    assert identity["schema_version"] == "paper-grading/thesis-eval-run-identity@1"
    assert identity["rubric_version_id"] == version_identity["rubric_version_id"]
    assert identity["business_profile_version"]
    assert identity["prompt_version"]
    assert identity["policy_snapshot_sha256"]
    assert identity["grade_scale_sha256"]
    assert identity["runtime_identity_sha256"]
    evaluation_identity = report["evaluation_identity"]
    assert evaluation_identity["rubric_version_id"] == version_identity[
        "rubric_version_id"
    ]
    assert evaluation_identity["business_profile_key"] == "thesis"
    assert evaluation_identity["policy_hash"] == identity["policy_hash"]
    assert evaluation_identity["grade_scale_sha256"] == identity[
        "grade_scale_sha256"
    ]
    assert evaluation_identity["rounding"] == {
        "mode": "half_up",
        "digits": 2,
    }
    assert report["grade_scale_sha256"] == identity["grade_scale_sha256"]
    _summary, issues = summarize_run_identities(report)
    assert issues == []
