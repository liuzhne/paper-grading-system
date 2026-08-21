"""M4 import, compilation and dual-write executable specification.

Only the exact absence of a named M4 public capability becomes a strict
XFAIL.  Once a symbol exists, signature errors, internal import errors,
transaction leaks and behavioral defects are ordinary failures.

The contract is intentionally two-phase:

* ``prepare_*`` parses files and may call the compiler/LLM without a database
  session, returning an immutable complete in-memory graph.
* ``persist_prepared_import`` writes that graph in one short transaction.  The
  provenance/AtomicRule graph is authoritative; ``RubricCriterion`` remains a
  compatibility projection only.
"""

from __future__ import annotations

from collections.abc import Mapping
from copy import deepcopy
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from decimal import Decimal
import importlib
from io import BytesIO
from types import ModuleType
from unittest.mock import Mock
import xml.etree.ElementTree as ET
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, event, func, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.app.db import models
from backend.app.db.sqlite import enable_sqlite_foreign_keys
from backend.app.services.scoring.adapters.rubric_snapshot import (
    CompiledRubricSnapshotLoader,
)
from backend.app.services.scoring.core.canonical import canonical_sha256
from backend.app.services.scoring.core.checker_registry import (
    VersionedCheckerRegistry,
)
from backend.app.services.scoring.core.execution_plan import (
    RuleExecutionPlanBuilder,
)
from backend.app.tests.m2_contract_fixtures import (
    PROFILE_KEY,
    PROFILE_VERSION,
    technical_policy_snapshot_payload,
)
from backend.app.tests.m3_contract_fixtures import (
    CHECKER_KEY,
    CHECKER_VERSION,
    ENGINE_CONTRACT_VERSION,
    POLICY_COMPILER_VERSION,
    checker_registration_payload,
    deterministic_missing_owner_observation,
    idempotency_projection,
    scoring_request_payload,
    semantic_high_band_response,
)
from backend.app.tests.m4_import_contract_fixtures import (
    COMMENT_ANCHOR,
    COMMENT_ID,
    COMMENT_TEXT,
    COMPILER_VERSION,
    FIT_RULE_CODE,
    PARSER_VERSION,
    PREPARED_SCHEMA_VERSION,
    PROMPT_VERSION,
    RISK_RULE_CODE,
    ambiguous_rules_xlsx_bytes,
    canonical_json_bytes,
    file_import_command,
    hybrid_import_command,
    hybrid_rules_xlsx_bytes,
    inspect_word_comment_without_production_parser,
    legacy_upgrade_command,
    manual_import_command,
    manual_rubric_payload,
    real_rules_xlsx_bytes,
    real_template_docx_bytes,
    sha256_bytes,
    weighted_policy_payload,
)


PIPELINE_MODULE = "backend.app.services.rubric_import.pipeline"
LIFECYCLE_MODULE = "backend.app.services.rubrics.lifecycle"
RULE_EXECUTOR_MODULE = "backend.app.services.scoring.core.rule_executor"
PIPELINE_SYMBOLS = (
    "prepare_file_import",
    "prepare_manual_json_import",
    "prepare_legacy_draft_upgrade",
    "persist_prepared_import",
    "rebuild_legacy_projection",
)
LIFECYCLE_SYMBOLS = (
    "submit_atomic_rule_for_review",
    "approve_atomic_rule",
    "review_template_link",
    "submit_for_review",
    "return_to_draft",
    "publish_rubric",
)


class M4ImportCapabilityUnavailable(RuntimeError):
    """The sole exception allowed to turn an M4 import gate into XFAIL."""


@dataclass(frozen=True, slots=True)
class _CapabilityProbe:
    target: str
    module: ModuleType | None
    symbols: Mapping[str, object]
    import_error: ModuleNotFoundError | None = None

    def has(self, name: str) -> bool:
        return self.module is not None and name in self.symbols

    def require(self, name: str):
        if not self.has(name):
            error = M4ImportCapabilityUnavailable(
                f"{self.target} must expose M4 capability {name}"
            )
            if self.import_error is not None:
                raise error from self.import_error
            raise error
        return self.symbols[name]


def _probe_module(target: str, required: tuple[str, ...]) -> _CapabilityProbe:
    try:
        module = importlib.import_module(target)
    except ModuleNotFoundError as exc:
        target_missing = exc.name == target or (
            exc.name is not None and target.startswith(exc.name + ".")
        )
        if not target_missing:
            raise
        return _CapabilityProbe(target, None, {}, exc)
    namespace = vars(module)
    return _CapabilityProbe(
        target,
        module,
        {name: namespace[name] for name in required if name in namespace},
    )


_PIPELINE = _probe_module(PIPELINE_MODULE, PIPELINE_SYMBOLS)
_LIFECYCLE = _probe_module(LIFECYCLE_MODULE, LIFECYCLE_SYMBOLS)
_RULE_EXECUTOR = _probe_module(RULE_EXECUTOR_MODULE, ("execute_rule_plan",))
_UNSET = object()


def _requires(*requirements: tuple[_CapabilityProbe, str]):
    missing = [
        f"{probe.target}.{name}"
        for probe, name in requirements
        if not probe.has(name)
    ]
    return pytest.mark.xfail(
        condition=bool(missing),
        reason="M4 import capability is not implemented: " + ", ".join(missing),
        raises=M4ImportCapabilityUnavailable,
        strict=True,
    )


requires_file_prepare = _requires((_PIPELINE, "prepare_file_import"))
requires_file_persist = _requires(
    (_PIPELINE, "prepare_file_import"),
    (_PIPELINE, "persist_prepared_import"),
)
requires_projection_rebuild = _requires(
    (_PIPELINE, "prepare_file_import"),
    (_PIPELINE, "persist_prepared_import"),
    (_PIPELINE, "rebuild_legacy_projection"),
)
requires_manual = _requires(
    (_PIPELINE, "prepare_manual_json_import"),
    (_PIPELINE, "persist_prepared_import"),
)
requires_legacy_upgrade = _requires(
    (_PIPELINE, "prepare_legacy_draft_upgrade"),
    (_PIPELINE, "persist_prepared_import"),
)
_SIGNING_SYMBOLS = (
    "submit_atomic_rule_for_review",
    "approve_atomic_rule",
    "review_template_link",
    "submit_for_review",
    "publish_rubric",
)
requires_manual_signing = _requires(
    (_PIPELINE, "prepare_manual_json_import"),
    (_PIPELINE, "persist_prepared_import"),
    *[(_LIFECYCLE, name) for name in _SIGNING_SYMBOLS],
)
requires_legacy_signing = _requires(
    (_PIPELINE, "prepare_legacy_draft_upgrade"),
    (_PIPELINE, "persist_prepared_import"),
    (_LIFECYCLE, "return_to_draft"),
    *[(_LIFECYCLE, name) for name in _SIGNING_SYMBOLS],
)
requires_vertical = _requires(
    (_PIPELINE, "prepare_file_import"),
    (_PIPELINE, "persist_prepared_import"),
    (_RULE_EXECUTOR, "execute_rule_plan"),
    *[(_LIFECYCLE, name) for name in LIFECYCLE_SYMBOLS],
)


def _thaw(value):
    if isinstance(value, Mapping):
        return {str(key): _thaw(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_thaw(item) for item in value]
    return deepcopy(value)


def _mapping(value, *, label: str) -> dict:
    if isinstance(value, Mapping):
        return _thaw(value)
    method = getattr(value, "to_mapping", None)
    assert callable(method), f"{label} must expose to_mapping()"
    mapped = method()
    assert isinstance(mapped, Mapping), f"{label}.to_mapping() must return a mapping"
    return _thaw(mapped)


def _prepare_file(*, command=None, rules=_UNSET, template=_UNSET, scorer=None):
    prepare = _PIPELINE.require("prepare_file_import")
    value = prepare(
        command=deepcopy(command or file_import_command()),
        rules_bytes=real_rules_xlsx_bytes() if rules is _UNSET else rules,
        template_bytes=(
            real_template_docx_bytes() if template is _UNSET else template
        ),
        scorer=scorer,
    )
    graph = _mapping(value, label="prepared import graph")
    assert graph["schema_version"] == PREPARED_SCHEMA_VERSION
    return value, graph


def _persist(db, prepared, *, actor_id: str, target_rubric_id: str | None = None):
    persist = _PIPELINE.require("persist_prepared_import")
    kwargs = {
        "session": db,
        "prepared": prepared,
        "actor_id": actor_id,
    }
    if target_rubric_id is not None:
        kwargs["target_rubric_id"] = target_rubric_id
    result = persist(**kwargs)
    mapped = _mapping(result, label="persisted import identity")
    required = {"rubric_id", "compilation_id", "rubric_version_id"}
    assert required.issubset(mapped), "persist result must expose all graph root IDs"
    return mapped


def _rows(graph: Mapping[str, object], name: str) -> list[dict]:
    value = graph.get(name)
    assert isinstance(value, list), f"prepared graph.{name} must be a list"
    assert all(isinstance(item, dict) for item in value)
    return value


def _artifact(graph: dict, token: str) -> dict:
    matches = [
        item
        for item in _rows(graph, "artifacts")
        if token in str(item.get("artifact_type", "")).lower()
    ]
    assert len(matches) == 1
    return matches[0]


def _blocker_codes(graph: dict) -> set[str]:
    compilation = graph.get("compilation")
    assert isinstance(compilation, dict)
    blockers = compilation.get("blockers")
    assert isinstance(blockers, list)
    return {
        str(item["code"] if isinstance(item, Mapping) else item)
        for item in blockers
    }


@pytest.fixture()
def m4_db():
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    enable_sqlite_foreign_keys(engine)
    models.Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine, autoflush=False)()
    actor = models.User(
        id="00000000-0000-0000-0000-000000000441",
        username="m4-import-author",
        display_name="M4 import author",
        role="rubric_editor",
    )
    reviewer = models.User(
        id="00000000-0000-0000-0000-000000000442",
        username="m4-import-reviewer",
        display_name="M4 import reviewer",
        role="rubric_reviewer",
    )
    session.add_all([actor, reviewer])
    session.commit()
    try:
        yield session, actor, reviewer
    finally:
        session.close()
        engine.dispose()


class _DraftRuleScorer:
    def __init__(self, db=None):
        self.db = db
        self.calls: list[tuple[str, dict]] = []

    def complete_json(self, instructions, payload):
        if self.db is not None:
            assert not self.db.in_transaction(), (
                "LLM compilation must finish before the persistence transaction starts"
            )
        self.calls.append((instructions, deepcopy(payload)))
        return {
            "rules": [
                {
                    "match": "论证薄弱",
                    "points": 2,
                    "reason": "论证链不完整",
                }
            ]
        }


def test_m4_real_excel_word_fixture_has_independent_hashes_and_locators():
    rules = real_rules_xlsx_bytes()
    template = real_template_docx_bytes()
    workbook = load_workbook(filename=BytesIO(rules), data_only=True)
    sheet = workbook["Atomic Rules"]

    assert sheet.max_row == 3
    assert file_import_command()["schema_version"] == "rubric-import-command@2"
    assert sheet["A2"].value == RISK_RULE_CODE
    assert sheet["A3"].value == FIT_RULE_CODE
    assert len(sha256_bytes(rules)) == len(sha256_bytes(template)) == 64
    assert sha256_bytes(rules) != sha256_bytes(template)
    assert inspect_word_comment_without_production_parser(template) == {
        "comment_id": COMMENT_ID,
        "comment_text": COMMENT_TEXT,
        "anchor_text": COMMENT_ANCHOR,
    }
    reopened = Document(BytesIO(template))
    assert COMMENT_ANCHOR in [paragraph.text for paragraph in reopened.paragraphs]
    with zipfile.ZipFile(BytesIO(template)) as archive:
        content_types = ET.fromstring(archive.read("[Content_Types].xml"))
        relationships = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
    assert any(
        item.get("PartName") == "/word/comments.xml"
        for item in content_types.iter()
    )
    assert any(
        item.get("Type", "").endswith("/comments")
        and item.get("Target") == "comments.xml"
        for item in relationships.iter()
    )


def test_m4_manual_json_fixture_has_stable_content_address():
    payload = manual_rubric_payload()
    first = canonical_json_bytes(payload)
    reordered = {key: payload[key] for key in reversed(payload)}

    assert canonical_json_bytes(reordered) == first
    assert sha256_bytes(first) == sha256_bytes(canonical_json_bytes(reordered))


@requires_file_prepare
def test_file_prepare_retains_excel_word_and_compiler_provenance_dto_fields():
    rules = real_rules_xlsx_bytes()
    template = real_template_docx_bytes()
    prepared, graph = _prepare_file(rules=rules, template=template)

    assert graph["source_kind"] == "file_import"
    compilation = graph["compilation"]
    assert {
        "status",
        "parser_version",
        "compiler_version",
        "prompt_version",
        "model_provider",
        "model_name",
        "sampling_params",
        "raw_parse_output",
        "raw_model_output",
        "validation_result",
        "blockers",
        "warnings",
        "human_changes",
    }.issubset(compilation)
    assert compilation["parser_version"] == PARSER_VERSION
    assert compilation["compiler_version"] == COMPILER_VERSION
    assert compilation["prompt_version"] == PROMPT_VERSION
    assert compilation["human_changes"] == []

    excel = _artifact(graph, "excel")
    word = _artifact(graph, "word")
    assert excel["file_name"] == "m4-rules.xlsx"
    assert excel["file_hash"] == sha256_bytes(rules)
    assert excel["file_size_bytes"] == len(rules)
    assert word["file_name"] == "m4-template.docx"
    assert word["file_hash"] == sha256_bytes(template)
    assert word["file_size_bytes"] == len(template)

    source_rules = _rows(graph, "source_rules")
    by_code = {item["source_rule_code"]: item for item in source_rules}
    assert {RISK_RULE_CODE, FIT_RULE_CODE}.issubset(by_code)
    assert by_code[RISK_RULE_CODE]["sheet_name"] == "Atomic Rules"
    assert by_code[RISK_RULE_CODE]["row_number"] == 2
    assert "A2" in by_code[RISK_RULE_CODE]["cell_locator"]
    assert by_code[RISK_RULE_CODE]["raw_text"]
    assert by_code[FIT_RULE_CODE]["row_number"] == 3
    assert "A3" in by_code[FIT_RULE_CODE]["cell_locator"]

    template_items = _rows(graph, "template_items")
    comment_items = [item for item in template_items if item["kind"] == "comment"]
    assert len(comment_items) == 1
    comment = comment_items[0]
    assert comment["raw_text"] == COMMENT_TEXT
    assert comment["source_locator"]["comment_id"] == COMMENT_ID
    assert comment["source_locator"]["anchor_text"] == COMMENT_ANCHOR
    assert "风险控制" in comment["section_path"]
    assert comment["source_hash"] == sha256_bytes(COMMENT_TEXT.encode("utf-8"))

    # A prepared graph is an immutable DTO; a thawed transport copy may change
    # without mutating the object that will be persisted.
    before = _mapping(prepared, label="prepared import graph")
    transport_copy = _mapping(prepared, label="prepared import graph")
    transport_copy["compilation"]["parser_version"] = "tampered"
    assert _mapping(prepared, label="prepared import graph") == before
    with pytest.raises((AttributeError, TypeError)):
        prepared["source_kind"] = "tampered"


@requires_file_persist
def test_llm_compiled_rules_remain_draft_and_never_become_scoring_authority(m4_db):
    db, actor, _ = m4_db
    command = file_import_command()
    command["rubric"]["name"] = "M4 LLM draft rule"
    command["rubric"]["global_policy"] = technical_policy_snapshot_payload(
        total_score="10", rounding_digits=2
    )
    command["files"]["template_file_name"] = None
    command["compiler"].update(
        {
            "model_provider": "openai_compatible",
            "model_name": "rubric-compiler-test-model",
            "sampling_params": {"temperature": "0", "seed": 19},
        }
    )
    assert not db.in_transaction()
    scorer = _DraftRuleScorer(db)
    prepared, graph = _prepare_file(
        command=command,
        rules=ambiguous_rules_xlsx_bytes(),
        template=None,
        scorer=scorer,
    )

    assert scorer.calls
    assert not db.in_transaction()
    assert graph["compilation"]["model_provider"] == "openai_compatible"
    assert graph["compilation"]["model_name"] == "rubric-compiler-test-model"
    assert graph["compilation"]["sampling_params"] == {
        "temperature": "0",
        "seed": 19,
    }
    llm_rules = [
        rule
        for rule in _rows(graph, "atomic_rules")
        if rule["creation_method"] == "llm"
    ]
    assert llm_rules
    assert {rule["status"] for rule in llm_rules} == {"draft"}
    assert all(rule.get("reviewed_by") is None for rule in llm_rules)
    assert all(rule.get("reviewed_at") is None for rule in llm_rules)

    # A blocked compilation is still an auditable draft graph; only after all
    # parser/compiler/LLM work has completed may the short write transaction
    # begin.
    identity = _persist(db, prepared, actor_id=actor.id)
    assert not db.in_transaction()
    persisted_compilation = db.get(
        models.RubricCompilation,
        identity["compilation_id"],
    )
    assert persisted_compilation.status != "published"
    assert any(
        (item.get("code") if isinstance(item, Mapping) else item)
        == "MISSING_EXECUTABLE_SCORING_MODE"
        for item in persisted_compilation.blockers
    )


@requires_file_prepare
def test_missing_band_or_deduct_contract_creates_blocked_review_only_draft():
    command = file_import_command()
    command["rubric"]["name"] = "M4 ambiguous draft"
    command["rubric"]["global_policy"] = technical_policy_snapshot_payload(
        total_score="10", rounding_digits=2
    )
    command["files"]["template_file_name"] = None
    _, graph = _prepare_file(
        command=command,
        rules=ambiguous_rules_xlsx_bytes(),
        template=None,
        scorer=_DraftRuleScorer(),
    )

    assert "MISSING_EXECUTABLE_SCORING_MODE" in _blocker_codes(graph)
    criteria = _rows(graph, "criteria")
    assert len(criteria) == 1
    assert criteria[0]["scoring_mode"] == "review_only"
    assert criteria[0]["scoring_mode"] != "llm_direct"
    projection = graph["legacy_projection"]
    assert projection["criteria"][0]["scoring_mode"] == "review_only"
    assert all(
        rule.get("direction") != "llm_direct"
        for rule in _rows(graph, "atomic_rules")
    )


@requires_file_prepare
def test_empty_criterion_description_is_an_independent_publish_blocker():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "评分规则"
    sheet.append(["编号", "评分项", "分值", "评分说明", "评分模式"])
    sheet.append(["T02", "分析与解决问题", 20, None, "review_only"])
    rules = BytesIO()
    workbook.save(rules)
    command = file_import_command()
    command["rubric"]["name"] = "M4 missing criterion description"
    command["rubric"]["global_policy"] = technical_policy_snapshot_payload(
        total_score="20", rounding_digits=2
    )
    command["files"]["template_file_name"] = None

    _, graph = _prepare_file(
        command=command,
        rules=rules.getvalue(),
        template=None,
    )

    assert "MISSING_CRITERION_DESCRIPTION" in _blocker_codes(graph)
    blocker = next(
        item
        for item in graph["compilation"]["blockers"]
        if item["code"] == "MISSING_CRITERION_DESCRIPTION"
    )
    assert blocker["criterion_code"] == "T02"
    assert "scoring description" in blocker["message"]


@requires_file_prepare
def test_formal_hybrid_expands_to_stable_single_semantic_leaf_criteria_in_points_mode():
    rules = hybrid_rules_xlsx_bytes()
    command = hybrid_import_command(weighted=False)
    _, first = _prepare_file(command=command, rules=rules, template=None)
    _, second = _prepare_file(command=command, rules=rules, template=None)

    leaves = [
        item
        for item in _rows(first, "criteria")
        if item.get("parent_criterion_code") == "IMPLEMENTATION"
    ]
    second_leaves = [
        item
        for item in _rows(second, "criteria")
        if item.get("parent_criterion_code") == "IMPLEMENTATION"
    ]
    assert len(leaves) == len(second_leaves) == 2
    assert [item["code"] for item in leaves] == [
        item["code"] for item in second_leaves
    ]
    assert len({item["code"] for item in leaves}) == 2
    assert all(item["code"] != "IMPLEMENTATION" for item in leaves)
    assert {item["scoring_mode"] for item in leaves} == {"deductive", "banded"}
    assert all(item["weight"] is None for item in leaves)
    assert sum(Decimal(str(item["max_score"])) for item in leaves) == Decimal("20")
    assert not any(
        item["code"] == "IMPLEMENTATION"
        for item in graph_scoring_projection(first)
    )


def graph_scoring_projection(graph: dict) -> list[dict]:
    projection = graph.get("legacy_projection")
    assert isinstance(projection, dict)
    criteria = projection.get("criteria")
    assert isinstance(criteria, list)
    return criteria


@requires_file_prepare
def test_formal_hybrid_derives_exact_leaf_weights_without_parent_double_weighting():
    command = hybrid_import_command(weighted=True)
    assert command["rubric"]["global_policy"]["aggregation"]["mode"] == (
        "weighted_normalized"
    )
    rules = hybrid_rules_xlsx_bytes(parent_weight="60")
    _, graph = _prepare_file(command=command, rules=rules, template=None)
    leaves = [
        item
        for item in _rows(graph, "criteria")
        if item.get("parent_criterion_code") == "IMPLEMENTATION"
    ]

    assert len(leaves) == 2
    assert [Decimal(str(item["weight"])) for item in leaves] == [
        Decimal("24"),
        Decimal("36"),
    ]
    projection = graph_scoring_projection(graph)
    assert {item["code"] for item in projection} == {
        item["code"] for item in leaves
    }
    assert sum(Decimal(str(item["weight"])) for item in projection) == Decimal("60")
    assert "IMPLEMENTATION" not in {item["code"] for item in projection}


@requires_file_prepare
def test_formal_hybrid_blocks_inexact_derived_leaf_weights_instead_of_rounding():
    command = hybrid_import_command(weighted=True)
    command["rubric"]["global_policy"] = weighted_policy_payload(
        total_score="3"
    )
    rules = hybrid_rules_xlsx_bytes(
        parent_weight="1",
        first_points="1",
        second_points="2",
    )
    _, graph = _prepare_file(command=command, rules=rules, template=None)

    assert "HYBRID_WEIGHT_NOT_EXACT" in _blocker_codes(graph)
    leaves = [
        item
        for item in _rows(graph, "criteria")
        if item.get("parent_criterion_code") == "IMPLEMENTATION"
    ]
    assert not any(
        Decimal(str(item["weight"] or "0")) in {Decimal("0.33"), Decimal("0.67")}
        for item in leaves
    )


@requires_file_persist
def test_prepare_runs_before_one_short_transaction_that_atomically_writes_full_graph(m4_db):
    db, actor, _ = m4_db
    scorer = _DraftRuleScorer(db)
    command = file_import_command()
    rules = real_rules_xlsx_bytes()
    template = real_template_docx_bytes()

    assert not db.in_transaction()
    prepared, graph = _prepare_file(
        command=command,
        rules=rules,
        template=template,
        scorer=scorer,
    )
    assert not db.in_transaction()

    transactions = {"begin": 0, "commit": 0}

    def after_begin(session, transaction, connection):
        del session, transaction, connection
        transactions["begin"] += 1

    def after_commit(session):
        del session
        transactions["commit"] += 1

    event.listen(db, "after_begin", after_begin)
    event.listen(db, "after_commit", after_commit)
    persisted = _persist(db, prepared, actor_id=actor.id)
    event.remove(db, "after_begin", after_begin)
    event.remove(db, "after_commit", after_commit)

    assert transactions == {"begin": 1, "commit": 1}
    assert not db.in_transaction()
    assert db.get(models.Rubric, persisted["rubric_id"]).status == "draft"
    assert db.scalar(
        select(func.count()).select_from(models.RubricCompilation)
    ) == 1
    assert db.scalar(select(func.count()).select_from(models.SourceArtifact)) == 2
    assert db.scalar(select(func.count()).select_from(models.SourceRule)) >= 2
    assert db.scalar(select(func.count()).select_from(models.TemplateItem)) >= 1
    assert db.scalar(select(func.count()).select_from(models.RubricVersion)) == 1
    assert db.scalar(select(func.count()).select_from(models.AtomicRule)) == len(
        _rows(graph, "atomic_rules")
    )
    assert db.scalar(select(func.count()).select_from(models.RuleLevel)) == 2

    rules_by_code = {
        item.rule_code: item
        for item in db.scalars(select(models.AtomicRule)).all()
    }
    assert {RISK_RULE_CODE, FIT_RULE_CODE} == set(rules_by_code)
    assert rules_by_code[RISK_RULE_CODE].source_rules
    assert rules_by_code[FIT_RULE_CODE].source_rules
    assert rules_by_code[FIT_RULE_CODE].levels


@requires_file_persist
def test_full_graph_persistence_rolls_back_every_projection_and_provenance_row(m4_db):
    db, actor, _ = m4_db
    prepared, _ = _prepare_file()

    def explode(session, flush_context, instances):
        del session, flush_context, instances
        raise RuntimeError("M4 injected graph write failure")

    event.listen(db, "before_flush", explode, once=True)
    with pytest.raises(RuntimeError, match="injected graph write failure"):
        _persist(db, prepared, actor_id=actor.id)
    assert not db.in_transaction(), "persistence service must roll back its failed transaction"

    table_models = (
        models.Rubric,
        models.RubricCriterion,
        models.RubricCompilation,
        models.SourceArtifact,
        models.SourceRule,
        models.TemplateItem,
        models.RubricVersion,
        models.AtomicRule,
        models.RuleLevel,
        models.RuleTemplateLink,
    )
    assert {
        model.__tablename__: db.scalar(select(func.count()).select_from(model))
        for model in table_models
    } == {model.__tablename__: 0 for model in table_models}


@requires_projection_rebuild
def test_legacy_criterion_is_rebuilt_from_authoritative_graph_not_a_second_truth(m4_db):
    db, actor, _ = m4_db
    prepared, graph = _prepare_file()
    identity = _persist(db, prepared, actor_id=actor.id)
    authoritative = {
        item["code"]: item for item in graph_scoring_projection(graph)
    }
    stored = db.scalars(
        select(models.RubricCriterion).where(
            models.RubricCriterion.rubric_id == identity["rubric_id"]
        )
    ).all()
    assert {item.code for item in stored} == set(authoritative)
    for item in stored:
        assert Decimal(str(item.max_score)) == Decimal(
            str(authoritative[item.code]["max_score"])
        )

    stored[0].max_score = Decimal("1")
    stored[0].scoring_mode = "llm_direct"
    db.commit()
    rebuild = _PIPELINE.require("rebuild_legacy_projection")
    rebuild(session=db, rubric_id=identity["rubric_id"])
    db.commit()
    refreshed = db.get(models.RubricCriterion, stored[0].id)
    assert Decimal(str(refreshed.max_score)) == Decimal(
        str(authoritative[refreshed.code]["max_score"])
    )
    assert refreshed.scoring_mode == authoritative[refreshed.code]["scoring_mode"]
    assert refreshed.scoring_mode != "llm_direct"


@requires_manual
def test_manual_json_create_builds_source_artifact_source_rules_and_same_draft_graph(m4_db):
    db, actor, _ = m4_db
    prepare = _PIPELINE.require("prepare_manual_json_import")
    command = manual_import_command()
    prepared = prepare(command=deepcopy(command))
    graph = _mapping(prepared, label="manual JSON prepared graph")
    assert graph["schema_version"] == PREPARED_SCHEMA_VERSION
    assert graph["source_kind"] == "manual_json"
    artifact = _artifact(graph, "manual_json")
    source_bytes = canonical_json_bytes(manual_rubric_payload())
    assert artifact["file_hash"] == sha256_bytes(source_bytes)
    assert artifact["file_size_bytes"] == len(source_bytes)
    source_rules = _rows(graph, "source_rules")
    assert source_rules
    assert all(item["sheet_name"] == "manual_json" for item in source_rules)
    assert all(item["cell_locator"].startswith("/criteria/") for item in source_rules)

    identity = _persist(db, prepared, actor_id=actor.id)
    compilation = db.get(models.RubricCompilation, identity["compilation_id"])
    assert compilation.parser_version == "manual-json-parser@1"
    assert compilation.human_changes == []
    assert db.scalar(select(func.count()).select_from(models.SourceArtifact)) == 1
    assert db.scalar(select(func.count()).select_from(models.AtomicRule)) >= 1


@requires_legacy_upgrade
def test_legacy_draft_upgrade_is_explicit_and_creates_blocked_full_provenance_graph(m4_db):
    db, actor, _ = m4_db
    legacy = models.Rubric(
        name="Legacy draft requiring upgrade",
        version="legacy-v1",
        total_score=10,
        status="draft",
        description="No provenance exists yet.",
        created_by=actor.id,
        owner_id=actor.id,
    )
    legacy.criteria.append(
        models.RubricCriterion(
            code="LEGACY_DIRECT",
            name="Legacy direct score",
            max_score=10,
            weight=None,
            criterion_type="llm_judgment",
            scoring_mode="llm_direct",
            applies_to="global",
            display_order=0,
        )
    )
    db.add(legacy)
    db.commit()
    legacy_id = legacy.id

    prepare = _PIPELINE.require("prepare_legacy_draft_upgrade")
    prepared = prepare(command=legacy_upgrade_command(rubric_id=legacy_id))
    graph = _mapping(prepared, label="legacy upgrade prepared graph")
    assert graph["source_kind"] == "legacy_draft_upgrade"
    assert "MISSING_EXECUTABLE_SCORING_MODE" in _blocker_codes(graph)
    assert all(
        item["scoring_mode"] != "llm_direct"
        for item in graph_scoring_projection(graph)
    )
    assert "legacy" in _artifact(graph, "legacy")["artifact_type"]

    identity = _persist(
        db,
        prepared,
        actor_id=actor.id,
        target_rubric_id=legacy_id,
    )
    assert identity["rubric_id"] == legacy_id
    assert db.scalar(select(func.count()).select_from(models.Rubric)) == 1
    assert db.get(models.Rubric, legacy_id).status == "draft"
    assert db.scalar(
        select(func.count())
        .select_from(models.RubricCompilation)
        .where(models.RubricCompilation.rubric_id == legacy_id)
    ) == 1
    assert db.scalar(select(func.count()).select_from(models.SourceArtifact)) >= 1
    assert db.scalar(select(func.count()).select_from(models.SourceRule)) >= 1
    assert db.scalar(select(func.count()).select_from(models.RubricVersion)) == 1
    assert db.scalar(select(func.count()).select_from(models.AtomicRule)) >= 1


def _assert_approval_event(
    event_payload: dict,
    *,
    rule_code: str,
    reviewer_id: str,
    reviewed_at: datetime,
):
    assert set(event_payload) == {
        "change_id",
        "rule_code",
        "field_path",
        "action",
        "before",
        "after",
        "actor_id",
        "occurred_at",
        "reason",
    }
    assert event_payload["change_id"]
    assert event_payload["rule_code"] == rule_code
    assert event_payload["field_path"] == f"/atomic_rules/{rule_code}/status"
    assert event_payload["action"] == "approve"
    assert event_payload["before"] == "review"
    assert event_payload["after"] == "approved"
    assert event_payload["actor_id"] == reviewer_id
    assert event_payload["reason"]
    assert event_payload["occurred_at"] == reviewed_at.isoformat(
        timespec="microseconds"
    )
    assert datetime.fromisoformat(event_payload["occurred_at"]).tzinfo is None


def _validate_required_fields_params(params):
    assert isinstance(params, Mapping)
    assert set(params) == {"required_fields"}
    assert params["required_fields"] == ["owner"]


def _required_fields_checker(*, document, params):
    del document, params
    return deterministic_missing_owner_observation()


class _SemanticRuntime:
    def __init__(self):
        self.calls = []

    def score(self, *, envelope):
        self.calls.append(envelope)
        return semantic_high_band_response()


@dataclass(frozen=True, slots=True)
class _TechnicalProfile:
    profile_key: str = PROFILE_KEY
    profile_version: str = PROFILE_VERSION
    prompt_version: str = "technical-proposal-prompt@1"

    def select_prompt_metadata(self, *, metadata):
        del metadata
        return {}

    def build_prompt_extensions(self, *, submission_snapshot, document_snapshot):
        del submission_snapshot, document_snapshot
        return {}


def _approve_publish_and_load_plan(
    db,
    *,
    identity: dict,
    actor_id: str,
    reviewer_id: str,
    expected_rule_codes: set[str],
    checker_registry=None,
):
    """Run the one ordinary signing chain shared by every source kind."""

    submit_rule = _LIFECYCLE.require("submit_atomic_rule_for_review")
    approve_rule = _LIFECYCLE.require("approve_atomic_rule")
    review_link = _LIFECYCLE.require("review_template_link")
    submit_rubric = _LIFECYCLE.require("submit_for_review")
    publish = _LIFECYCLE.require("publish_rubric")
    rules = db.scalars(
        select(models.AtomicRule)
        .where(models.AtomicRule.rubric_version_id == identity["rubric_version_id"])
        .order_by(models.AtomicRule.rule_code)
    ).all()
    assert {item.rule_code for item in rules} == expected_rule_codes
    compilation = db.get(models.RubricCompilation, identity["compilation_id"])
    assert compilation.status == "validated"
    assert compilation.blockers == []

    # Keep the replay clock relative to the persisted graph.  A fixed wall-clock
    # date eventually becomes earlier than the objects it is supposed to audit.
    created_times = [compilation.created_at, *(rule.created_at for rule in rules)]
    base_time = max(created_times) + timedelta(minutes=1)
    for offset, rule in enumerate(rules):
        submit_rule(
            db,
            identity["rubric_id"],
            rule.rule_code,
            actor_id,
            "提交规则进入统一审核链。",
            now=base_time + timedelta(minutes=offset),
        )
        approve_rule(
            db,
            identity["rubric_id"],
            rule.rule_code,
            reviewer_id,
            "确认来源、证据策略和计分语义。",
            now=base_time + timedelta(minutes=10 + offset),
        )

    links = db.scalars(
        select(models.RuleTemplateLink)
        .join(models.AtomicRule, models.RuleTemplateLink.rule_id == models.AtomicRule.id)
        .where(models.AtomicRule.rubric_version_id == identity["rubric_version_id"])
    ).all()
    for offset, link in enumerate(links):
        if link.review_status == "pending":
            review_link(
                db,
                identity["rubric_id"],
                link.id,
                reviewer_id,
                "confirmed",
                "确认模板来源映射。",
                now=base_time + timedelta(minutes=20 + offset),
            )
    db.commit()

    db.refresh(compilation)
    events = {
        item["rule_code"]: item
        for item in compilation.human_changes
        if item.get("action") == "approve"
    }
    assert set(events) == expected_rule_codes
    for rule in rules:
        db.refresh(rule)
        _assert_approval_event(
            events[rule.rule_code],
            rule_code=rule.rule_code,
            reviewer_id=reviewer_id,
            reviewed_at=rule.reviewed_at,
        )

    submit_rubric(db, identity["rubric_id"])
    db.commit()
    version = publish(
        db,
        identity["rubric_id"],
        identity["compilation_id"],
        reviewer_id,
        now=base_time + timedelta(hours=1),
    )
    db.commit()
    assert version.id == identity["rubric_version_id"]
    assert db.get(models.Rubric, identity["rubric_id"]).status == "published"

    snapshot = CompiledRubricSnapshotLoader().load_from_session(
        session=db,
        rubric_version_id=version.id,
        expected_profile_key=PROFILE_KEY,
    )
    registry = checker_registry or VersionedCheckerRegistry()
    plan = RuleExecutionPlanBuilder(
        checker_registry=registry,
        policy_compiler_version=POLICY_COMPILER_VERSION,
        engine_contract_version=ENGINE_CONTRACT_VERSION,
    ).build(
        rubric=snapshot,
        profile=_TechnicalProfile(),
        document_schema_version="document-snapshot@1",
    )
    plan_mapping = _mapping(plan, label="signed source execution plan")
    assert plan_mapping["rubric_version_id"] == version.id
    assert set(plan_mapping["dependency_order"]) == expected_rule_codes
    return version, rules, plan_mapping


@requires_manual_signing
def test_manual_json_valid_rubric_uses_the_ordinary_approve_publish_and_plan_chain(m4_db):
    db, actor, reviewer = m4_db
    prepare = _PIPELINE.require("prepare_manual_json_import")
    prepared = prepare(command=manual_import_command())
    graph = _mapping(prepared, label="manual JSON prepared graph")
    assert graph["source_kind"] == "manual_json"
    assert _artifact(graph, "manual_json")["artifact_type"] == "manual_json"
    assert _blocker_codes(graph) == set()
    rule_codes = {item["rule_code"] for item in _rows(graph, "atomic_rules")}
    assert rule_codes
    identity = _persist(db, prepared, actor_id=actor.id)

    version, rules, plan = _approve_publish_and_load_plan(
        db,
        identity=identity,
        actor_id=actor.id,
        reviewer_id=reviewer.id,
        expected_rule_codes=rule_codes,
    )

    assert version.compilation_id == identity["compilation_id"]
    assert {item.rule_code for item in rules} == rule_codes
    assert all(item.status == "approved" for item in rules)
    assert set(plan["dependency_order"]) == rule_codes
    artifact = db.scalar(
        select(models.SourceArtifact).where(
            models.SourceArtifact.compilation_id == identity["compilation_id"]
        )
    )
    assert artifact.artifact_type == "manual_json"
    assert artifact.source_rules


@requires_legacy_signing
def test_legacy_draft_cannot_publish_until_explicit_mapping_recompiles_the_same_chain(m4_db):
    db, actor, reviewer = m4_db
    legacy = models.Rubric(
        name="Legacy draft requiring upgrade",
        version="legacy-v1",
        total_score=10,
        status="draft",
        description="No provenance exists yet.",
        created_by=actor.id,
        owner_id=actor.id,
    )
    legacy.criteria.append(
        models.RubricCriterion(
            code="LEGACY_DIRECT",
            name="Legacy direct score",
            max_score=10,
            weight=None,
            criterion_type="llm_judgment",
            scoring_mode="llm_direct",
            applies_to="global",
            description="The old draft omitted band and deduct rules.",
            dimension="content",
            display_order=0,
        )
    )
    db.add(legacy)
    db.commit()
    legacy_id = legacy.id

    prepare = _PIPELINE.require("prepare_legacy_draft_upgrade")
    blocked_prepared = prepare(
        command=legacy_upgrade_command(rubric_id=legacy_id)
    )
    blocked_graph = _mapping(blocked_prepared, label="blocked legacy upgrade")
    assert _artifact(blocked_graph, "legacy")["artifact_type"] == "legacy_draft"
    assert "MISSING_EXECUTABLE_SCORING_MODE" in _blocker_codes(blocked_graph)
    blocked_identity = _persist(
        db,
        blocked_prepared,
        actor_id=actor.id,
        target_rubric_id=legacy_id,
    )

    submit_rubric = _LIFECYCLE.require("submit_for_review")
    publish = _LIFECYCLE.require("publish_rubric")
    lifecycle_error = getattr(
        _LIFECYCLE.module,
        "RubricLifecycleError",
        ValueError,
    )
    blocked_error = None
    try:
        submit_rubric(db, legacy_id)
        db.commit()
    except lifecycle_error as exc:
        db.rollback()
        blocked_error = exc
    if blocked_error is None:
        try:
            publish(
                db,
                legacy_id,
                blocked_identity["compilation_id"],
                reviewer.id,
                now=(
                    db.get(
                        models.RubricCompilation,
                        blocked_identity["compilation_id"],
                    ).created_at
                    + timedelta(minutes=1)
                ),
            )
            db.commit()
        except lifecycle_error as exc:
            db.rollback()
            blocked_error = exc
    assert blocked_error is not None, "unresolved llm_direct legacy graph must not publish"
    assert db.get(models.Rubric, legacy_id).status != "published"

    if db.get(models.Rubric, legacy_id).status == "review":
        return_to_draft = _LIFECYCLE.require("return_to_draft")
        return_to_draft(db, legacy_id)
        db.commit()

    executable_command = legacy_upgrade_command(
        rubric_id=legacy_id,
        explicit_mapping=True,
    )
    executable_command["draft_recompile"] = {
        "mode": "supersede_unpublished",
        "supersedes_compilation_id": blocked_identity["compilation_id"],
    }
    executable_prepared = prepare(command=executable_command)
    executable_graph = _mapping(
        executable_prepared,
        label="explicitly mapped legacy upgrade",
    )
    assert executable_graph["source_kind"] == "legacy_draft_upgrade"
    assert _artifact(executable_graph, "legacy")["artifact_type"] == "legacy_draft"
    assert _blocker_codes(executable_graph) == set()
    assert {
        item["scoring_mode"] for item in graph_scoring_projection(executable_graph)
    } == {"banded"}
    executable_rule_codes = {
        item["rule_code"] for item in _rows(executable_graph, "atomic_rules")
    }
    assert executable_rule_codes
    executable_identity = _persist(
        db,
        executable_prepared,
        actor_id=actor.id,
        target_rubric_id=legacy_id,
    )
    assert executable_identity["compilation_id"] != blocked_identity["compilation_id"]

    version, rules, plan = _approve_publish_and_load_plan(
        db,
        identity=executable_identity,
        actor_id=actor.id,
        reviewer_id=reviewer.id,
        expected_rule_codes=executable_rule_codes,
    )
    assert version.compilation_id == executable_identity["compilation_id"]
    assert {item.rule_code for item in rules} == executable_rule_codes
    assert set(plan["dependency_order"]) == executable_rule_codes
    blocked_compilation = db.get(
        models.RubricCompilation,
        blocked_identity["compilation_id"],
    )
    assert blocked_compilation.blockers, "superseded draft remains auditable"


@requires_vertical
def test_real_excel_word_import_reviews_publishes_plans_and_scores_once(
    m4_db,
    monkeypatch,
):
    execute_rule_plan = _RULE_EXECUTOR.require("execute_rule_plan")
    db, actor, reviewer = m4_db
    prepared, _ = _prepare_file()
    identity = _persist(db, prepared, actor_id=actor.id)
    rules = db.scalars(
        select(models.AtomicRule)
        .where(models.AtomicRule.rubric_version_id == identity["rubric_version_id"])
        .order_by(models.AtomicRule.rule_code)
    ).all()
    assert {item.rule_code for item in rules} == {RISK_RULE_CODE, FIT_RULE_CODE}

    submit_rule = _LIFECYCLE.require("submit_atomic_rule_for_review")
    approve_rule = _LIFECYCLE.require("approve_atomic_rule")
    review_link = _LIFECYCLE.require("review_template_link")
    submit_rubric = _LIFECYCLE.require("submit_for_review")
    publish = _LIFECYCLE.require("publish_rubric")
    compilation = db.get(models.RubricCompilation, identity["compilation_id"])
    reviewed_at = max(
        [compilation.created_at, *(rule.created_at for rule in rules)]
    ) + timedelta(minutes=11)
    for offset, rule in enumerate(rules):
        submit_rule(
            db,
            identity["rubric_id"],
            rule.rule_code,
            actor.id,
            "提交原子规则供独立审核。",
            now=reviewed_at - timedelta(minutes=10 - offset),
        )
        approve_rule(
            db,
            identity["rubric_id"],
            rule.rule_code,
            reviewer.id,
            "来源与执行语义已人工核对。",
            now=reviewed_at + timedelta(minutes=offset),
        )

    links = db.scalars(
        select(models.RuleTemplateLink)
        .join(models.AtomicRule, models.RuleTemplateLink.rule_id == models.AtomicRule.id)
        .where(models.AtomicRule.rubric_version_id == identity["rubric_version_id"])
    ).all()
    for offset, link in enumerate(links, start=10):
        review_link(
            db,
            identity["rubric_id"],
            link.id,
            reviewer.id,
            "confirmed",
            "批注与原子规则具有直接约束关系。",
            now=reviewed_at + timedelta(minutes=offset),
        )
    db.commit()

    compilation = db.get(models.RubricCompilation, identity["compilation_id"])
    approval_events = [
        item for item in compilation.human_changes if item.get("action") == "approve"
    ]
    assert len(approval_events) == len(rules)
    by_rule = {item["rule_code"]: item for item in approval_events}
    for rule in rules:
        db.refresh(rule)
        _assert_approval_event(
            by_rule[rule.rule_code],
            rule_code=rule.rule_code,
            reviewer_id=reviewer.id,
            reviewed_at=rule.reviewed_at,
        )
        assert rule.status == "approved"
        assert rule.reviewed_by == reviewer.id
        assert rule.reviewed_at is not None

    submit_rubric(db, identity["rubric_id"])
    db.commit()
    version = publish(
        db,
        identity["rubric_id"],
        identity["compilation_id"],
        reviewer.id,
        now=reviewed_at + timedelta(hours=1),
    )
    db.commit()

    snapshot = CompiledRubricSnapshotLoader().load_from_session(
        session=db,
        rubric_version_id=version.id,
        expected_profile_key=PROFILE_KEY,
    )
    registry = VersionedCheckerRegistry()
    registry.register(
        registration=checker_registration_payload(),
        checker=_required_fields_checker,
        validate_params=_validate_required_fields_params,
    )
    plan = RuleExecutionPlanBuilder(
        checker_registry=registry,
        policy_compiler_version=POLICY_COMPILER_VERSION,
        engine_contract_version=ENGINE_CONTRACT_VERSION,
    ).build(
        rubric=snapshot,
        profile=_TechnicalProfile(),
        document_schema_version="document-snapshot@1",
    )
    plan_mapping = _mapping(plan, label="M4 execution plan")
    assert plan_mapping["rubric_version_id"] == version.id
    assert plan_mapping["rubric_version_hash"] == version.version_hash
    assert set(plan_mapping["dependency_order"]) == {RISK_RULE_CODE, FIT_RULE_CODE}
    nodes = {item["rule_code"]: item for item in plan_mapping["nodes"]}
    assert nodes[RISK_RULE_CODE]["atomic_rule_snapshot"]["checker_version"] == CHECKER_VERSION
    assert nodes[FIT_RULE_CODE]["atomic_rule_snapshot"]["levels"]
    assert all(rule.source_rules for rule in rules)

    # Public scoring keeps its existing four ports, but M4 must delegate the
    # actual plan exactly once to the separately testable executor entrypoint.
    executor_module = _RULE_EXECUTOR.module
    assert executor_module is not None
    engine_module = importlib.import_module("backend.app.services.scoring.core.engine")
    delegated = Mock(wraps=execute_rule_plan)
    monkeypatch.setattr(executor_module, "execute_rule_plan", delegated)
    for name, value in tuple(vars(engine_module).items()):
        if value is execute_rule_plan:
            monkeypatch.setattr(engine_module, name, delegated)

    request = scoring_request_payload()
    request["plan"] = plan_mapping
    request["idempotency_key"] = canonical_sha256(idempotency_projection(request))
    runtime = _SemanticRuntime()
    outcome = engine_module.score_submission(
        request=request,
        checker_registry=registry,
        llm_runtime=runtime,
        profile=_TechnicalProfile(),
    )
    result = _mapping(outcome, label="M4 scoring outcome")
    assert delegated.call_count == 1
    assert len(runtime.calls) == 1
    assert result["schema_version"] == "scoring-outcome@2"
    assert result["status"] == "completed"
    assert result["final_total"] == "90"
    assert {item["rule_code"] for item in result["rule_decisions"]} == {
        RISK_RULE_CODE,
        FIT_RULE_CODE,
    }

    risk_decision = next(
        item
        for item in result["rule_decisions"]
        if item["rule_code"] == RISK_RULE_CODE
    )
    assert risk_decision["status"] == "triggered"
    assert risk_decision["version_hash"] == version.version_hash
    assert risk_decision["calculated_effect"] == "-10"
    risk_contribution = next(
        item
        for item in result["score_contributions"]
        if item["rule_code"] == RISK_RULE_CODE and item["kind"] == "deduction"
    )
    assert risk_contribution["amount"] == "-10"

    # Close the audit chain from the actual applied deduction, rather than
    # merely proving that unrelated provenance rows exist somewhere.
    risk_rule = next(rule for rule in rules if rule.rule_code == RISK_RULE_CODE)
    assert risk_rule.rubric_version_id == version.id
    assert len(risk_rule.source_rules) == 1
    excel_source = risk_rule.source_rules[0]
    assert excel_source.sheet_name == "Atomic Rules"
    assert excel_source.row_number == 2
    assert "A2" in excel_source.cell_locator
    confirmed_comment_links = [
        link
        for link in risk_rule.template_links
        if link.review_status == "confirmed"
        and link.template_item.kind == "comment"
        and link.template_item.source_locator.get("comment_id") == COMMENT_ID
    ]
    assert len(confirmed_comment_links) == 1
    template_item = confirmed_comment_links[0].template_item
    assert template_item.raw_text == COMMENT_TEXT
    assert template_item.source_locator["anchor_text"] == COMMENT_ANCHOR
    assert "风险控制" in template_item.section_path
