from io import BytesIO

import pytest
from openpyxl import Workbook

from backend.app.services.rubric_import.coverage import compute_coverage
from backend.app.services.rubric_import.extraction.table_extractor import extract_table
from backend.app.services.rubric_import.sources.units import SourceLedger
from backend.app.services.rubric_import.sources.xlsx_adapter import load_xlsx
from backend.app.tests import rubric_parse_fixtures as fx


def _extract(data):
    ledger = SourceLedger()
    sheets = load_xlsx(data, ledger, doc_id="rules", doc_role="rules")
    return extract_table(sheets, ledger), ledger


def test_simple_table_is_fully_claimed():
    result, ledger = _extract(fx.simple_rules_xlsx())
    assert result.sheet_title == "评分规则"
    assert [row.criterion.name for row in result.records] == ["研究方法", "文献综述"]
    assert [row.row_number for row in result.records] == [2, 3]
    assert ledger.status("xlsx:评分规则!R1C2").status == "structural"
    assert ledger.status("xlsx:评分规则!R2C3").claimed_by == ("C01.max_score",)
    assert compute_coverage(ledger)["documents"][0]["ratio"] == 1.0
    assert result.unmapped_columns == []


def test_complex_table_reports_losses():
    result, ledger = _extract(fx.complex_rules_xlsx())
    assert [row.criterion.name for row in result.records] == ["选题意义", "文献综述", "方法设计"]
    # 合并的维度单元被两个评分项认领
    assert ledger.status("xlsx:评分表!R4C1").claimed_by == ("C01.dimension", "C02.dimension")
    # 未映射列“备注”的内容未被认领
    assert ledger.status("xlsx:评分表!R5C6").status == "unclaimed"
    # R5 的备注 + R8 整行合并（A:F）覆盖到备注列
    assert result.unmapped_columns == [{"column": 6, "header": "备注", "non_empty": 2, "data_rows": 5}]
    # 无满分行与整行合并的说明行被丢弃并记录
    assert [(row["row_number"], row["reason"]) for row in result.dropped_rows] == [
        (7, "未找到满分"), (8, "整行合并的说明行")]
    assert ledger.status("xlsx:评分表!R8C1").status == "unclaimed"
    assert ledger.status("xlsx:评分表!R7C2").status == "unclaimed"
    # 合计行为结构性内容，并给出声明总分
    assert ledger.status("xlsx:评分表!R9C1").status == "structural"
    assert result.total_row == {"row_number": 9, "declared_total": 50.0}
    # 表头上方的标题未被认领；其他工作表按规则忽略
    assert ledger.status("xlsx:评分表!R1C1").status == "unclaimed"
    assert ledger.status("xlsx:附加扣分!R2C3").status == "ignored_by_rule"
    assert result.ignored_sheets == [{"title": "附加扣分", "reason": "仅解析第一张识别到评分项的工作表"}]


def test_sheet_without_header_before_rules_sheet_is_ignored_with_warning():
    workbook = Workbook()
    workbook.active.title = "说明"
    workbook.active.append(["本表用于毕业论文评分"])
    sheet = workbook.create_sheet("规则")
    sheet.append(["评分项", "分值"])
    sheet.append(["选题", 10])
    buffer = BytesIO()
    workbook.save(buffer)
    result, ledger = _extract(buffer.getvalue())
    assert result.sheet_title == "规则"
    assert result.warnings == ["工作表 说明 未识别到评分规则表头，已跳过。"]
    assert ledger.status("xlsx:说明!R1C1").status == "ignored_by_rule"


def test_no_rules_sheet_raises_same_error_as_before():
    with pytest.raises(ValueError, match="Excel 未解析到有效评分项"):
        _extract(fx.unmapped_header_xlsx())


def test_custom_header_aliases_are_supported():
    ledger = SourceLedger()
    sheets = load_xlsx(fx.unmapped_header_xlsx(), ledger, doc_id="rules", doc_role="rules")
    result = extract_table(sheets, ledger, header_aliases={"name": ["考核内容"], "max_score": ["配分"]})
    assert [row.criterion.name for row in result.records] == ["选题"]
    assert result.unmapped_columns == [{"column": 3, "header": "细则", "non_empty": 1, "data_rows": 1}]


def test_header_aliases_are_resolved_per_profile_with_thesis_defaults():
    from backend.app.services.rubric_import.extraction.table_extractor import HEADER_ALIASES
    from backend.app.services.rubric_import.extraction.table_extractor import PROFILE_HEADER_ALIASES
    from backend.app.services.rubric_import.extraction.table_extractor import header_aliases_for

    assert header_aliases_for("thesis") == HEADER_ALIASES
    assert header_aliases_for(None) == HEADER_ALIASES
    PROFILE_HEADER_ALIASES["demo_profile"] = {"max_score": ["配分"]}
    try:
        merged = header_aliases_for("demo_profile")
        assert merged["max_score"][-1] == "配分"
        assert merged["name"] == HEADER_ALIASES["name"]
        assert HEADER_ALIASES["max_score"][-1] != "配分"  # 默认表不被修改
    finally:
        PROFILE_HEADER_ALIASES.pop("demo_profile")


def test_merged_parent_name_is_reinterpreted_as_dimension():
    result, ledger = _extract(fx.merged_parent_dimension_xlsx())

    assert result.mapping == {"item_label": 0, "description": 2, "dimension": 1}
    assert [row.criterion.code for row in result.records] == ["T02", "T03", "T04", "T05"]
    assert [row.criterion.name for row in result.records] == [
        "分析与解决问题1", "分析与解决问题2", "分析与解决问题3", "分析与解决问题4",
    ]
    assert {row.criterion.dimension for row in result.records} == {"分析与解决问题"}
    assert [row.criterion.max_score for row in result.records] == [20.0, 20.0, 10.0, 10.0]
    assert ledger.status("xlsx:Sheet!R2C2").claimed_by == (
        "T02.dimension", "T03.dimension", "T04.dimension", "T05.dimension",
    )
    assert result.structure_issues == []
    assert result.warnings == ["检测到纵向合并的父级评价项目，已按评分维度解析。"]


def test_independent_duplicate_names_are_not_reinterpreted_as_dimension():
    result, _ = _extract(fx.merged_parent_dimension_xlsx(independently_repeated=True))

    assert result.mapping["name"] == 1
    assert "dimension" not in result.mapping
    assert {row.criterion.name for row in result.records} == {"分析与解决问题"}
    assert result.structure_issues == []


def test_merged_name_does_not_overwrite_existing_dimension_mapping():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["打分项", "评价项目", "评分维度", "具体要求"])
    sheet.append(["指导教师成绩项2（20分）", "分析与解决问题", "专业能力", "要求A"])
    sheet.append(["指导教师成绩项3（10分）", "分析与解决问题", "专业能力", "要求B"])
    sheet.merge_cells("B2:B3")
    buffer = BytesIO()
    workbook.save(buffer)

    result, _ = _extract(buffer.getvalue())

    assert result.mapping == {"item_label": 0, "name": 1, "dimension": 2, "description": 3}
    assert {row.criterion.name for row in result.records} == {"分析与解决问题"}
    assert {row.criterion.dimension for row in result.records} == {"专业能力"}
    assert result.structure_issues[0]["code"] == "MERGED_NAME_AMBIGUOUS"
    assert result.structure_issues[0]["reason"] == "已经存在维度列，无法自动重新解释合并名称列。"
