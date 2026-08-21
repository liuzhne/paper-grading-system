"""Independent M4 import/compilation fixtures.

The factories in this module deliberately do not import the M4 pipeline.  Raw
file hashes, canonical manual-JSON bytes and the expected provenance locators
are derived here so the production importer cannot become its own test oracle.
"""

from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
from openpyxl import Workbook

from backend.app.services.scoring.core.canonical import canonical_sha256
from backend.app.tests.m2_contract_fixtures import (
    PROFILE_KEY,
    technical_policy_snapshot_payload,
)
from backend.app.tests.m3_contract_fixtures import CHECKER_KEY


# M4 extends the legacy import DTO; it must not silently mutate the frozen
# pre-provenance shape in place.
IMPORT_SCHEMA_VERSION = "rubric-import-command@2"
PREPARED_SCHEMA_VERSION = "prepared-rubric-graph@1"
PARSER_VERSION = "rubric-file-parser@2"
COMPILER_VERSION = "atomic-rule-compiler@1"
PROMPT_VERSION = "rubric-compilation-prompt@2"
# Workflow identity is separate from the business profile.  Reuse the key
# frozen by the M3 graph and the M4 publication registry fixtures.
WORKFLOW_PROFILE = "template_driven"
HASH_SCHEME = "rubric-content-v2"

RISK_RULE_CODE = "proposal.risk_owner.v1"
FIT_RULE_CODE = "proposal.solution_fit.v1"
COMMENT_ID = "7"
COMMENT_ANCHOR = "每项风险必须明确责任人"
COMMENT_TEXT = "缺少责任人时应按风险控制规则扣分。"

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_CONTENT_TYPES = "http://schemas.openxmlformats.org/package/2006/content-types"
_RELATIONSHIPS = "http://schemas.openxmlformats.org/package/2006/relationships"


def sha256_bytes(value: bytes) -> str:
    return sha256(value).hexdigest()


def canonical_json_bytes(value) -> bytes:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")


def _q(tag: str) -> str:
    return "{%s}%s" % (_W, tag)


def real_rules_xlsx_bytes() -> bytes:
    """Two executable leaves: one deterministic deduct and one semantic band."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Atomic Rules"
    sheet.append(
        [
            "原子规则编号",
            "评分项编号",
            "评分项",
            "分值",
            "权重",
            "类型",
            "评分模式",
            "评分说明",
            "证据提示",
            "扣分规则",
            "适用范围",
            "分档",
            "维度",
            "checker_key",
            "checker_params",
            "evidence_policy",
            "effect_type",
            "repeat_policy",
            "strictness",
        ]
    )
    sheet.append(
        [
            RISK_RULE_CODE,
            "RISK_CONTROL",
            "Risk control completeness",
            20,
            None,
            "deterministic",
            "deductive",
            "每项风险均应明确责任人。",
            "总体方案；风险控制",
            "缺少责任人扣10分",
            "风险控制",
            None,
            "content",
            CHECKER_KEY,
            json.dumps({"required_fields": ["owner"]}),
            json.dumps(
                {
                    "mode": "scoped_absence",
                    "requirement": "required",
                    "minimum_coverage": "1",
                }
            ),
            "score",
            "once",
            "required",
        ]
    )
    sheet.append(
        [
            FIT_RULE_CODE,
            "SOLUTION_FIT",
            "Solution fit",
            80,
            None,
            "semantic",
            "banded",
            "方案应明确回应需求。",
            "需求分析；总体方案",
            None,
            "总体方案",
            "FIT_HIGH:80;FIT_LOW:40",
            "content",
            None,
            "{}",
            json.dumps(
                {
                    "mode": "source_quote",
                    "requirement": "required",
                    "minimum_coverage": "1",
                }
            ),
            "score",
            None,
            "required",
        ]
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def real_template_docx_bytes() -> bytes:
    """Build a real DOCX with an anchored Word comment.

    ``python-docx`` currently has no stable cross-version comment writer, so
    the fixture adds the standard OOXML comment part directly.  The production
    parser sees a normal DOCX zip rather than a mocked parser structure.
    """

    document = Document()
    document.add_heading("总体方案", level=1)
    document.add_heading("风险控制", level=2)
    document.add_paragraph(COMMENT_ANCHOR)
    document.add_paragraph("方案必须对应需求并说明取舍依据。")
    original = BytesIO()
    document.save(original)

    source = zipfile.ZipFile(BytesIO(original.getvalue()), "r")
    entries = {name: source.read(name) for name in source.namelist()}
    source.close()

    document_root = ET.fromstring(entries["word/document.xml"])
    target = None
    for paragraph in document_root.iter(_q("p")):
        text = "".join(node.text or "" for node in paragraph.iter(_q("t")))
        if text == COMMENT_ANCHOR:
            target = paragraph
            break
    assert target is not None
    first_run = next(node for node in list(target) if node.tag == _q("r"))
    first_index = list(target).index(first_run)
    target.insert(first_index, ET.Element(_q("commentRangeStart"), {_q("id"): COMMENT_ID}))
    target.insert(first_index + 2, ET.Element(_q("commentRangeEnd"), {_q("id"): COMMENT_ID}))
    reference_run = ET.Element(_q("r"))
    ET.SubElement(reference_run, _q("commentReference"), {_q("id"): COMMENT_ID})
    target.insert(first_index + 3, reference_run)
    entries["word/document.xml"] = ET.tostring(
        document_root,
        encoding="utf-8",
        xml_declaration=True,
    )

    comments = ET.Element(_q("comments"))
    comment = ET.SubElement(
        comments,
        _q("comment"),
        {_q("id"): COMMENT_ID, _q("author"): "M4 reviewer"},
    )
    paragraph = ET.SubElement(comment, _q("p"))
    run = ET.SubElement(paragraph, _q("r"))
    text = ET.SubElement(run, _q("t"))
    text.text = COMMENT_TEXT
    entries["word/comments.xml"] = ET.tostring(
        comments,
        encoding="utf-8",
        xml_declaration=True,
    )

    content_types = ET.fromstring(entries["[Content_Types].xml"])
    override_tag = "{%s}Override" % _CONTENT_TYPES
    if not any(
        item.get("PartName") == "/word/comments.xml"
        for item in content_types.findall(override_tag)
    ):
        ET.SubElement(
            content_types,
            override_tag,
            {
                "PartName": "/word/comments.xml",
                "ContentType": (
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.comments+xml"
                ),
            },
        )
    entries["[Content_Types].xml"] = ET.tostring(
        content_types,
        encoding="utf-8",
        xml_declaration=True,
    )

    relationships = ET.fromstring(entries["word/_rels/document.xml.rels"])
    relationship_tag = "{%s}Relationship" % _RELATIONSHIPS
    if not any(
        item.get("Type", "").endswith("/comments")
        for item in relationships.findall(relationship_tag)
    ):
        ET.SubElement(
            relationships,
            relationship_tag,
            {
                "Id": "rIdM4Comments",
                "Type": (
                    "http://schemas.openxmlformats.org/officeDocument/2006/"
                    "relationships/comments"
                ),
                "Target": "comments.xml",
            },
        )
    entries["word/_rels/document.xml.rels"] = ET.tostring(
        relationships,
        encoding="utf-8",
        xml_declaration=True,
    )

    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    return output.getvalue()


def inspect_word_comment_without_production_parser(docx_bytes: bytes) -> dict:
    """Extract the fixture's comment with only ZIP + ElementTree."""

    with zipfile.ZipFile(BytesIO(docx_bytes)) as archive:
        document = ET.fromstring(archive.read("word/document.xml"))
        comments = ET.fromstring(archive.read("word/comments.xml"))
    comment = next(
        node
        for node in comments.iter(_q("comment"))
        if node.get(_q("id")) == COMMENT_ID
    )
    comment_text = "".join(node.text or "" for node in comment.iter(_q("t")))
    anchor = ""
    open_comment = False
    for node in document.iter():
        if node.tag == _q("commentRangeStart") and node.get(_q("id")) == COMMENT_ID:
            open_comment = True
        elif node.tag == _q("commentRangeEnd") and node.get(_q("id")) == COMMENT_ID:
            open_comment = False
        elif node.tag == _q("t") and open_comment:
            anchor += node.text or ""
    return {
        "comment_id": COMMENT_ID,
        "comment_text": comment_text,
        "anchor_text": anchor,
    }


def file_import_command() -> dict:
    return {
        "schema_version": IMPORT_SCHEMA_VERSION,
        "source_kind": "file_import",
        "rubric": {
            "name": "M4 technical proposal rubric",
            "version": "v1.0",
            "description": "Imported from the audited M4 fixture.",
            "business_profile_key": PROFILE_KEY,
            "workflow_profile": WORKFLOW_PROFILE,
            "global_policy": technical_policy_snapshot_payload(
                total_score="100",
                rounding_digits=2,
            ),
        },
        "files": {
            "rules_file_name": "m4-rules.xlsx",
            "template_file_name": "m4-template.docx",
        },
        "compiler": {
            "parser_version": PARSER_VERSION,
            "compiler_version": COMPILER_VERSION,
            "prompt_version": PROMPT_VERSION,
            "model_provider": None,
            "model_name": None,
            "sampling_params": {},
        },
        "version": {
            "hash_scheme": HASH_SCHEME,
        },
    }


def ambiguous_rules_xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "评分规则"
    sheet.append(["编号", "评分项", "分值", "评分说明", "扣分规则"])
    sheet.append(
        [
            "AMBIGUOUS",
            "论证质量",
            10,
            "论证应充分并有证据支持。",
            "论证薄弱时适当扣分",
        ]
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def hybrid_rules_xlsx_bytes(
    *,
    parent_weight: str | None = None,
    first_points: str = "8",
    second_points: str = "12",
) -> bytes:
    total = str(int(first_points) + int(second_points))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "评分规则"
    sheet.append(
        [
            "编号",
            "评分项",
            "分值",
            "权重",
            "类型",
            "评分模式",
            "子检查",
        ]
    )
    sheet.append(
        [
            "IMPLEMENTATION",
            "实施方案",
            total,
            parent_weight,
            "hybrid",
            "hybrid",
            "里程碑完整性 | deterministic | %s\n实施可行性 | semantic | %s"
            % (first_points, second_points),
        ]
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def hybrid_import_command(*, weighted: bool = False) -> dict:
    command = file_import_command()
    command["rubric"] = {
        **command["rubric"],
        "name": "M4 hybrid rubric weighted" if weighted else "M4 hybrid rubric points",
        "global_policy": (
            weighted_policy_payload(total_score="20")
            if weighted
            else technical_policy_snapshot_payload(
                total_score="20",
                rounding_digits=2,
            )
        ),
    }
    command["files"]["rules_file_name"] = "m4-hybrid.xlsx"
    command["files"]["template_file_name"] = None
    return command


def weighted_policy_payload(*, total_score: str) -> dict:
    policy = technical_policy_snapshot_payload(
        total_score=total_score,
        rounding_digits=2,
    )
    policy["aggregation"]["mode"] = "weighted_normalized"
    policy.pop("policy_hash")
    policy["policy_hash"] = canonical_sha256(policy)
    return policy


def manual_rubric_payload() -> dict:
    return {
        "name": "M4 manual rubric",
        "version": "v1.0",
        "description": "Created through the public JSON API.",
        "total_score": "10",
        "business_profile_key": PROFILE_KEY,
        "workflow_profile": WORKFLOW_PROFILE,
        "global_policy": technical_policy_snapshot_payload(
            total_score="10",
            rounding_digits=2,
        ),
        "criteria": [
            {
                "code": "MANUAL_FIT",
                "name": "Manual fit",
                "max_score": "10",
                "weight": None,
                "criterion_type": "llm_judgment",
                "scoring_mode": "banded",
                "applies_to": "global",
                "levels": [
                    {"code": "HIGH", "points": "10", "descriptor": "充分"},
                    {"code": "LOW", "points": "5", "descriptor": "部分充分"},
                ],
            }
        ],
    }


def manual_import_command() -> dict:
    return {
        "schema_version": IMPORT_SCHEMA_VERSION,
        "source_kind": "manual_json",
        "rubric": manual_rubric_payload(),
        "compiler": {
            "parser_version": "manual-json-parser@1",
            "compiler_version": COMPILER_VERSION,
            "prompt_version": "manual-json-no-llm@1",
            "model_provider": None,
            "model_name": None,
            "sampling_params": {},
        },
        "version": {"hash_scheme": HASH_SCHEME},
    }


def legacy_draft_payload(*, rubric_id: str) -> dict:
    return {
        "schema_version": "legacy-rubric-draft@1",
        "rubric_id": rubric_id,
        "name": "Legacy draft requiring upgrade",
        "version": "legacy-v1",
        "description": "No provenance exists yet.",
        "total_score": "10",
        "status": "draft",
        "criteria": [
            {
                "code": "LEGACY_DIRECT",
                "name": "Legacy direct score",
                "max_score": "10",
                "weight": None,
                "criterion_type": "llm_judgment",
                "scoring_mode": "llm_direct",
                "applies_to": "global",
                "description": "The old draft omitted band and deduct rules.",
                "evidence_hints": [],
                "deduction_rules": [],
                "rubric_levels": [],
                "sub_checks": [],
                "dimension": "content",
                "deduction_rules_structured": [],
            }
        ],
    }


def legacy_upgrade_command(
    *,
    rubric_id: str,
    explicit_mapping: bool = False,
) -> dict:
    command = {
        "schema_version": IMPORT_SCHEMA_VERSION,
        "source_kind": "legacy_draft_upgrade",
        "legacy_rubric": legacy_draft_payload(rubric_id=rubric_id),
        "rubric": {
            "business_profile_key": PROFILE_KEY,
            "workflow_profile": WORKFLOW_PROFILE,
            "global_policy": technical_policy_snapshot_payload(
                total_score="10",
                rounding_digits=2,
            ),
        },
        "compiler": {
            "parser_version": "legacy-draft-upgrader@1",
            "compiler_version": COMPILER_VERSION,
            "prompt_version": "legacy-upgrade-no-llm@1",
            "model_provider": None,
            "model_name": None,
            "sampling_params": {},
        },
        "version": {
            "hash_scheme": HASH_SCHEME,
            "version": (
                "legacy-v1-executable-v2"
                if explicit_mapping
                else "legacy-v1-blocked-v1"
            ),
        },
    }
    if explicit_mapping:
        command["upgrade_mapping"] = {
            "LEGACY_DIRECT": {
                "strategy": "band",
                "criterion_type": "llm_judgment",
                "scoring_mode": "banded",
                "judge_type": "semantic",
                "direction": "band",
                "effect_type": "score",
                "evidence_policy": {
                    "mode": "source_quote",
                    "requirement": "required",
                    "minimum_coverage": "1",
                },
                "levels": [
                    {
                        "level_code": "HIGH",
                        "points": "10",
                        "descriptor": "证据充分且论证完整。",
                        "display_order": 0,
                    },
                    {
                        "level_code": "LOW",
                        "points": "5",
                        "descriptor": "证据或论证仅部分充分。",
                        "display_order": 1,
                    },
                ],
            }
        }
    return command


def clone(value):
    return deepcopy(value)


__all__ = [
    "COMMENT_ANCHOR",
    "COMMENT_ID",
    "COMMENT_TEXT",
    "COMPILER_VERSION",
    "FIT_RULE_CODE",
    "HASH_SCHEME",
    "IMPORT_SCHEMA_VERSION",
    "PARSER_VERSION",
    "PREPARED_SCHEMA_VERSION",
    "PROMPT_VERSION",
    "RISK_RULE_CODE",
    "WORKFLOW_PROFILE",
    "ambiguous_rules_xlsx_bytes",
    "canonical_json_bytes",
    "clone",
    "file_import_command",
    "hybrid_import_command",
    "hybrid_rules_xlsx_bytes",
    "inspect_word_comment_without_production_parser",
    "legacy_draft_payload",
    "legacy_upgrade_command",
    "manual_import_command",
    "manual_rubric_payload",
    "real_rules_xlsx_bytes",
    "real_template_docx_bytes",
    "sha256_bytes",
    "weighted_policy_payload",
]
