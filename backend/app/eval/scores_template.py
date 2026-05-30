"""生成"教师成绩表"模板（.xlsx），列随 rubric 的评分项 code 自动展开，方便老师照填。

数据表只放表头（供直接填写）；列含义/示例/提示放在"填写说明"页，避免示例行被当成数据解析。
与 `eval/labeled_dataset.load_scores_table` 的列约定一致。
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.styles import Font
from openpyxl.styles import PatternFill


def build_scores_table_template(criteria):
    """criteria: 可迭代的评分项（对象或 dict，含 code/name/max_score）。返回 .xlsx 字节。"""
    rows = [(_attr(c, "code"), _attr(c, "name"), _attr(c, "max_score")) for c in criteria]
    headers = ["文件名", "总分"] + [str(code) for code, _, _ in rows]
    total_max = sum(float(max_score or 0) for _, _, max_score in rows) or 100

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "教师评分"
    sheet.append(headers)
    _style_header(sheet, len(headers))
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 10

    info = workbook.create_sheet("填写说明")
    info.append(["列", "含义"])
    info.append(["文件名", "与论文 .docx 文件名完全一致（含扩展名）"])
    info.append(["总分", "教师给的总分（本标准满分 %g）" % total_max])
    for code, name, max_score in rows:
        info.append([str(code), "评分项「%s」教师分（满分 %g）" % (name, float(max_score or 0))])
    info.append(["", ""])
    info.append(["示例", "文件名=张三.docx, 总分=86, 各 code 列填对应分项分"])
    info.append(["建议", "覆盖各分数段（优/良/中/及格/不及格），总量≥10，QWK 才稳定"])
    info.append(["注意", "分项 code 必须与评分标准一致；分项分可留空（则只算总分 QWK）"])
    _style_header(info, 2)
    info.column_dimensions["A"].width = 14
    info.column_dimensions["B"].width = 60
    for row in info.iter_rows(min_row=2, max_row=info.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _attr(criterion, name):
    if isinstance(criterion, dict):
        return criterion.get(name)
    return getattr(criterion, name, None)


def _style_header(sheet, columns):
    fill = PatternFill("solid", fgColor="DCEBFF")
    sheet.freeze_panes = "A2"
    for index in range(1, columns + 1):
        cell = sheet.cell(row=1, column=index)
        cell.fill = fill
        cell.font = Font(bold=True, color="17202A")
        cell.alignment = Alignment(horizontal="center", vertical="center")
