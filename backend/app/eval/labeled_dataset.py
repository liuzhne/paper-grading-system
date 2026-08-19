"""QWK 留出集：从"论文文件夹 + 教师成绩表"构建评估，比对系统分与真实人工分（设计§15）。

成绩表（.xlsx 或 .csv）列约定：
- 文件名列（表头含 文件名/文件/论文/filename/file/paper）
- 总分列（表头含 总分/总成绩/total/score）
- 其余表头 = 评分项 code 的列 → 视为该项的人工分（分项分，用于逐维度 bias）
"""

import csv
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from backend.app.db.models import GradingBatch
from backend.app.db.models import Rubric
from backend.app.db.models import RubricCompilation
from backend.app.db.models import RubricVersion
from backend.app.eval.run_eval import _items_by_code
from backend.app.eval.runner import EvalPrediction
from backend.app.eval.runner import EvalSample
from backend.app.eval.runner import evidence_quality_counts
from backend.app.eval.runner import evaluate
from backend.app.eval.gating import evaluation_sha256
from backend.app.services.calibration.library import get_anchors
from backend.app.services.papers.ingestion import ingest_file
from backend.app.services.scoring.engine import score_paper
from backend.app.services.scoring.core.canonical import canonical_sha256
from backend.app.services.scoring.core.policy import build_corrected_thesis_policy
from backend.app.services.scoring.core.policy import compile_scoring_policy
from backend.app.services.scoring.core.policy import validate_weight_configuration
from backend.app.services.scoring.profiles.thesis import ThesisProfile

FILENAME_HEADERS = ("文件名", "文件", "论文", "filename", "file", "paper")
TOTAL_HEADERS = ("总分", "总成绩", "总评", "total", "score")


def build_labeled_eval(
    db: Session,
    rubric_id: str,
    papers_dir,
    scores_path,
    scorer=None,
    sample_ids_by_filename=None,
    rubric_version_id=None,
):
    """从论文文件夹 + 教师成绩表构建评估：逐篇导入+评分，与人工真值比对（设计§15）。"""
    rubric = db.get(Rubric, rubric_id)
    if rubric is None:
        raise ValueError("rubric not found: %s" % rubric_id)
    evaluation_policy, rubric_version = _evaluation_policy(
        db,
        rubric,
        rubric_version_id,
    )
    rubric_version_id = (
        rubric_version.id if rubric_version is not None else None
    )
    rubric_codes = {criterion.code for criterion in rubric.criteria}
    table = load_scores_table(scores_path)
    if not table:
        raise ValueError("成绩表为空或未识别到有效行")

    batch = GradingBatch(
        name="QWK评估-%s-%s" % (rubric.name, rubric.version),
        rubric_id=rubric_id,
        rubric_version_id=rubric_version_id,
        status="draft",
    )
    db.add(batch)
    db.flush()

    base_dir = Path(papers_dir)
    samples = []
    predictions = []
    errors = []
    completed_runs = 0
    review_required_runs = 0
    blocked_runs = 0
    invalid_evidence_items = 0
    evaluated_items = 0
    per_sample = []
    sample_ids_by_filename = sample_ids_by_filename or {}
    for entry in table:
        source = base_dir / entry["filename"]
        sample_id = sample_ids_by_filename.get(entry["filename"])
        error_identity = (
            {"sample_id": sample_id}
            if sample_id
            else {"filename": entry["filename"]}
        )
        if not source.exists():
            errors.append({**error_identity, "error": "未找到论文文件"})
            continue
        try:
            paper = ingest_file(db, batch.id, str(source), entry["filename"])
            if paper.status != "parsed":
                errors.append(
                    {
                        **error_identity,
                        "error": paper.error_message or "解析失败",
                    }
                )
                continue
            run = score_paper(db, paper.id, scorer=scorer)
        except Exception as exc:  # 单篇失败不中断整批
            errors.append({**error_identity, "error": str(exc)})
            continue
        completed_runs += 1
        review_required_runs += int(bool(run.need_manual_review))
        invalid_count, item_count = evidence_quality_counts(run)
        invalid_evidence_items += invalid_count
        evaluated_items += item_count
        system_items = _items_by_code(run)
        human_items = {
            code: value
            for code, value in entry["items"].items()
            if code in rubric_codes
        }
        private_sample = {
            "sample_id": sample_id or paper.id,
            "human_total": entry["total"],
            "system_total": (
                float(run.final_total_score)
                if run.final_total_score is not None
                else None
            ),
            "human_items": human_items,
            "system_items": system_items,
            "need_manual_review": bool(run.need_manual_review),
            "blocked": run.final_total_score is None,
            "invalid_evidence_item_count": invalid_count,
            "evaluated_item_count": item_count,
            "run_identity": _run_identity_projection(
                run,
                require_frozen=bool(rubric_version_id),
            ),
        }
        per_sample.append(private_sample)
        if run.final_total_score is None:
            blocked_runs += 1
            errors.append(
                {
                    **error_identity,
                    "error": "自动总分因 invalid/blocked 评分项为空；该样本未进入指标",
                }
            )
            continue
        sample_key = sample_id or paper.id
        samples.append(
            EvalSample(
                key=sample_key,
                human_total=entry["total"],
                human_items=human_items,
            )
        )
        predictions.append(
            EvalPrediction(
                key=sample_key,
                system_total=float(run.final_total_score),
                system_items=system_items,
            )
        )
    db.commit()

    report = evaluate(
        predictions,
        samples,
        scoring_policy=evaluation_policy,
    )
    report["errors"] = errors
    report["rubric_id"] = rubric_id
    report["rubric_version_id"] = rubric_version_id
    report["dataset_size"] = len(table)
    report["completed_runs"] = completed_runs
    report["review_rate"] = (
        review_required_runs / completed_runs if completed_runs else None
    )
    report["blocked_rate"] = blocked_runs / completed_runs if completed_runs else None
    report["invalid_evidence_item_count"] = invalid_evidence_items
    report["evaluated_item_count"] = evaluated_items
    report["invalid_evidence_rate"] = (
        invalid_evidence_items / evaluated_items if evaluated_items else None
    )
    report["per_sample"] = per_sample
    anchors = []
    for code in sorted(rubric_codes):
        selected = get_anchors(db, rubric.id, code)
        anchors.append({"criterion_code": code, "anchors": selected})
    report["anchors_identity"] = {
        "schema": "paper-grading/anchor-manifest@1",
        "count": sum(len(item["anchors"]) for item in anchors),
        "manifest_sha256": evaluation_sha256(anchors),
        # CalibrationAnchor currently lacks a source-sample identity, so this
        # cannot be machine-proven and must be supplied by the approval review.
        "holdout_exclusion_proven": False,
    }
    report["evaluation_identity"] = _evaluation_identity(
        report=report,
        policy=evaluation_policy,
        rubric=rubric,
        rubric_version=rubric_version,
        per_sample=per_sample,
    )
    return report


def _evaluation_policy(db, rubric, rubric_version_id):
    versions = db.scalars(
        select(RubricVersion).where(RubricVersion.rubric_id == rubric.id)
    ).all()
    if rubric_version_id is not None or versions:
        candidates = [
            version
            for version in versions
            if _is_consistently_published_version(db, rubric, version)
        ]
        if rubric_version_id is not None:
            candidates = [
                version
                for version in candidates
                if version.id == rubric_version_id
            ]
            if len(candidates) != 1:
                raise ValueError(
                    "evaluation rubric version identity is not consistently published"
                )
        elif len(candidates) != 1:
            raise ValueError(
                "formal evaluation requires exactly one immutable published RubricVersion"
            )
        version = candidates[0]
        policy = compile_scoring_policy(
            version.global_policy,
            total_score=rubric.total_score,
        )
        return policy, version
    weights = validate_weight_configuration(
        rubric.criteria,
        total_score=rubric.total_score,
    )
    return (
        build_corrected_thesis_policy(
            rubric.total_score,
            weights.mode,
        ),
        None,
    )


def _is_consistently_published_version(db, rubric, version):
    compilation = db.get(RubricCompilation, version.compilation_id)
    return bool(
        rubric.status == "published"
        and rubric.published_at is not None
        and compilation is not None
        and compilation.rubric_id == rubric.id
        and compilation.status == "validated"
        and compilation.reviewed_by is not None
        and compilation.reviewed_at is not None
        and compilation.reviewed_at == rubric.published_at
        and compilation.published_at == rubric.published_at
        and compilation.final_version_hash == version.version_hash
    )


def _evaluation_identity(*, report, policy, rubric, rubric_version, per_sample):
    identities = [
        item["run_identity"]
        for item in per_sample
        if isinstance(item.get("run_identity"), dict)
    ]
    first = identities[0] if identities else {}

    def unique(key):
        return sorted(
            {
                str(identity[key])
                for identity in identities
                if identity.get(key) is not None
            }
        )

    return {
        "schema_version": "grading-core/evaluation-identity@1",
        "rubric_id": rubric.id,
        "rubric_version_id": (
            rubric_version.id if rubric_version is not None else None
        ),
        "rubric_version_hash": (
            rubric_version.version_hash if rubric_version is not None else None
        ),
        "business_profile_key": (
            rubric_version.business_profile_key
            if rubric_version is not None
            else first.get("business_profile_key") or "thesis"
        ),
        "business_profile_version": first.get("business_profile_version"),
        "policy_hash": policy.policy_hash,
        "policy_snapshot_sha256": canonical_sha256(policy.to_mapping()),
        "grade_scale_sha256": report["grade_scale_sha256"],
        "rounding": {
            "mode": policy.rounding.mode,
            "digits": policy.rounding.digits,
        },
        "total_score": str(policy.aggregation.total_score),
        "execution_plan_hashes": unique("execution_plan_hash"),
        "document_snapshot_hashes": unique("document_snapshot_hash"),
        "runtime_identity_hashes": unique("runtime_identity_sha256"),
        "models": sorted(
            {
                (
                    identity.get("model_provider"),
                    identity.get("model_name"),
                    identity.get("model_version"),
                )
                for identity in identities
            },
            key=lambda item: tuple(
                "" if value is None else str(value) for value in item
            ),
        ),
    }


def _run_identity_projection(run, *, require_frozen=False):
    return ThesisProfile().build_eval_run_identity(
        run=run,
        require_frozen=require_frozen,
    )


def load_scores_table(path):
    """返回 [{filename, total, items:{code:score}}]。"""
    suffix = Path(path).suffix.lower()
    rows = _read_xlsx(path) if suffix in (".xlsx", ".xlsm") else _read_csv(path)
    return _rows_to_samples(rows)


def _read_xlsx(path):
    workbook = load_workbook(str(path), data_only=True)
    sheet = workbook.active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def _read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as handle:
        return [row for row in csv.reader(handle)]


def _rows_to_samples(rows):
    rows = [row for row in rows if row and any(_norm(cell) for cell in row)]
    if not rows:
        return []
    raw_header = ["" if cell is None else str(cell).strip() for cell in rows[0]]
    norm_header = [cell.lower() for cell in raw_header]  # 仅用于别名匹配
    filename_idx = _find_column(norm_header, FILENAME_HEADERS)
    total_idx = _find_column(norm_header, TOTAL_HEADERS)
    if filename_idx is None or total_idx is None:
        raise ValueError("成绩表需包含『文件名』列与『总分』列")

    # 其余列视为评分项 code → 保留原始大小写（与 rubric 的 criterion code 严格匹配）。
    code_columns = {
        index: raw_header[index]
        for index in range(len(raw_header))
        if index not in (filename_idx, total_idx) and raw_header[index]
    }

    samples = []
    for row in rows[1:]:
        filename = _cell(row, filename_idx)
        total = _num(_cell(row, total_idx))
        if not filename or total is None:
            continue
        items = {}
        for index, code in code_columns.items():
            value = _num(_cell(row, index))
            if value is not None:
                items[code] = value
        samples.append({"filename": str(filename).strip(), "total": total, "items": items})
    return samples


def _find_column(header, aliases):
    for index, cell in enumerate(header):
        if cell and any(alias.lower() in cell for alias in aliases):
            return index
    return None


def _cell(row, index):
    if index is None or index >= len(row):
        return None
    return row[index]


def _num(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip())
    except ValueError:
        return None


def _norm(value):
    return "" if value is None else str(value).strip().lower()
