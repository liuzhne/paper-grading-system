"""Authenticated deployment smoke test for the running HTTP service.

The caller is responsible for applying migrations and seeding the disposable
smoke database.  This script uses synthetic content only and exercises the
same login, upload, score, report, and export routes used by the Web UI.
"""

from __future__ import annotations

import argparse
from io import BytesIO
import json

from docx import Document
import httpx

from backend.app.core.config import settings


DEFAULT_RUBRIC_NAME = "本科毕业论文通用评分标准"
DEFAULT_BATCH_NAME = "2026 届论文评分开发批次"


def _parser():
    parser = argparse.ArgumentParser(prog="smoke_deployment")
    parser.add_argument("--base-url", default="http://127.0.0.1:8000")
    parser.add_argument("--rubric-name", default=DEFAULT_RUBRIC_NAME)
    parser.add_argument("--batch-name", default=DEFAULT_BATCH_NAME)
    return parser


def _response(response, label):
    try:
        response.raise_for_status()
    except httpx.HTTPStatusError as exc:
        body = response.text[:500]
        raise RuntimeError(
            f"{label} failed with HTTP {response.status_code}: {body}"
        ) from exc
    return response


def _sample_docx():
    document = Document()
    document.add_paragraph("部署冒烟测试论文")
    document.add_paragraph("姓名：测试用户")
    document.add_paragraph("学号：SMOKE-0001")
    document.add_paragraph("中文摘要")
    document.add_paragraph("本文使用合成内容验证部署后的上传、解析和评分链路。")
    document.add_paragraph("关键词：部署；冒烟测试；评分")
    document.add_paragraph("目录")
    document.add_paragraph("第一章 绪论")
    document.add_paragraph("本章说明研究背景、研究意义和相关工作。")
    document.add_paragraph("第二章 文献综述")
    document.add_paragraph("本章综述与系统部署和质量验证有关的研究。")
    document.add_paragraph("第三章 研究方法")
    document.add_paragraph("本文采用可重复的接口测试方法验证服务行为。")
    document.add_paragraph("第四章 实验结果与分析")
    document.add_paragraph("测试覆盖登录、上传、解析、评分、报告和导出。")
    document.add_paragraph("第五章 创新点")
    document.add_paragraph("测试只使用合成文档，不包含真实学生信息。")
    document.add_paragraph("结论")
    document.add_paragraph("部署链路在预期接口上完成闭环。")
    document.add_paragraph("参考文献")
    document.add_paragraph("[1] Synthetic deployment smoke fixture, 2026.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def run_smoke(*, base_url, rubric_name, batch_name):
    if not settings.AUTH_ENABLED:
        raise RuntimeError("deployment smoke requires AUTH_ENABLED=true")
    if not settings.AUTH_PASSWORD:
        raise RuntimeError("deployment smoke requires AUTH_PASSWORD")

    with httpx.Client(
        base_url=base_url.rstrip("/"),
        timeout=120.0,
        trust_env=False,
    ) as client:
        integrations = _response(
            client.get("/api/system/integrations"),
            "integration status",
        ).json()
        if not integrations.get("frontend", {}).get("static_web_ready"):
            raise RuntimeError("static Web frontend is not ready")

        auth_status = _response(
            client.get("/api/auth/status"),
            "auth status",
        ).json()
        if auth_status.get("auth_required") is not True:
            raise RuntimeError("protected deployment did not require authentication")

        login = _response(
            client.post(
                "/api/auth/login",
                json={
                    "username": settings.AUTH_USERNAME,
                    "password": settings.AUTH_PASSWORD,
                },
            ),
            "login",
        ).json()
        token = login.get("token")
        if not token:
            raise RuntimeError("login response did not contain a token")
        headers = {"Authorization": f"Bearer {token}"}

        _response(
            client.get("/api/auth/me", headers=headers),
            "authenticated identity",
        )
        readiness = _response(
            client.get("/api/system/ops-readiness", headers=headers),
            "operations readiness",
        ).json()
        if readiness.get("signals", {}).get("security", {}).get("status") != "pass":
            raise RuntimeError("deployment security readiness did not pass")

        rubrics = _response(
            client.get("/api/rubrics", headers=headers),
            "rubric listing",
        ).json()
        rubric = next(
            (item for item in rubrics if item.get("name") == rubric_name),
            None,
        )
        if rubric is None:
            raise RuntimeError(f"smoke rubric not found: {rubric_name}")

        batches = _response(
            client.get("/api/batches", headers=headers),
            "batch listing",
        ).json()
        batch = next(
            (
                item
                for item in batches
                if item.get("name") == batch_name
                and item.get("rubric_id") == rubric.get("id")
            ),
            None,
        )
        if batch is None:
            raise RuntimeError(f"smoke batch not found: {batch_name}")

        uploaded = _response(
            client.post(
                "/api/papers/upload",
                headers=headers,
                data={"batch_id": batch["id"]},
                files={
                    "file": (
                        "deployment-smoke.docx",
                        _sample_docx(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            ),
            "synthetic document upload",
        ).json()
        if uploaded.get("status") != "parsed":
            raise RuntimeError("synthetic document was not parsed")

        run = _response(
            client.post(
                f"/api/papers/{uploaded['id']}/score",
                headers=headers,
            ),
            "mock scoring",
        ).json()
        if run.get("final_total_score") is None:
            raise RuntimeError("scoring response did not contain a final score")

        report = _response(
            client.get(
                f"/api/scoring-runs/{run['id']}/report",
                headers=headers,
            ),
            "HTML report",
        )
        if "text/html" not in report.headers.get("content-type", ""):
            raise RuntimeError("report endpoint did not return HTML")

        export = _response(
            client.get(
                f"/api/batches/{batch['id']}/export.xlsx",
                headers=headers,
            ),
            "Excel export",
        )
        if "spreadsheetml" not in export.headers.get("content-type", ""):
            raise RuntimeError("batch export endpoint did not return XLSX")

    return {
        "schema_version": "deployment-smoke@1",
        "status": "passed",
        "frontend": "ready",
        "authentication": "passed",
        "database": readiness.get("signals", {}).get("database", {}).get("dialect"),
        "llm_provider": integrations.get("llm", {}).get("provider"),
        "paper_status": uploaded.get("status"),
        "scoring_run_id": run.get("id"),
        "report": "passed",
        "excel_export": "passed",
    }


def main(argv=None):
    args = _parser().parse_args(argv)
    result = run_smoke(
        base_url=args.base_url,
        rubric_name=args.rubric_name,
        batch_name=args.batch_name,
    )
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
